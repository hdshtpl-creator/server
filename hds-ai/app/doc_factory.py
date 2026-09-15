"""
doc_factory.py — "Tạo BỘ file" từ khung chat: AI đọc hồ sơ đính kèm + lịch sử
hội thoại, tự lên danh sách văn bản cần soạn (biên bản nghiệm thu, giấy đề
nghị thanh toán từng đợt…) rồi tạo TỪNG file .docx — yêu cầu chủ dự án
26/08/2026, mô phỏng luồng "gửi hợp đồng đợt 1, nhờ làm đợt 2-3 + nghiệm thu"
trên Claude/ChatGPT.

Khác luồng "Tạo file mẫu" (template_fill — điền chủ thể vào ĐÚNG một mẫu,
chính xác tối đa): ở đây model được chủ động hơn — quyết định số file, chọn
khuôn cho từng file (mẫu kho đã chọn, file .docx đã tải lên, hoặc soạn mới),
và LẤY DỮ LIỆU TỪ CHÍNH HỒ SƠ ĐÍNH KÈM (số hợp đồng, số tài khoản, các đợt
thanh toán…). Đổi lại, câu trả lời BẮT BUỘC có mục "CHỖ AI TỰ QUYẾT ĐỊNH —
KIỂM TRA BẮT BUỘC" liệt kê mọi suy đoán, và lời nhắc rà soát toàn văn — vì
văn bản sẽ gửi cho khách, trách nhiệm cuối vẫn là người soát.

Mở rộng 15/09/2026 (yêu cầu chủ dự án):
  - BỘ MẪU (app/bo_mau.py): chọn một bộ .docx mẫu tải lên từ Quản trị → kế
    hoạch KHÔNG hỏi model nữa mà là "mỗi file mẫu trong bộ → một file điền" —
    bộ 30-100 file mà để model tự kê danh sách JSON là sót/lặp. Gọi tên bộ
    ngay trong câu chat cũng được (bo_mau.match_set).
  - KHÔNG TRẦN số file (MAX_FILES = 0); admin đặt trần qua cài đặt
    doc_factory_max_files nếu máy yếu.
  - Dữ liệu điền lấy thêm từ LỊCH SỬ HỘI THOẠI (tóm tắt + các lượt gần nhất):
    "đang chat về khách Minh, bảo tạo bộ hồ sơ" là có dữ liệu ngay.
  - Nhận lệnh từ chat thường (detect_request) — không cần bấm nút.
  - Từ 2 file trở lên: thêm gói .zip cả bộ để tải một lần.

Hai cách tạo một file:
  - dien_mau: mở khuôn .docx (mẫu kho, file tải lên, hoặc file trong bộ mẫu),
    thay chuỗi theo run — giữ nguyên định dạng gốc (tái dùng toàn bộ máy của
    template_fill).
  - soan_moi: model viết Markdown, render .docx bằng drafting.render_docx
    (Times New Roman 12pt) — cho văn bản không có khuôn sẵn.
"""
import io
import json
import re
import unicodedata
import zipfile
from pathlib import Path

from app import db, settings, template_fill

# 0 = KHÔNG giới hạn số file một lượt (chủ dự án 15/09/2026: "nới từ 8 lên vô
# hạn"). Admin muốn đặt trần thì sửa cài đặt doc_factory_max_files trên web.
MAX_FILES = 0
MAX_ATTACH_PROMPT_CHARS = 12_000
# Tổng ngân sách ký tự cho MỌI file đính kèm trong một prompt — giữ dưới trần
# llm_num_ctx (32768 token) sau khi trừ văn bản mẫu (24k ký tự) + phần chỉ dẫn.
ATTACH_TOTAL_BUDGET = 40_000
MAX_PLAN_DESC_CHARS = 500
# Lịch sử hội thoại đưa vào prompt: giữ các lượt GẦN NHẤT khi vượt ngân sách.
HISTORY_BUDGET = 8_000
MAX_HISTORY_MSG_CHARS = 1_500
MAX_SUMMARY_CHARS = 2_500


def max_files() -> int:
    """Trần số file một lượt: 0 = không giới hạn."""
    try:
        return max(0, settings.get_int("doc_factory_max_files", MAX_FILES))
    except Exception:  # noqa: BLE001 — không có CSDL (test) thì dùng hằng
        return MAX_FILES


# ---------------------------------------------------------------------------
# Nhận lệnh từ chat thường
# ---------------------------------------------------------------------------
def _fold(s: str) -> str:
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "d")
    return re.sub(r"[\s_]+", " ", s).strip().lower()


_VERB = r"(?:tao|lam|dien|soan|xuat|sinh|lap|chuan bi)"
_DOI_TUONG_BO = (r"(?:bo ho so|bo file|bo tai lieu|bo van ban|bo giay to|bo mau|"
                 r"bo hop dong|tron bo|ca bo)")
# "tạo bộ hồ sơ …", "làm giúp tôi các file …", "xuất hồ sơ cho khách …"
_RE_TAO_DAU_CAU = re.compile(
    r"^\s*(?:hay\s+|nho\s+|giup toi\s+|giup\s+|ban\s+|toi muon\s+|toi can\s+|"
    r"can\s+|muon\s+|lam on\s+)?"
    + _VERB + r"\s+"
    r"(?:giup toi\s+|giup\s+|ho toi\s+|ho\s+|cho toi\s+|toi\s+|mot\s+|ca\s+|"
    r"toan bo\s+|day du\s+|luon\s+|lai\s+|san\s+|nhanh\s+|ngay\s+|them\s+)*"
    r"(?:" + _DOI_TUONG_BO + r"|ho so|cac file|file|giay to|cac van ban)\b(.*)$",
    re.S)
# "điền thông tin khách vào bộ mẫu Thuê nhà" — động từ ở đâu cũng được, miễn
# câu có nhắc tới MỘT BỘ.
_RE_TAO_CO_BO = re.compile(r"\b" + _VERB + r"\b.*\b" + _DOI_TUONG_BO + r"\b", re.S)
# Câu HỎI về cách làm ("tạo bộ hồ sơ như thế nào?") là tra cứu, không phải lệnh.
_RE_HOI = re.compile(
    r"\?\s*$|\b(?:the nao|lam sao|cach nao|nhu the nao|ra sao|duoc khong|"
    r"co duoc khong|huong dan|la gi)\b")
# "tạo file excel/pdf/ảnh" — không phải bộ văn bản .docx; tránh cướp câu hỏi
# thường về file kỹ thuật.
_RE_KHONG_PHAI = re.compile(r"\b(?:excel|xlsx|pdf|anh|hinh|zip|json|csv|powerpoint|pptx)\b")


def detect_request(question: str) -> dict | None:
    """Câu chat có phải lệnh TẠO BỘ FILE không (không cần bấm nút).

    Trả {"nhac_bo_mau": bool} hoặc None. Khuôn cố ý KHÁC chat_draft: bên kia
    là "tạo <loại giấy> cho <tên>" (một văn bản, đi luồng bản nháp), bên này
    là "tạo bộ hồ sơ / các file / file … " (nhiều văn bản, điền khuôn .docx).
    """
    folded = _fold(question)
    if not folded or _RE_HOI.search(folded) or _RE_KHONG_PHAI.search(folded):
        return None
    if _RE_TAO_DAU_CAU.search(folded) or _RE_TAO_CO_BO.search(folded):
        return {"nhac_bo_mau": "bo mau" in folded}
    return None


# ---------------------------------------------------------------------------
# Dữ liệu từ lịch sử hội thoại
# ---------------------------------------------------------------------------
def build_history_block(summary: str, history, budget: int = HISTORY_BUDGET) -> str:
    """Khối 'DỮ LIỆU TỪ HỘI THOẠI' cho prompt: tóm tắt phần đầu + các lượt gần
    nhất. Vượt ngân sách thì bỏ lượt CŨ trước — thông tin khách thường nằm ở
    những lượt vừa nói."""
    parts = []
    summary = " ".join((summary or "").split())
    if summary:
        parts.append("TÓM TẮT PHẦN ĐẦU HỘI THOẠI:\n" + summary[:MAX_SUMMARY_CHARS])
    msgs = []
    for item in history or []:
        try:
            role, content = item[0], item[1]
        except (TypeError, IndexError, KeyError):
            continue
        body = " ".join(str(content or "").split())
        if not body:
            continue
        if len(body) > MAX_HISTORY_MSG_CHARS:
            body = body[:MAX_HISTORY_MSG_CHARS] + "…"
        label = "Người dùng" if role == "user" else "Trợ lý"
        msgs.append(f"- {label}: {body}")
    kept, total = [], 0
    for m in reversed(msgs):
        if total + len(m) > budget:
            break
        kept.append(m)
        total += len(m)
    kept.reverse()
    if kept:
        parts.append("CÁC LƯỢT TRAO ĐỔI GẦN NHẤT:\n" + "\n".join(kept))
    return "\n\n".join(parts)


def user_history_text(history) -> str:
    """Phần NGƯỜI DÙNG đã gõ trong lịch sử — vùng tin cậy để cắm cờ ⚠ (câu
    trợ lý có thể chép từ file khách gửi, không tính là tin cậy)."""
    out = []
    for item in history or []:
        try:
            if item[0] == "user":
                out.append(str(item[1] or ""))
        except (TypeError, IndexError, KeyError):
            continue
    return " ".join(out)


# ---------------------------------------------------------------------------
# Kế hoạch: hỏi model cần tạo những file nào
# ---------------------------------------------------------------------------
def build_plan_prompt(question: str, attachments, kho_template_title: str | None,
                      du_lieu_hoi_thoai: str = "", max_files: int = 0):
    """attachments: [(filename, text, has_docx)] — has_docx = còn bản gốc làm khuôn."""
    lines = []
    for fname, text, has_docx in attachments:
        tag = " (CÓ BẢN .docx GỐC — dùng được làm khuôn)" if has_docx else ""
        body = (text or "").strip()[:MAX_ATTACH_PROMPT_CHARS]
        lines.append(f"### FILE ĐÍNH KÈM «{fname}»{tag}\n{body}")
    if du_lieu_hoi_thoai:
        lines.append("### DỮ LIỆU TỪ HỘI THOẠI (những gì hai bên đã trao đổi)\n"
                     + du_lieu_hoi_thoai)
    khuon_note = []
    if kho_template_title:
        khuon_note.append(f'- "mau_kho": mẫu đã chọn trong kho — «{kho_template_title}»')
    for fname, _text, has_docx in attachments:
        if has_docx:
            khuon_note.append(f'- "{fname}": file .docx người dùng tải lên')
    khuon_note.append('- "soan_moi": không có khuôn phù hợp, soạn mới từ đầu')
    tran = (f"- Tối đa {max_files} file; chỉ tạo file mà yêu cầu thật sự cần.\n"
            if max_files and max_files > 0
            else "- Tạo ĐỦ số file mà yêu cầu thật sự cần, không thêm file thừa.\n")

    system = ("Bạn là trợ lý soạn thảo của công ty luật. Bạn chỉ trả về DUY "
              "NHẤT một khối JSON, không giải thích gì bên ngoài JSON.")
    prompt = (
        f"YÊU CẦU CỦA NGƯỜI DÙNG: {question.strip()}\n\n"
        + ("\n\n".join(lines) + "\n\n" if lines else "")
        + "Nhiệm vụ: lên DANH SÁCH các văn bản cần tạo để đáp ứng đúng yêu cầu "
          "(ví dụ: mỗi đợt thanh toán một giấy đề nghị, mỗi mốc bàn giao một "
          "biên bản nghiệm thu). Với mỗi văn bản, chọn khuôn trong các lựa "
          "chọn sau:\n" + "\n".join(khuon_note) + "\n\n"
        "Trả về JSON đúng khung:\n"
        "{\n"
        '  "files": [{"ten_file": "DNTT-02 Giấy đề nghị thanh toán đợt 2",\n'
        '             "khuon": "mau_kho" | "<tên file .docx tải lên>" | "soan_moi",\n'
        '             "yeu_cau": "văn bản này chứa gì, dữ liệu lấy từ đâu, số tiền/đợt nào"}],\n'
        '  "ghi_chu": "các điểm bạn TỰ QUYẾT ĐỊNH hoặc còn thiếu dữ liệu, để người dùng kiểm tra"\n'
        "}\n"
        "Quy tắc:\n"
        + tran +
        "- Dữ liệu (số hợp đồng, tên bên, số tài khoản, số tiền, tỷ lệ đợt) "
        "phải lấy từ file đính kèm, hội thoại hoặc câu lệnh; thiếu thì ghi vào "
        "ghi_chu, không bịa.\n"
        "- Ngày tháng chưa biết thì để trống và ghi chú, không tự đặt.\n"
    )
    return prompt, system


def parse_plan(raw_answer: str, max_files: int = 0):
    """Bóc + kiểm tra kế hoạch. Trả về (files, ghi_chu). Raise ValueError khi hỏng.
    max_files: 0 = không cắt."""
    payload = template_fill.parse_llm_json(raw_answer)
    files = []
    items = payload.get("files") or []
    if max_files and max_files > 0:
        items = items[:max_files]
    for item in items:
        if not isinstance(item, dict):
            continue
        ten = str(item.get("ten_file") or "").strip()
        if not ten:
            continue
        khuon = str(item.get("khuon") or "soan_moi").strip()
        yeu_cau = str(item.get("yeu_cau") or "").strip()[:MAX_PLAN_DESC_CHARS]
        files.append({"ten_file": ten[:120], "khuon": khuon, "yeu_cau": yeu_cau})
    if not files:
        raise ValueError("Kế hoạch không có file nào")
    return files, str(payload.get("ghi_chu") or "").strip()


def plan_tu_bo_mau(bo: dict, file_ids=None, max_files: int = 0):
    """Kế hoạch TẤT ĐỊNH từ bộ mẫu: mỗi file mẫu → một file điền (giữ thứ tự
    trong bộ). file_ids (nếu có) thu hẹp về các file người dùng tích chọn.
    Không hỏi model — bộ 30-100 file mà để model kê JSON là sót/lặp/sai tên."""
    wanted = {int(i) for i in (file_ids or []) if i is not None}
    files = []
    for f in bo.get("files") or []:
        if wanted and int(f["id"]) not in wanted:
            continue
        ten = Path(f["ten_file"]).stem or f["ten_file"]
        files.append({"ten_file": ten[:120], "khuon": f"bo_mau:{f['id']}",
                      "yeu_cau": "", "duong_dan": f["duong_dan"]})
    if max_files and max_files > 0:
        files = files[:max_files]
    return files


# ---------------------------------------------------------------------------
# Soạn mới bằng Markdown
# ---------------------------------------------------------------------------
def clean_markdown(raw_answer: str) -> str:
    body = re.sub(r"<think>.*?</think>", "", raw_answer or "", flags=re.DOTALL)
    body = re.sub(r"^```(?:markdown|md)?\s*$", "", body.strip(),
                  flags=re.MULTILINE)
    return body.strip()


def build_draft_prompt(ten_file: str, yeu_cau: str, question: str,
                       attachments, ghi_chu_plan: str, du_lieu_hoi_thoai: str = ""):
    parts = []
    for fname, text, _has in attachments:
        body = (text or "").strip()[:MAX_ATTACH_PROMPT_CHARS]
        if body:
            parts.append(f"### DỮ LIỆU TỪ «{fname}»\n{body}")
    if du_lieu_hoi_thoai:
        parts.append("### DỮ LIỆU TỪ HỘI THOẠI\n" + du_lieu_hoi_thoai)
    system = ("Bạn là chuyên viên soạn thảo văn bản của công ty luật Việt Nam. "
              "Bạn trả về DUY NHẤT nội dung văn bản ở dạng Markdown (tiêu đề "
              "dùng #, ##; danh sách dùng -), không lời dẫn, không giải thích.")
    prompt = (
        f"Soạn văn bản: {ten_file}\n"
        f"Yêu cầu riêng: {yeu_cau or '(theo thông lệ loại văn bản này)'}\n"
        f"Bối cảnh từ người dùng: {question.strip()}\n\n"
        + ("\n\n".join(parts) + "\n\n" if parts else "")
        + "Quy tắc bắt buộc:\n"
        "- Đúng thể thức văn bản Việt Nam: quốc hiệu, tiêu ngữ, tên văn bản, "
        "căn cứ, nội dung từng điều/mục, phần ký tên hai bên (nếu là văn bản "
        "song phương).\n"
        "- Mọi con số, tên bên, số tài khoản, số hợp đồng phải lấy từ DỮ LIỆU "
        "ở trên hoặc câu lệnh; thông tin chưa có thì ghi [CẦN BỔ SUNG: mô tả] "
        "— tuyệt đối không bịa.\n"
        "- Ngày tháng chưa chốt thì viết 'ngày … tháng … năm 2026'.\n"
    )
    return prompt, system


# ---------------------------------------------------------------------------
# Nguyên liệu
# ---------------------------------------------------------------------------
def _attachments(conversation_id, use_temp):
    """[(filename, text, source_path|None)] từ temp_files của hội thoại."""
    if not conversation_id or not use_temp:
        return []
    with db.session(role="internal", admin=True) as conn:
        with conn.cursor() as cur:
            cur.execute("""SELECT filename, content, source_path FROM temp_files
                            WHERE conversation_id=%s AND expires_at > now()
                            ORDER BY id""", (conversation_id,))
            return [(r[0], r[1], r[2]) for r in cur.fetchall()]


def _load_skeleton(path_str: str):
    """Mở khuôn .docx đã lưu trong DATA_WORK (chat_uploads hoặc bo_mau) — rào
    trong DATA_WORK, đường dẫn đọc từ CSDL không được dắt ra ngoài."""
    import docx
    path = Path(path_str)
    if not path.is_absolute():
        path = Path.cwd() / path
    work = template_fill.DATA_WORK
    work_root = (work if work.is_absolute() else Path.cwd() / work).resolve()
    resolved = path.resolve(strict=True)
    resolved.relative_to(work_root)
    return docx.Document(str(resolved))


def bundle_zip(entries) -> bytes:
    """Gói các file kết quả thành một .zip trong bộ nhớ. entries: [(tên trong
    zip, Path)]. Tên trùng thì đánh số để không ghi đè nhau."""
    buf = io.BytesIO()
    seen = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for arcname, path in entries:
            base, n = arcname, 1
            while arcname in seen:
                n += 1
                stem, suffix = Path(base).stem, Path(base).suffix
                arcname = f"{stem} ({n}){suffix}"
            seen.add(arcname)
            zf.write(str(path), arcname)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Luồng chính
# ---------------------------------------------------------------------------
def handle(question, *, user_id, dept_ids=None, is_banqt=False, can_finance=False,
           conversation_id=None, template_doc_id=None, model=None, on_status=None,
           role_level=None, dept_codes=None, use_temp=True,
           bo_mau_id=None, bo_mau_file_ids=None, history=None, summary=""):
    """Tạo bộ file, trả về dict kiểu direct-answer cho rag.prepare.

    bo_mau_id: bộ mẫu người dùng chọn trên giao diện (hoặc None — khi đó thử
    khớp TÊN bộ trong câu lệnh). bo_mau_file_ids: thu hẹp về vài file trong bộ.
    history/summary: lịch sử hội thoại (rag.get_history/get_summary) — dữ liệu
    khách đã nói trong chat được dùng để điền.
    """
    from app import bo_mau as bo_mau_mod
    from app import models
    from app.drafting import render_docx

    def note(label):
        if on_status:
            try:
                on_status(label)
            except Exception:  # noqa: BLE001 — tiến trình chỉ để hiển thị
                pass

    def _fail(answer_md):
        return {"answer": answer_md, "answer_mode": "structured",
                "grounding_status": "not_applicable", "evidence": [], "state": {}}

    # Bộ file luôn mang dữ liệu định danh của một người/một khách cụ thể —
    # model_soan_thao giữ nó ở lại máy chủ trừ khi admin đã mở phạm vi rộng nhất.
    chosen_model = models.model_soan_thao(
        (model or "").strip() or models.effective_llm_model())

    def llm_full(prompt, system, temperature):
        """Gọi model qua STREAM rồi ghép lại — một lượt tạo bộ file có tới
        1+N lời gọi model, mỗi lời gọi không-stream chỉ có trần 300 giây tường;
        stream tính sự sống theo từng mẩu chữ nên văn bản dài trên máy chậm
        không chết oan giữa chừng (rà soát 26/08/2026)."""
        from app.models import llm_stream
        return "".join(llm_stream(prompt, system=system,
                                  temperature=temperature, model=chosen_model))

    tran = max_files()

    # ---- Bộ mẫu: chọn trên giao diện, hoặc gọi tên trong câu ---------------
    bo = None
    if bo_mau_id:
        bo = bo_mau_mod.get_set(bo_mau_id, role=role_level, dept_ids=dept_ids,
                                is_banqt=is_banqt)
        if not bo:
            return _fail("Mình không mở được bộ mẫu đã chọn (đã bị gỡ hoặc tài "
                         "khoản chưa được xem bộ này). Chọn bộ khác hoặc bỏ chọn "
                         "rồi thử lại nhé.")
    else:
        try:
            visible = bo_mau_mod.list_sets(role=role_level, dept_ids=dept_ids,
                                           is_banqt=is_banqt)
        except Exception:  # noqa: BLE001 — bảng chưa migrate thì coi như không có bộ
            visible = []
        bo = bo_mau_mod.match_set(question, visible)
        if bo is None and bo_mau_mod.nhac_bo_mau(question):
            if visible:
                ten_bo = "\n".join(f"- **{s['ten']}** ({s.get('so_file', 0)} file)"
                                   for s in visible[:30])
                return _fail("Mình chưa nhận ra bộ mẫu nào trong câu. Các bộ bạn "
                             f"dùng được:\n{ten_bo}\n\nGõ lại kèm ĐÚNG tên bộ (ví "
                             f"dụ: \"tạo bộ hồ sơ theo bộ mẫu {visible[0]['ten']} "
                             "cho khách …\"), hoặc chọn bộ ở nút **Bộ mẫu** dưới "
                             "khung chat.")
            return _fail("Chưa có bộ mẫu nào trong hệ thống. Người có quyền duyệt "
                         "vào **Quản trị → Bộ mẫu hồ sơ** để tạo bộ và tải các file "
                         ".docx mẫu lên, rồi quay lại chat gọi tên bộ đó.")
    if bo is not None and not (bo.get("files") or []):
        return _fail(f"Bộ mẫu «{bo['ten']}» chưa có file .docx nào. Vào **Quản trị → "
                     "Bộ mẫu hồ sơ** tải file mẫu vào bộ trước nhé.")

    # ---- Nguyên liệu -----------------------------------------------------
    note("Đang đọc hồ sơ đính kèm…")
    temp_rows = _attachments(conversation_id, use_temp)
    attachments = [(fname, text, bool(src)) for fname, text, src in temp_rows]
    docx_uploads = {fname: src for fname, _text, src in temp_rows if src}
    du_lieu_hoi_thoai = build_history_block(summary, history)

    # NGÂN SÁCH TỔNG cho phần hồ sơ trong prompt: trần vật lý là llm_num_ctx —
    # vượt là Ollama cắt PHẦN ĐẦU prompt (đúng chỗ đặt văn bản mẫu) một cách
    # câm lặng. Nhiều file đính kèm thì chia nhau ngân sách, và phải NÓI RA
    # việc cắt trong mục kiểm tra.
    attachments_truncated = []
    if attachments:
        share = min(MAX_ATTACH_PROMPT_CHARS,
                    max(2_000, ATTACH_TOTAL_BUDGET // len(attachments)))
        cut = []
        for fname, text, has_docx in attachments:
            body = text or ""
            if len(body) > share:
                attachments_truncated.append(fname)
            cut.append((fname, body[:share], has_docx))
        attachments = cut

    kho_doc = kho_text = None
    template_warning = None
    if template_doc_id:
        row = template_fill._template_row(
            template_doc_id, dept_ids, is_banqt, can_finance,
            role_level=role_level, dept_codes=dept_codes)
        if not row:
            return _fail("Mình không mở được file mẫu đã chọn (không còn trong "
                         "kệ mẫu hoặc tài khoản chưa được xem). Bỏ chọn mẫu "
                         "hoặc chọn mẫu khác rồi thử lại nhé.")
        _kid, kho_title, _ktype, kho_source = row
        if kho_source and str(kho_source).lower().endswith(".docx"):
            try:
                kho_doc_obj, _p = template_fill._load_docx(kho_source)
                kho_doc = (kho_title, kho_source)
                kho_text = template_fill.document_text(kho_doc_obj)
            except (ValueError, FileNotFoundError, OSError):
                template_warning = (f"Mẫu đã chọn «{kho_title}» không mở được "
                                    "(tệp gốc hỏng hoặc không còn) — các file "
                                    "được SOẠN MỚI, KHÔNG theo mẫu kho.")
        else:
            # Đừng im lặng bỏ mẫu người dùng đã chọn — họ sẽ tưởng bộ file đi
            # theo khuôn chuẩn của công ty trong khi thực tế là soạn mới.
            template_warning = (f"Mẫu đã chọn «{kho_title}» không phải file "
                                ".docx nên không dùng làm khuôn được — các "
                                "file được SOẠN MỚI, KHÔNG theo mẫu kho.")
    if (not attachments and not kho_doc and bo is None and not du_lieu_hoi_thoai
            and not question.strip()):
        return _fail("Chưa có gì để tạo file: đính kèm hồ sơ, chọn bộ mẫu/mẫu, hoặc "
                     "mô tả rõ cần những văn bản gì nhé.")

    # VÙNG TIN CẬY để cắm cờ ⚠: câu lệnh + những gì NGƯỜI DÙNG đã gõ trong hội
    # thoại + các trường bóc bằng regex truy vết được. Luồng này KHÔNG cách ly
    # giá trị ngoài vùng (dữ liệu điền vốn lấy từ hồ sơ đính kèm theo đúng yêu
    # cầu), nhưng giá trị nào ngoài vùng sẽ mang cờ nhắc đối chiếu bản gốc.
    trusted_fields = {}
    for _fname, text, _h in attachments:
        try:
            from app import autofill
            trusted_fields = autofill.merge_missing(
                trusted_fields, autofill.extract_person_fields(text))
        except Exception:  # noqa: BLE001
            pass
    trusted_fold = template_fill._fold_for_match(
        question + " " + user_history_text(history) + " "
        + " ".join(v for v in trusted_fields.values() if v))

    # ---- Bước 1: kế hoạch ------------------------------------------------
    ghi_chu_plan = ""
    if bo is not None:
        note(f"Đang lấy danh sách file trong bộ mẫu «{bo['ten']}»…")
        plan_files = plan_tu_bo_mau(bo, bo_mau_file_ids, tran)
        if not plan_files:
            return _fail("Không có file nào trong bộ mẫu được chọn để điền. Bỏ tích "
                         "chọn từng file hoặc chọn lại bộ nhé.")
    else:
        note("Đang lên danh sách file cần tạo…")
        plan_prompt, plan_system = build_plan_prompt(
            question, attachments, kho_doc[0] if kho_doc else None,
            du_lieu_hoi_thoai, tran)
        raw_plan = llm_full(plan_prompt, plan_system, 0.0)
        try:
            plan_files, ghi_chu_plan = parse_plan(raw_plan, tran)
        except (ValueError, json.JSONDecodeError):
            return _fail("Model chưa lên được danh sách file hợp lệ. Bạn mô tả rõ "
                         "hơn: cần những văn bản gì, mỗi văn bản chứa dữ liệu nào?")

    # ---- Bước 2: tạo từng file ------------------------------------------
    # LƯU Ý AN TOÀN: khác luồng "Tạo file mẫu", ở đây dữ liệu điền LẤY TỪ CHÍNH
    # HỒ SƠ ĐÍNH KÈM / HỘI THOẠI (người dùng yêu cầu vậy — số hợp đồng, số tài
    # khoản, các đợt tiền đều nằm trong đó), nên KHÔNG áp được chốt
    # trusted_text. Bù lại:
    # (1) mọi thay thế vẫn qua bộ lọc cấu trúc của sanitize_replacements;
    # (2) câu trả lời liệt kê từng chỗ đã thay + mục "TỰ QUYẾT ĐỊNH" bắt buộc;
    # (3) file là hàng tạm 24 giờ, không vào kho tri thức.
    results = []      # (ten_file, token, out_path, mo_ta, warnings, details)
    total = len(plan_files)
    party_note = (f"YÊU CẦU CHUNG: {question.strip()}\n"
                  f"GHI CHÚ KẾ HOẠCH: {ghi_chu_plan[:1000]}" if ghi_chu_plan
                  else f"YÊU CẦU CHUNG: {question.strip()}")
    du_lieu_dien = "\n\n".join(f"NỘI DUNG «{f}»:\n{t}" for f, t, _h in attachments)
    if du_lieu_hoi_thoai:
        du_lieu_dien += ("\n\n" if du_lieu_dien else "") + \
            "DỮ LIỆU TỪ HỘI THOẠI:\n" + du_lieu_hoi_thoai
    for i, spec in enumerate(plan_files, 1):
        ten_file, khuon, yeu_cau = spec["ten_file"], spec["khuon"], spec["yeu_cau"]
        note(f"Đang soạn file {i}/{total}: {ten_file}…")
        warnings = []
        details = []   # bảng đối chiếu «cũ» → «mới» của CHÍNH file này
        skeleton_kind = None
        if khuon == "mau_kho" and kho_doc:
            skeleton_path, skeleton_text, skeleton_kind = kho_doc[1], kho_text, "kho"
        elif khuon.startswith("bo_mau:") and spec.get("duong_dan"):
            skeleton_path, skeleton_text, skeleton_kind = spec["duong_dan"], None, "bo"
        elif khuon in docx_uploads:
            skeleton_path, skeleton_text, skeleton_kind = docx_uploads[khuon], None, "upload"
        else:
            skeleton_path = skeleton_text = None
            if khuon not in ("soan_moi", "", None):
                warnings.append(f"khuôn «{khuon}» không mở được — đã soạn mới")

        def _flag(value):
            """⚠ cho giá trị KHÔNG nằm trong câu lệnh/lịch sử người dùng/trường
            đã bóc — tức là model tự lấy từ file đính kèm hoặc câu trả lời
            cũ; người soát phải đối chiếu bản gốc."""
            if template_fill._fold_for_match(value) in trusted_fold:
                return ""
            return " ⚠ (lấy từ file đính kèm/hội thoại — đối chiếu bản gốc)"

        try:
            if skeleton_path:
                # Điền vào khuôn giữ định dạng
                if skeleton_kind == "kho":
                    doc, _p = template_fill._load_docx(skeleton_path)
                elif skeleton_kind == "bo":
                    import docx as _docx
                    doc = _docx.Document(str(bo_mau_mod.resolve_path(skeleton_path)))
                else:
                    doc = _load_skeleton(skeleton_path)
                if skeleton_text is None:
                    skeleton_text = template_fill.document_text(doc)
                placeholders = template_fill.scan_placeholders(doc)
                fill_prompt, fill_system = template_fill.build_fill_prompt(
                    skeleton_text,
                    f"FILE CẦN TẠO: {ten_file}\n"
                    + (f"YÊU CẦU RIÊNG: {yeu_cau}\n" if yeu_cau else "")
                    + f"{party_note}\n\n{du_lieu_dien}",
                    placeholders)
                raw_fill = llm_full(fill_prompt, fill_system, 0.0)
                payload = template_fill.parse_llm_json(raw_fill)
                ph_values = payload.get("placeholders") or {}
                ph_pairs = []
                for literal, key in placeholders:
                    for k, v in ph_values.items():
                        if template_fill.normalize_key(str(k)) == key and str(v).strip():
                            ph_pairs.append((literal, str(v).strip()))
                            break
                reps, dropped = template_fill.sanitize_replacements(
                    payload.get("replacements"), skeleton_text)
                counts = template_fill.apply_replacements(doc, ph_pairs + reps)
                changed = sum(counts.values())
                # BẢNG ĐỐI CHIẾU TỪNG CHỖ THAY — đây là chốt bù cho việc luồng
                # này không cách ly giá trị ngoài vùng tin cậy: người soát phải
                # NHÌN THẤY từng chỗ đã thay, không chỉ một con số tổng.
                for literal, value in ph_pairs:
                    if counts.get(literal, 0):
                        details.append(f"chỗ trống `{literal}` → **{value}**"
                                       f" ({counts[literal]} chỗ){_flag(value)}")
                for old, new in reps:
                    n = counts.get(old, 0)
                    if n:
                        details.append(f"«{old}» → **{new}** ({n} chỗ){_flag(new)}")
                for old, reason in dropped:
                    warnings.append(f"bỏ «{old[:60]}»: {reason}")
                # Chỗ trống {{…}} trong mẫu mà model không điền được: NÓI RA,
                # đừng để người soát tự dò 30 trang tìm ngoặc còn sót.
                con_trong = [lit for lit, _k in placeholders
                             if not counts.get(lit, 0)]
                if con_trong:
                    warnings.append("chưa có dữ liệu cho "
                                    + ", ".join(f"`{c}`" for c in con_trong[:12])
                                    + (" …" if len(con_trong) > 12 else "")
                                    + " — vẫn để nguyên trong file")
                extra_note = str(payload.get("ghi_chu") or "").strip()
                if extra_note:
                    warnings.append(extra_note)
                token, out_path = template_fill.save_filled(doc, ten_file)
                mo_ta = f"điền khuôn ({changed} chỗ thay)"
            else:
                # Soạn mới bằng Markdown → .docx
                d_prompt, d_system = build_draft_prompt(
                    ten_file, yeu_cau, question, attachments, ghi_chu_plan,
                    du_lieu_hoi_thoai)
                raw_md = llm_full(d_prompt, d_system, 0.1)
                md = clean_markdown(raw_md)
                if not md:
                    raise ValueError("model trả nội dung rỗng")
                n_holes = md.count("[CẦN BỔ SUNG")
                payload_bytes = render_docx(md, ten_file, [])
                token, out_path = template_fill.save_filled_bytes(payload_bytes, ten_file)
                mo_ta = ("soạn mới"
                         + (f", {n_holes} chỗ [CẦN BỔ SUNG]" if n_holes else ""))
        except Exception as exc:  # noqa: BLE001 — một file hỏng không huỷ cả bộ
            results.append((ten_file, None, None,
                            f"KHÔNG TẠO ĐƯỢC ({type(exc).__name__})", warnings, []))
            continue
        results.append((ten_file, token, out_path, mo_ta, warnings, details))

    template_fill.cleanup_old_fills()
    made = [r for r in results if r[1]]
    if not made:
        return _fail("Mình chưa tạo được file nào từ yêu cầu này. Bạn thử mô "
                     "tả lại danh sách file cần soạn và dữ liệu cho từng file.")

    # Gói .zip cả bộ khi có từ 2 file — bộ 30-100 file bấm từng nút là không ai
    # làm nổi. Gói hỏng thì vẫn còn từng nút lẻ, không chặn kết quả.
    zip_evidence = None
    if len(made) >= 2:
        try:
            ten_goi = (bo["ten"] if bo is not None else "bo-file") + " - da dien"
            zip_bytes = bundle_zip([(p.name, p) for _t, _tok, p, _m, _w, _d in made])
            zip_token, zip_path = template_fill.save_filled_bytes(
                zip_bytes, ten_goi, extension="zip")
            zip_evidence = {"kind": "system",
                            "title": f"Tải cả bộ ({len(made)} file)",
                            "source_locator": f"template_fill#{zip_token}",
                            "quote": zip_path.name, "as_of": None}
        except Exception:  # noqa: BLE001
            zip_evidence = None

    # ---- Câu trả lời ------------------------------------------------------
    dau = (f"Đã điền bộ mẫu **«{bo['ten']}»**: **{len(made)}/{total} file**:"
           if bo is not None else
           f"Đã tạo **{len(made)}/{total} file** theo yêu cầu:")
    lines = [dau]
    for ten_file, token, _path, mo_ta, _warn, details in results:
        if token:
            lines.append(f"- **{ten_file}** — {mo_ta}")
            for d in details:
                lines.append(f"  - {d}")
        else:
            lines.append(f"- ~~{ten_file}~~ — {mo_ta}")

    review_lines = []
    if template_warning:
        review_lines.append(f"- {template_warning}")
    if attachments_truncated:
        review_lines.append(
            "- Hồ sơ đính kèm dài, hệ thống chỉ đọc được phần đầu của: "
            + ", ".join(f"«{f}»" for f in attachments_truncated)
            + " — dữ liệu ở các trang sau có thể chưa vào file.")
    if ghi_chu_plan:
        review_lines.append(f"- {ghi_chu_plan}")
    for ten_file, token, _path, _mo_ta, warn, _details in results:
        for w in warn:
            review_lines.append(f"- [{ten_file}] {w}")
    if review_lines:
        lines.append("\n**CHỖ AI TỰ QUYẾT ĐỊNH / CÒN THIẾU — KIỂM TRA BẮT BUỘC:**")
        lines.extend(review_lines)

    lines.append("\nBấm các nút **Tải** ngay dưới tin nhắn này để lấy từng "
                 "file" + (" (hoặc tải cả bộ .zip)" if zip_evidence else "")
                 + ". Văn bản sẽ gửi ra ngoài — vui lòng RÀ TOÀN VĂN từng "
                 "file (tên bên, con số, ngày tháng, điều khoản) trước khi dùng; "
                 "file tự xoá sau 24 giờ.")

    evidence = [{
        "kind": "system",
        "title": f"File đã tạo: {ten_file}",
        "source_locator": f"template_fill#{token}",
        "quote": out_path.name,
        "as_of": None,
    } for ten_file, token, out_path, _m, _w, _d in results if token]
    if zip_evidence:
        evidence.insert(0, zip_evidence)

    with db.session(role="internal", admin=True) as conn:
        db.audit(conn, user_id, "doc_factory", "conversation", conversation_id,
                 {"planned": total, "made": len(made),
                  "template_doc_id": template_doc_id,
                  "bo_mau_id": bo["id"] if bo is not None else None})
    return {"answer": "\n".join(lines), "answer_mode": "structured",
            "grounding_status": "verified", "evidence": evidence,
            "state": {}}
