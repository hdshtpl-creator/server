"""
bo_mau.py — BỘ MẪU HỒ SƠ: nhóm file .docx mẫu đi cùng nhau (hợp đồng + phụ
lục + biên bản bàn giao + giấy đề nghị…), tải lên từ Quản trị, để AI điền dữ
liệu khách vào TỪNG file của bộ trong một lượt chat (yêu cầu chủ dự án
15/09/2026: "bộ hồ sơ hợp đồng mẫu bằng Word + một file tổng hợp thông tin →
AI tự điền vào bộ hồ sơ").

Khác kệ HỢP ĐỒNG MẪU / THƯ MẪU trong kho tri thức: mẫu ở đây KHÔNG học vào
vector, không duyệt nhãn — chỉ giữ bản gốc .docx để điền. Tệp nằm ở
data/work/bo_mau/<id>/ (trong DATA_WORK, cùng rào với khuôn .docx tải lên
chat mà doc_factory._load_skeleton đang kiểm).

Phạm vi: department_id NULL = cả công ty; có giá trị = phòng đó + Ban quản
trị + admin. Tạo/sửa/xoá: người có quyền duyệt (require_reviewer ở api.py);
xem/dùng: mọi vai nội bộ trong phạm vi.
"""
import json
import os
import re
import shutil
import unicodedata
import uuid
from pathlib import Path

from app import db

DATA_WORK = Path(os.getenv("DATA_WORK", "./data/work"))
THU_MUC = "bo_mau"
# Trần số bộ đang hoạt động (chủ dự án: "tới 100 bộ") và số file mỗi bộ.
MAX_BO = 100
MAX_FILE_MOI_BO = 100
MAX_TEN_CHARS = 120
MAX_MO_TA_CHARS = 1000
CHI_NHAN = {".docx"}


class LoiBoMau(ValueError):
    """Lỗi nghiệp vụ nói thẳng cho người dùng (400)."""


def _fold(s: str) -> str:
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "d")
    return re.sub(r"[\s_]+", " ", s).strip().lower()


def goc_bo_mau() -> Path:
    root = DATA_WORK if DATA_WORK.is_absolute() else Path.cwd() / DATA_WORK
    return root / THU_MUC


def thu_muc_bo(bo_id: int) -> Path:
    return goc_bo_mau() / str(int(bo_id))


# ---------------------------------------------------------------------------
# Phạm vi nhìn thấy
# ---------------------------------------------------------------------------
def thay_het(role: str | None, is_banqt: bool) -> bool:
    return role == "admin" or bool(is_banqt)


def _visible_clause(role, dept_ids, is_banqt):
    """(mệnh đề WHERE, tham số) lọc bộ mẫu theo phạm vi người hỏi."""
    if thay_het(role, is_banqt):
        return "b.active", ()
    ids = [int(d) for d in (dept_ids or [])]
    if ids:
        return ("b.active AND (b.department_id IS NULL OR b.department_id = ANY(%s))",
                (ids,))
    return "b.active AND b.department_id IS NULL", ()


def _rows_to_sets(rows):
    return [{"id": r[0], "ten": r[1], "mo_ta": r[2] or "", "department_id": r[3],
             "created_by": r[4], "active": bool(r[5]),
             "created_at": str(r[6])[:19] if r[6] else None,
             "files": []} for r in rows]


def _attach_files(conn, sets):
    if not sets:
        return sets
    by_id = {s["id"]: s for s in sets}
    with conn.cursor() as cur:
        cur.execute("""SELECT id, bo_mau_id, ten_file, duong_dan, thu_tu,
                              placeholders, so_ky_tu
                         FROM bo_mau_file
                        WHERE bo_mau_id = ANY(%s)
                        ORDER BY bo_mau_id, thu_tu, id""", (list(by_id),))
        for fid, bid, ten, path, thu_tu, ph, so_ky_tu in cur.fetchall():
            ph_list = ph if isinstance(ph, list) else []
            by_id[bid]["files"].append({
                "id": fid, "ten_file": ten, "duong_dan": path, "thu_tu": thu_tu,
                "placeholders": ph_list, "so_placeholder": len(ph_list),
                "so_ky_tu": so_ky_tu or 0})
    for s in sets:
        s["so_file"] = len(s["files"])
    return sets


def list_sets(role=None, dept_ids=None, is_banqt=False, include_inactive=False):
    """Các bộ mẫu người hỏi được thấy, kèm danh sách file của từng bộ."""
    where, params = _visible_clause(role, dept_ids, is_banqt)
    if include_inactive and thay_het(role, is_banqt):
        where, params = "true", ()
    with db.session(role="internal") as conn:
        with conn.cursor() as cur:
            cur.execute(f"""SELECT b.id, b.ten, b.mo_ta, b.department_id, b.created_by,
                                   b.active, b.created_at
                              FROM bo_mau b
                             WHERE {where}
                             ORDER BY b.ten, b.id""", params)
            sets = _rows_to_sets(cur.fetchall())
        return _attach_files(conn, sets)


def get_set(bo_id: int, role=None, dept_ids=None, is_banqt=False):
    """Một bộ mẫu TRONG PHẠM VI người hỏi (None nếu không thấy)."""
    where, params = _visible_clause(role, dept_ids, is_banqt)
    with db.session(role="internal") as conn:
        with conn.cursor() as cur:
            cur.execute(f"""SELECT b.id, b.ten, b.mo_ta, b.department_id, b.created_by,
                                   b.active, b.created_at
                              FROM bo_mau b
                             WHERE b.id=%s AND {where}""", (int(bo_id), *params))
            rows = cur.fetchall()
        sets = _attach_files(conn, _rows_to_sets(rows))
    return sets[0] if sets else None


# ---------------------------------------------------------------------------
# Tạo / sửa / xoá bộ
# ---------------------------------------------------------------------------
def _clean_ten(ten: str) -> str:
    ten = " ".join((ten or "").split())[:MAX_TEN_CHARS].strip()
    if len(ten) < 2:
        raise LoiBoMau("Tên bộ mẫu cần ít nhất 2 ký tự")
    return ten


def create_set(ten: str, mo_ta: str | None, department_id: int | None, user_id):
    ten = _clean_ten(ten)
    mo_ta = " ".join((mo_ta or "").split())[:MAX_MO_TA_CHARS]
    with db.session(role="internal") as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT count(*) FROM bo_mau WHERE active")
            if int(cur.fetchone()[0]) >= MAX_BO:
                raise LoiBoMau(f"Đã đủ {MAX_BO} bộ mẫu đang hoạt động — gỡ bớt bộ "
                               "cũ trước khi thêm")
            cur.execute("SELECT 1 FROM bo_mau WHERE active AND lower(ten)=lower(%s)",
                        (ten,))
            if cur.fetchone():
                raise LoiBoMau(f"Đã có bộ mẫu tên «{ten}» — đặt tên khác để gọi "
                               "trong chat không nhầm")
            cur.execute("""INSERT INTO bo_mau (ten, mo_ta, department_id, created_by)
                           VALUES (%s,%s,%s,%s) RETURNING id""",
                        (ten, mo_ta or None, department_id, user_id))
            bo_id = cur.fetchone()[0]
        db.audit(conn, user_id, "bo_mau_create", "bo_mau", bo_id,
                 {"ten": ten, "department_id": department_id})
    thu_muc_bo(bo_id).mkdir(parents=True, exist_ok=True)
    return bo_id


_KHONG_DOI = object()


def update_set(bo_id: int, user_id, ten=None, mo_ta=None,
               department_id=_KHONG_DOI, active=None):
    fields, params = [], []
    ten_moi = None
    if ten is not None:
        ten_moi = _clean_ten(ten)
        fields.append("ten=%s")
        params.append(ten_moi)
    if mo_ta is not None:
        fields.append("mo_ta=%s")
        params.append(" ".join(mo_ta.split())[:MAX_MO_TA_CHARS] or None)
    if department_id is not _KHONG_DOI:
        fields.append("department_id=%s")
        params.append(department_id)
    if active is not None:
        fields.append("active=%s")
        params.append(bool(active))
    if not fields:
        return False
    fields.append("updated_at=now()")
    with db.session(role="internal") as conn:
        with conn.cursor() as cur:
            if ten_moi is not None:
                cur.execute("""SELECT 1 FROM bo_mau
                                WHERE active AND lower(ten)=lower(%s) AND id<>%s""",
                            (ten_moi, int(bo_id)))
                if cur.fetchone():
                    raise LoiBoMau(f"Đã có bộ mẫu khác tên «{ten_moi}»")
            if active:
                cur.execute("SELECT count(*) FROM bo_mau WHERE active AND id<>%s",
                            (int(bo_id),))
                if int(cur.fetchone()[0]) >= MAX_BO:
                    raise LoiBoMau(f"Đã đủ {MAX_BO} bộ mẫu đang hoạt động")
            cur.execute("UPDATE bo_mau SET " + ", ".join(fields) + " WHERE id=%s",
                        (*params, int(bo_id)))
            changed = cur.rowcount
        db.audit(conn, user_id, "bo_mau_update", "bo_mau", int(bo_id),
                 {"fields": [f.split("=")[0] for f in fields]})
    return bool(changed)


def delete_set(bo_id: int, user_id):
    """Xoá bộ + toàn bộ tệp gốc trên đĩa. CASCADE chỉ xoá dòng, phải dọn tay."""
    with db.session(role="internal") as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM bo_mau WHERE id=%s", (int(bo_id),))
            changed = cur.rowcount
        db.audit(conn, user_id, "bo_mau_delete", "bo_mau", int(bo_id), {})
    folder = thu_muc_bo(bo_id)
    if folder.exists():
        shutil.rmtree(folder, ignore_errors=True)
    return bool(changed)


# ---------------------------------------------------------------------------
# File trong bộ
# ---------------------------------------------------------------------------
def safe_name(name: str) -> str:
    name = Path(name or "").name
    name = re.sub(r"[^\w\s.\-()À-ỹ]", "_", name, flags=re.UNICODE).strip()
    return name[:150] or "mau.docx"


def scan_docx(path: Path):
    """(danh sách placeholder nguyên văn, số ký tự) — quét bằng chính máy của
    template_fill để lúc điền và lúc kê khai không lệch nhau."""
    import docx
    from app import template_fill
    doc = docx.Document(str(path))
    ph = [literal for literal, _key in template_fill.scan_placeholders(doc)]
    text = template_fill.document_text(doc)
    return ph, len(text)


def add_file(bo_id: int, filename: str, tmp_path: Path, user_id) -> dict:
    """Đưa một file .docx vào bộ: chép vào thư mục bộ, quét {{…}}, ghi bản ghi.
    tmp_path là file tạm đã nhận xong từ upload; người gọi tự dọn."""
    safe = safe_name(filename)
    if Path(safe).suffix.lower() not in CHI_NHAN:
        raise LoiBoMau(f"«{safe}»: chỉ nhận file .docx (Word mới). File .doc cũ "
                       "hãy mở bằng Word và Lưu thành .docx trước.")
    with db.session(role="internal") as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM bo_mau WHERE id=%s", (int(bo_id),))
            if not cur.fetchone():
                raise LoiBoMau("Bộ mẫu không tồn tại")
            cur.execute("""SELECT count(*), coalesce(max(thu_tu),0)
                             FROM bo_mau_file WHERE bo_mau_id=%s""", (int(bo_id),))
            n, max_thu_tu = cur.fetchone()
            if int(n) >= MAX_FILE_MOI_BO:
                raise LoiBoMau(f"Bộ đã đủ {MAX_FILE_MOI_BO} file")
            cur.execute("SELECT 1 FROM bo_mau_file WHERE bo_mau_id=%s AND ten_file=%s",
                        (int(bo_id), safe))
            if cur.fetchone():
                raise LoiBoMau(f"Bộ đã có file tên «{safe}» — gỡ bản cũ trước nếu "
                               "muốn thay")
    folder = thu_muc_bo(bo_id)
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / f"{uuid.uuid4().hex[:8]}_{safe}"
    shutil.copyfile(tmp_path, dest)
    try:
        ph, so_ky_tu = scan_docx(dest)
    except Exception as exc:  # noqa: BLE001 — file hỏng thì không nhận
        dest.unlink(missing_ok=True)
        raise LoiBoMau(f"«{safe}»: không mở được bằng bộ đọc Word "
                       f"({type(exc).__name__}) — file có thể hỏng hoặc không phải "
                       ".docx thật") from exc
    with db.session(role="internal") as conn:
        with conn.cursor() as cur:
            cur.execute("""INSERT INTO bo_mau_file
                               (bo_mau_id, ten_file, duong_dan, thu_tu, placeholders,
                                so_ky_tu)
                           VALUES (%s,%s,%s,%s,%s,%s) RETURNING id""",
                        (int(bo_id), safe, str(dest), int(max_thu_tu) + 1,
                         json.dumps(ph, ensure_ascii=False), so_ky_tu))
            fid = cur.fetchone()[0]
            cur.execute("UPDATE bo_mau SET updated_at=now() WHERE id=%s", (int(bo_id),))
        db.audit(conn, user_id, "bo_mau_add_file", "bo_mau_file", fid,
                 {"bo_mau_id": int(bo_id), "ten_file": safe, "placeholders": len(ph)})
    return {"id": fid, "ten_file": safe, "placeholders": ph,
            "so_placeholder": len(ph), "so_ky_tu": so_ky_tu}


def delete_file(bo_id: int, file_id: int, user_id) -> bool:
    with db.session(role="internal") as conn:
        with conn.cursor() as cur:
            cur.execute("""DELETE FROM bo_mau_file WHERE id=%s AND bo_mau_id=%s
                           RETURNING duong_dan""", (int(file_id), int(bo_id)))
            row = cur.fetchone()
            if row:
                cur.execute("UPDATE bo_mau SET updated_at=now() WHERE id=%s",
                            (int(bo_id),))
        db.audit(conn, user_id, "bo_mau_delete_file", "bo_mau_file", int(file_id),
                 {"bo_mau_id": int(bo_id)})
    if not row:
        return False
    try:
        resolve_path(row[0]).unlink(missing_ok=True)
    except (ValueError, OSError):
        pass
    return True


def resolve_path(path_str: str) -> Path:
    """Mở tệp mẫu TRONG RÀO data/work/bo_mau — đường dẫn đọc từ CSDL không được
    dắt ra ngoài (cùng nguyên tắc với doc_factory._load_skeleton)."""
    path = Path(path_str)
    if not path.is_absolute():
        path = Path.cwd() / path
    resolved = path.resolve(strict=True)
    resolved.relative_to(goc_bo_mau().resolve())
    return resolved


# ---------------------------------------------------------------------------
# Gọi tên bộ mẫu trong câu chat
# ---------------------------------------------------------------------------
def match_set(question: str, sets) -> dict | None:
    """Bộ mẫu được GỌI TÊN trong câu ("tạo bộ hồ sơ theo bộ mẫu Thuê nhà cho
    khách Minh"): tên bộ (bỏ dấu) là chuỗi con của câu (bỏ dấu). Nhiều bộ khớp
    thì lấy tên DÀI nhất — "HĐ thuê nhà" thắng "HĐ" (thuần, test được)."""
    q = _fold(question)
    if not q:
        return None
    best, best_len = None, 0
    for s in sets or []:
        name = _fold(s.get("ten") or "")
        if len(name) >= 2 and name in q and len(name) > best_len:
            best, best_len = s, len(name)
    return best


def nhac_bo_mau(question: str) -> bool:
    """Câu có nhắc tới 'bộ mẫu' (để hỏi lại khi không khớp tên nào)."""
    return "bo mau" in _fold(question)
