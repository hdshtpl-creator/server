"""xem_truoc.py — Bản XEM THẲNG TRONG TRÌNH DUYỆT của một tệp gốc.

Dùng chung cho nút "Xem trước" của web HDS (api._preview_pdf) và cho CRM
(/integration/v1/…/preview, link xem tạm — 01/10/2026: "CRM mở file khách bất
kỳ đều xem nhanh được, không cần tải").

  · PDF, ảnh PNG/JPG/WEBP, văn bản thuần: trả nguyên tệp, trình duyệt tự hiện.
  · Word / Excel / CSV: đổi sang PDF MỘT LẦN bằng LibreOffice (có sẵn trên máy
    chủ cho khâu đọc .doc) rồi giữ ở data/work/preview; lần sau trả ngay.
  · Loại khác (TIFF, BMP, ZIP…): không xem thẳng được — báo để tải về.

CHỐT AN TOÀN: chỉ trả inline đúng các loại trên, văn bản luôn là text/plain.
Tệp khách có thể là HTML đội lốt .doc (Confluence, trang Word web) — không bao
giờ được phát ra dưới dạng text/html trên tên miền của hệ thống.
"""
from __future__ import annotations

import hashlib
import os
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from app import chay_soffice

HIEN_THANG = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".txt": "text/plain; charset=utf-8",
    ".md": "text/plain; charset=utf-8",
}
DOI_SANG_PDF = {".docx", ".doc", ".xlsx", ".csv"}
TIEN_TO_CACHE_TICH_HOP = "th-"
GIU_CACHE_NGAY = 30


class LoiXemTruoc(Exception):
    def __init__(self, msg: str, status: int = 409):
        super().__init__(msg)
        self.status = status


def thu_muc_cache() -> Path:
    p = Path(os.getenv("DATA_WORK", "./data/work")) / "preview"
    p.mkdir(parents=True, exist_ok=True)
    return p


def khoa_cache(path: Path) -> str:
    """Khoá cache theo (đường dẫn, kích thước, mtime) — không phải băm cả tệp
    lớn mỗi lần xem; tệp đổi nội dung là mtime/kích thước đổi → sinh lại."""
    st = path.stat()
    h = hashlib.sha1(f"{path.resolve()}|{st.st_size}|{st.st_mtime_ns}".encode()).hexdigest()
    return TIEN_TO_CACHE_TICH_HOP + h[:32]


def _cache_pdf(resolved: Path, cache_key) -> tuple[Path, bool]:
    """(đường dẫn PDF trong cache, còn mới không) — mới = sinh sau tệp gốc."""
    out = thu_muc_cache() / (re.sub(r"[^A-Za-z0-9_-]", "_", str(cache_key)) + ".pdf")
    try:
        return out, out.exists() and out.stat().st_mtime >= resolved.stat().st_mtime
    except OSError:
        return out, False


def _soffice_sang_pdf(src: Path, tmp: str, out: Path) -> None:
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice:
        raise LoiXemTruoc("Máy chủ chưa có LibreOffice để tạo bản xem trước — hãy tải về.")
    # Hồ sơ LibreOffice riêng cho mỗi lượt — hai lượt chuyển song song
    # không giẫm profile của nhau.
    cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp,
           f"-env:UserInstallation=file://{tmp}/profile", str(src)]
    try:
        chay_soffice.chay(cmd, timeout=120)
    except (subprocess.SubprocessError, OSError):
        raise LoiXemTruoc("Chưa chuyển được tệp này sang PDF để xem trước — hãy tải về.")
    produced = Path(tmp) / (src.stem + ".pdf")
    if not produced.exists():
        raise LoiXemTruoc("Chưa chuyển được tệp này sang PDF để xem trước — hãy tải về.")
    shutil.move(str(produced), str(out))


def pdf_tu_office(resolved: Path, cache_key) -> Path:
    """Bản PDF của một tệp Office, sinh một lần rồi dùng lại (theo mtime)."""
    out, moi = _cache_pdf(resolved, cache_key)
    if moi:
        return out
    with tempfile.TemporaryDirectory() as tmp:
        _soffice_sang_pdf(resolved, tmp, out)
    return out


# ---- Văn bản chỉ có BẢN CHỮ .md ----
# Kệ luật nạp lô 17–19/09/2026 có ~24.700 văn bản chỉ ở dạng .md (chữ đã bóc
# sẵn, KHÔNG kèm PDF/Word gốc). Phát text/plain thì người xem thấy nguyên
# "## Điều 1." và thẻ </p> sót lại — trông như tệp hỏng. Dàn lại thành trang
# văn bản (Word → PDF qua LibreOffice) cho nút Xem, và thành .docx cho nút Tải.
# Chỉ dàn trang, KHÔNG sửa chữ: số khoản "1." và gạch đầu dòng "- " giữ nguyên
# (đánh số tự động của Word sẽ đánh lại số khoản luật).
_TIEU_DE_MD = re.compile(r"^#{1,6}\s+(.*?)\s*#*\s*$")
_CA_DONG_DAM = re.compile(r"^\*\*\s*((?:(?!\*\*).)+?)\s*\*\*$")
# Ký tự điều khiển (chữ OCR hay có) — python-docx từ chối cả tệp nếu gặp.
_KY_TU_CAM_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")
# Chỉ thẻ HTML sót của bộ bóc; "<tên cơ quan>" trong biểu mẫu là chữ thật.
_THE_HTML_SOT = re.compile(
    r"</?(?:p|td|tr|th|table|tbody|thead|div|span|br|font|b|i|u|strong|em|sup|sub)\b[^>]*>",
    re.I)
_CAN_GIUA = re.compile(r"^(chương|mục|phần|phụ lục)\b", re.I)
_DIEU = re.compile(r"^điều\s+\d", re.I)
PHIEN_BAN_DAN_TRANG_MD = 1  # tăng khi đổi cách dàn → bản PDF cache cũ bị bỏ qua


def docx_tu_markdown(text: str, tieu_de: str = "") -> bytes:
    """Dàn bản chữ .md thành .docx thể thức văn bản hành chính (Times 13pt).

    Dòng trống ngăn đoạn; các dòng liền nhau trong một đoạn giữ xuống dòng như
    tệp gốc. Tiêu đề "##": in đậm; căn giữa phần đầu văn bản (trước Điều đầu
    tiên), dòng viết hoa và Chương/Mục/Phần/Phụ lục; "Điều N." căn trái.
    """
    import io

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin, s.right_margin = Cm(3), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal.paragraph_format.space_after = Pt(6)
    if tieu_de:
        doc.core_properties.title = tieu_de

    da_toi_dieu = False
    doan: list[str] = []

    def xa_doan():
        if doan:
            p = doc.add_paragraph()
            for i, dong in enumerate(doan):
                run = p.add_run(dong)
                if i < len(doan) - 1:
                    run.add_break()
            doan.clear()

    text = _KY_TU_CAM_XML.sub("", (text or "").replace("\r\n", "\n"))
    for raw in text.split("\n"):
        dong = _THE_HTML_SOT.sub("", raw).strip()
        if not dong:
            xa_doan()
            continue
        m = _TIEU_DE_MD.match(dong) or _CA_DONG_DAM.match(dong)
        if not m:
            doan.append(dong)
            continue
        xa_doan()
        noi_dung = m.group(1)
        if not noi_dung:
            continue
        la_dieu = bool(_DIEU.match(noi_dung))
        da_toi_dieu = da_toi_dieu or la_dieu
        co_chu = any(c.isalpha() for c in noi_dung)
        can_giua = not la_dieu and (
            not da_toi_dieu or _CAN_GIUA.match(noi_dung)
            or (co_chu and noi_dung == noi_dung.upper()))
        p = doc.add_paragraph()
        p.add_run(noi_dung).bold = True
        if can_giua:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    xa_doan()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def doc_markdown(path: Path) -> str:
    return path.read_bytes().decode("utf-8", errors="replace")


def pdf_tu_markdown(resolved: Path, cache_key) -> Path:
    """Bản PDF dàn trang của một tệp .md, sinh một lần rồi dùng lại."""
    out, moi = _cache_pdf(resolved, f"md{PHIEN_BAN_DAN_TRANG_MD}-{cache_key}")
    if moi:
        return out
    with tempfile.TemporaryDirectory() as tmp:
        # Tên tệp trung gian ASCII — LibreOffice đặt tên PDF theo tên nguồn.
        src = Path(tmp) / "van-ban.docx"
        src.write_bytes(docx_tu_markdown(doc_markdown(resolved), resolved.stem))
        _soffice_sang_pdf(src, tmp, out)
    return out


def xem_duoc(path: Path) -> bool:
    s = path.suffix.lower()
    return s in HIEN_THANG or s in DOI_SANG_PDF


def ban_xem(path: Path) -> tuple[Path, str, str]:
    """(tệp để phát, media type, tên hiển thị) cho chế độ xem thẳng."""
    s = path.suffix.lower()
    if s in HIEN_THANG:
        return path, HIEN_THANG[s], path.name
    if s in DOI_SANG_PDF:
        pdf = pdf_tu_office(path, khoa_cache(path))
        return pdf, "application/pdf", path.stem + ".pdf"
    raise LoiXemTruoc(f"Định dạng {s or '(không đuôi)'} chưa xem thẳng được trong trình duyệt — hãy tải về.", 415)


def don_cache(bay_gio: float | None = None, ngay: int = GIU_CACHE_NGAY) -> int:
    """Xoá bản xem trước của CRM (tiền tố th-) cũ hơn `ngay` ngày — xem lại thì
    sinh lại. Không đụng bản xem trước của web HDS (khoá theo id tài liệu)."""
    bay_gio = bay_gio or time.time()
    n = 0
    for p in thu_muc_cache().glob(TIEN_TO_CACHE_TICH_HOP + "*.pdf"):
        try:
            if bay_gio - p.stat().st_mtime > ngay * 86400:
                p.unlink()
                n += 1
        except OSError:
            continue
    return n
