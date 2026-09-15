"""
Test RÀ SOÁT RỦI RO TÀI LIỆU (app/ra_soat_rui_ro.py) — phần THUẦN, không cần
PostgreSQL/Ollama.

Chay: python -m unittest tests.test_ra_soat_rui_ro -v

Bốn thứ phải đứng vững:
  1. Nhận diện loại hợp đồng từ tiêu đề / phần đầu (có dấu, không dấu, hoa
     thường); văn bản trung tính → "khac".
  2. Tách điều khoản theo "Điều n." / "ĐIỀU n -" / "n." đầu dòng; văn bản trơn
     → một mục.
  3. ra_soat trên hợp đồng thật: ngưỡng luật (thử việc, lương thử việc, lãi
     suất) bắt đúng; điều khoản bắt buộc thiếu → "thieu"; đủ mọi mục → "thap";
     thứ tự thieu → canh_bao → dat.
  4. Báo cáo .docx mở được, có bảng, có tiêu đề.
"""
import io
import unittest
from datetime import date

from app import ra_soat_rui_ro as rs


HDLD = """CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM
Độc lập - Tự do - Hạnh phúc
HỢP ĐỒNG LAO ĐỘNG
Số: 12/2026/HĐLĐ
Hôm nay, ngày 01 tháng 03 năm 2026, tại Hà Nội, chúng tôi gồm:
BÊN A – NGƯỜI SỬ DỤNG LAO ĐỘNG: Công ty TNHH ABC; Địa chỉ: 12 Lê Lợi, Hà Nội; Mã số doanh nghiệp: 0101234567
Đại diện: Ông Nguyễn Văn A – Chức vụ: Giám đốc
BÊN B – NGƯỜI LAO ĐỘNG: Bà Trần Thị B; Sinh ngày 02/02/1995; Giới tính: Nữ; Nơi cư trú: 5 Trần Phú, Hà Nội; CCCD số 001195001234
Điều 1. Công việc và địa điểm làm việc
Chức danh: Nhân viên kế toán. Địa điểm làm việc: trụ sở Công ty.
Điều 2. Thời hạn hợp đồng
Loại hợp đồng: xác định thời hạn 24 tháng, từ 01/03/2026 đến 28/02/2028. Thời gian thử việc là 90 ngày; tiền lương thử việc bằng 80% mức lương chính thức.
Điều 3. Tiền lương
Mức lương: 15.000.000 đồng/tháng; phụ cấp ăn trưa 730.000 đồng. Hình thức trả lương: chuyển khoản; thời hạn trả lương: ngày 05 hàng tháng. Chế độ nâng bậc, nâng lương theo quy chế Công ty.
Điều 4. Thời giờ làm việc, nghỉ ngơi
8 giờ/ngày, 48 giờ/tuần; nghỉ hàng tuần chủ nhật; nghỉ phép năm theo luật. Làm thêm giờ không quá 40 giờ/tháng và 200 giờ/năm.
Điều 5. Bảo hộ lao động
Được trang bị bảo hộ lao động và bảo đảm an toàn, vệ sinh lao động.
Điều 6. Bảo hiểm và đào tạo
Đóng bảo hiểm xã hội, bảo hiểm y tế, bảo hiểm thất nghiệp theo quy định. Được đào tạo, bồi dưỡng nâng cao trình độ kỹ năng nghề.
Điều 7. Điều khoản chung
Việc chấm dứt hợp đồng thực hiện theo Bộ luật Lao động. Tranh chấp được giải quyết tại Toà án có thẩm quyền. Hợp đồng có hiệu lực từ ngày ký.
ĐẠI DIỆN BÊN A (ký tên)                NGƯỜI LAO ĐỘNG (ký, ghi rõ họ tên)
"""

HD_VAY = """HỢP ĐỒNG VAY TIỀN
Hôm nay, ngày 10/01/2026, tại Hà Nội.
Bên cho vay (Bên A): Ông Lê Văn C, CCCD 001080001111, địa chỉ 1 Hàng Bài, Hà Nội.
Bên vay (Bên B): Bà Phạm Thị D, CCCD 001090002222, địa chỉ 2 Hàng Gai, Hà Nội.
Điều 1. Số tiền vay
Bên A cho Bên B vay số tiền 500.000.000 đồng (năm trăm triệu đồng).
Điều 2. Lãi suất và thời hạn vay
Lãi suất: 30%/năm. Thời hạn vay: 12 tháng kể từ ngày giải ngân.
Điều 3. Phương thức trả nợ
Trả gốc và lãi vào ngày cuối kỳ; Bên B được trả nợ trước hạn.
Điều 4. Giải quyết tranh chấp
Tranh chấp giải quyết tại Toà án nơi Bên A cư trú. Hợp đồng có hiệu lực từ ngày ký.
Đại diện hai bên ký tên.
"""

HD_DV_THIEU = """HỢP ĐỒNG DỊCH VỤ TƯ VẤN
Bên A (Bên sử dụng dịch vụ): Công ty X. Bên B (Bên cung cấp dịch vụ): Công ty Luật Y.
Điều 1. Đối tượng hợp đồng
Bên B cung cấp dịch vụ tư vấn pháp lý thường xuyên cho Bên A.
Điều 2. Phí dịch vụ và thanh toán
Phí dịch vụ 20.000.000 đồng/tháng, thanh toán trước ngày 05 hàng tháng.
Điều 3. Thời hạn
Thời hạn hợp đồng 12 tháng kể từ ngày ký.
"""

HD_DV_DU = """HỢP ĐỒNG DỊCH VỤ TƯ VẤN PHÁP LÝ
Số: 05/2026/HĐDV
Hôm nay, ngày 05 tháng 01 năm 2026, tại Hà Nội.
Bên A (Bên sử dụng dịch vụ): Công ty Cổ phần X; Địa chỉ: 1 Tràng Tiền; Đại diện: Ông A – Chức vụ: Tổng giám đốc.
Bên B (Bên cung cấp dịch vụ): Công ty Luật TNHH Y; Địa chỉ: 2 Lý Thường Kiệt; Đại diện: Bà B – Chức vụ: Giám đốc.
Điều 1. Đối tượng và phạm vi dịch vụ
Bên B cung cấp dịch vụ tư vấn pháp lý thường xuyên: rà soát hợp đồng, tư vấn lao động, doanh nghiệp.
Điều 2. Phí dịch vụ và thanh toán
Phí dịch vụ 20.000.000 đồng/tháng (chưa gồm VAT); thanh toán chuyển khoản trước ngày 05 hàng tháng. Lãi chậm thanh toán 10%/năm trên số tiền chậm trả.
Điều 3. Thời hạn thực hiện
Thời hạn hợp đồng 12 tháng kể từ ngày ký; gia hạn bằng văn bản.
Điều 4. Quyền và nghĩa vụ của các bên
Bên A cung cấp tài liệu đầy đủ; Bên B thực hiện đúng tiến độ, bảo mật thông tin của Bên A.
Điều 5. Nghiệm thu
Kết quả dịch vụ được nghiệm thu bằng biên bản hoặc email xác nhận trong 3 ngày làm việc.
Điều 6. Phạt vi phạm, bồi thường và bất khả kháng
Bên vi phạm chịu phạt vi phạm 8% giá trị phần nghĩa vụ bị vi phạm và bồi thường thiệt hại thực tế. Sự kiện bất khả kháng được miễn trách theo quy định.
Điều 7. Chấm dứt hợp đồng
Mỗi bên được đơn phương chấm dứt khi báo trước 30 ngày.
Điều 8. Giải quyết tranh chấp và hiệu lực
Tranh chấp được thương lượng; không thành thì đưa ra Toà án có thẩm quyền. Hợp đồng có hiệu lực từ ngày ký.
ĐẠI DIỆN BÊN A (ký, ghi rõ họ tên)      ĐẠI DIỆN BÊN B (ký, ghi rõ họ tên)
"""


def _theo_ma(kq, ma):
    return next((m for m in kq["muc"] if m["ma"] == ma), None)


class BoDauTests(unittest.TestCase):
    def test_bo_dau_giu_do_dai(self):
        s = "Hợp ĐỒNG lao động — Điều 5: đặt cọc"
        f = rs.bo_dau(s)
        self.assertEqual(f, "hop dong lao dong — dieu 5: dat coc")
        self.assertEqual(len(f), len(s))


class NhanDienLoaiTests(unittest.TestCase):
    def test_sau_loai_khac_nhau(self):
        mau = [
            ("HỢP ĐỒNG LAO ĐỘNG", "Người sử dụng lao động và người lao động thoả thuận…",
             "hop_dong_lao_dong"),
            ("Hop dong thue nha", "Bên cho thuê giao nhà, Bên thuê trả tiền thuê hàng tháng.",
             "hop_dong_thue"),
            ("HỢP ĐỒNG VAY TIỀN", "Bên cho vay giao cho Bên vay khoản vay 200 triệu, lãi suất…",
             "hop_dong_vay"),
            ("", "THOẢ THUẬN BẢO MẬT THÔNG TIN\nBên tiết lộ và Bên nhận thông tin cam kết bảo mật.",
             "nda"),
            ("Hợp đồng nhượng quyền thương mại", "Bên nhượng quyền cấp cho Bên nhận quyền… "
             "phí nhượng quyền ban đầu", "hop_dong_nhuong_quyen"),
            ("HỢP ĐỒNG CHUYỂN NHƯỢNG CỔ PHẦN", "Bên chuyển nhượng bán cho Bên nhận chuyển "
             "nhượng 10.000 cổ phần.", "hop_dong_chuyen_nhuong_von"),
            ("Hợp đồng mua bán hàng hoá", "Bên bán giao hàng cho Bên mua tại kho; hàng hoá "
             "gồm…", "hop_dong_mua_ban"),
        ]
        for tieu_de, text, mong in mau:
            with self.subTest(tieu_de=tieu_de or text[:30]):
                ma, tin = rs.nhan_dien_loai(tieu_de, text)
                self.assertEqual(ma, mong)
                self.assertGreater(tin, 0.5)

    def test_van_ban_trung_tinh_ra_khac(self):
        ma, tin = rs.nhan_dien_loai("Thông báo nghỉ lễ",
                                    "Công ty thông báo lịch nghỉ lễ Quốc khánh 2/9: nghỉ từ "
                                    "ngày 01/09 đến hết 02/09. Đề nghị các phòng sắp xếp.")
        self.assertEqual((ma, tin), ("khac", 0.0))

    def test_tieu_de_nang_hon_than(self):
        _ma, tin_td = rs.nhan_dien_loai("HỢP ĐỒNG LAO ĐỘNG", "")
        _ma, tin_than = rs.nhan_dien_loai("", "x" * 500 + " hợp đồng lao động")
        self.assertGreater(tin_td, tin_than)


class TachDieuKhoanTests(unittest.TestCase):
    def test_ba_kieu_ghi_dieu(self):
        text = ("Điều 1. Đối tượng\nnội dung 1\nĐiều 2: Giá cả\nnội dung 2\n"
                "ĐIỀU 3 - Thanh toán\nnội dung 3\n")
        muc = rs.tach_dieu_khoan(text)
        self.assertEqual([m["so"] for m in muc], ["1", "2", "3"])
        self.assertEqual([m["tieu_de"] for m in muc], ["Đối tượng", "Giá cả", "Thanh toán"])
        self.assertIn("nội dung 1", muc[0]["noi_dung"])
        self.assertNotIn("nội dung 2", muc[0]["noi_dung"])
        self.assertTrue(muc[2]["noi_dung"].startswith("ĐIỀU 3"))

    def test_khong_dau_va_khong_bat_cau_trich_dan(self):
        text = "DIEU 1. Pham vi\nTheo Dieu 5 cua hop dong nay, cac ben…\nDieu 2) Gia\nabc"
        muc = rs.tach_dieu_khoan(text)
        self.assertEqual([m["so"] for m in muc], ["1", "2"])

    def test_danh_so_khong_co_dieu(self):
        muc = rs.tach_dieu_khoan("1. Phạm vi\nabc\n2. Giá\nxyz\n3) Thanh toán\nqqq")
        self.assertEqual([m["so"] for m in muc], ["1", "2", "3"])
        self.assertEqual(muc[1]["tieu_de"], "Giá")
        self.assertIn("xyz", muc[1]["noi_dung"])

    def test_van_ban_tron_mot_muc(self):
        text = "Hai bên thoả thuận như sau. Bên A giao hàng, Bên B trả tiền."
        muc = rs.tach_dieu_khoan(text)
        self.assertEqual(muc, [{"so": "", "tieu_de": "", "noi_dung": text, "bat_dau": 0,
                                "kieu": "toan_van"}])

    def test_phan_mo_dau_khong_tinh_la_dieu(self):
        self.assertEqual(len(rs.tach_dieu_khoan(HDLD)), 7)


class RaSoatHDLDTests(unittest.TestCase):
    def setUp(self):
        self.kq = rs.ra_soat(HDLD, "HĐLĐ Trần Thị B")

    def test_nhan_dien_va_dem(self):
        self.assertEqual(self.kq["loai"], "hop_dong_lao_dong")
        self.assertEqual(self.kq["so_dieu_khoan"], 7)
        self.assertGreaterEqual(self.kq["do_tin_cay"], 0.8)

    def test_nguong_thu_viec_va_luong_thu_viec(self):
        tv = _theo_ma(self.kq, "thu_viec_ngay")
        self.assertEqual(tv["trang_thai"], "canh_bao")
        self.assertEqual(tv["gia_tri"], 90)
        self.assertIn("Điều 25", tv["can_cu"])
        self.assertEqual(tv["dieu_khoan"], "Điều 2. Thời hạn hợp đồng")
        self.assertIn("90 ngày", tv["trich"])
        luong = _theo_ma(self.kq, "luong_thu_viec_pct")
        self.assertEqual(luong["trang_thai"], "canh_bao")
        self.assertEqual(luong["gia_tri"], 80)
        self.assertIn("Điều 26", luong["can_cu"])

    def test_nguong_trong_gioi_han_la_dat(self):
        for ma, gia_tri in [("hd_xac_dinh_thang", 24), ("gio_ngay", 8), ("gio_tuan", 48),
                            ("lam_them_thang", 40), ("lam_them_nam", 200)]:
            with self.subTest(ma=ma):
                m = _theo_ma(self.kq, ma)
                self.assertIsNotNone(m)
                self.assertEqual(m["trang_thai"], "dat")
                self.assertEqual(m["gia_tri"], gia_tri)

    def test_nguong_khong_neu_thi_khong_bao(self):
        self.assertIsNone(_theo_ma(self.kq, "thu_viec_thang"))

    def test_dieu_khoan_bat_buoc_deu_dat(self):
        for dk in rs.LOAI_HOP_DONG["hop_dong_lao_dong"]["dieu_khoan"]:
            with self.subTest(ma=dk["ma"]):
                m = _theo_ma(self.kq, dk["ma"])
                self.assertEqual(m["trang_thai"], "dat", m)
                self.assertTrue(m["dieu_khoan"])
                self.assertTrue(m["trich"])
                self.assertLessEqual(len(m["trich"]), rs.TRICH_MAX)
        self.assertEqual(_theo_ma(self.kq, "cong_viec")["dieu_khoan"],
                         "Điều 1. Công việc và địa điểm làm việc")
        self.assertEqual(_theo_ma(self.kq, "nsdld")["dieu_khoan"], "Phần mở đầu")

    def test_muc_chung_dat_va_khong_lap(self):
        for ma in ("cac_ben", "ngay_ky_hieu_luc", "tranh_chap", "chu_ky_dai_dien"):
            with self.subTest(ma=ma):
                self.assertEqual(_theo_ma(self.kq, ma)["trang_thai"], "dat")
        self.assertEqual(sum(1 for m in self.kq["muc"] if m["ma"] == "tranh_chap"), 1)

    def test_tong_ket_va_muc_rui_ro(self):
        tk = self.kq["tong_ket"]
        self.assertEqual(tk["thieu"], 0)
        self.assertEqual(tk["canh_bao"], 2)
        self.assertEqual(tk["muc_rui_ro"], "trung_binh")
        self.assertEqual(tk["dat"] + tk["canh_bao"] + tk["thieu"], len(self.kq["muc"]))

    def test_thieu_bao_hiem_len_dau_va_rui_ro_cao(self):
        text = HDLD.replace(
            "Đóng bảo hiểm xã hội, bảo hiểm y tế, bảo hiểm thất nghiệp theo quy định. ", "")
        kq = rs.ra_soat(text, loai="hop_dong_lao_dong")
        self.assertEqual(kq["muc"][0]["ma"], "bao_hiem")
        self.assertEqual(kq["muc"][0]["trang_thai"], "thieu")
        self.assertIn("Điều 21", kq["muc"][0]["can_cu"])
        self.assertEqual(kq["tong_ket"]["muc_rui_ro"], "cao")
        thu_tu = [rs._THU_TU[m["trang_thai"]] for m in kq["muc"]]
        self.assertEqual(thu_tu, sorted(thu_tu))
        self.assertEqual(kq["do_tin_cay"], 1.0)  # loai ép sẵn


class RaSoatLoaiKhacTests(unittest.TestCase):
    def test_vay_lai_30_phan_tram(self):
        kq = rs.ra_soat(HD_VAY, "HĐ vay")
        self.assertEqual(kq["loai"], "hop_dong_vay")
        lai = _theo_ma(kq, "lai_suat_nam")
        self.assertEqual(lai["trang_thai"], "canh_bao")
        self.assertEqual(lai["gia_tri"], 30)
        self.assertIn("Điều 468", lai["can_cu"])
        self.assertIn("20", lai["giai_thich"])
        self.assertEqual(lai["dieu_khoan"], "Điều 2. Lãi suất và thời hạn vay")
        self.assertEqual(_theo_ma(kq, "so_tien")["trang_thai"], "dat")
        self.assertEqual(_theo_ma(kq, "tra_truoc_han")["trang_thai"], "dat")
        self.assertEqual(_theo_ma(kq, "bao_dam")["trang_thai"], "canh_bao")

    def test_dich_vu_thieu_tranh_chap(self):
        kq = rs.ra_soat(HD_DV_THIEU, "HĐ dịch vụ")
        self.assertEqual(kq["loai"], "hop_dong_dich_vu")
        tc = [m for m in kq["muc"] if m["ma"] == "tranh_chap"]
        self.assertEqual(len(tc), 1)
        self.assertIn(tc[0]["trang_thai"], ("thieu", "canh_bao"))
        self.assertTrue(tc[0]["de_xuat"])
        self.assertEqual(_theo_ma(kq, "doi_tuong")["trang_thai"], "dat")
        self.assertEqual(kq["tong_ket"]["muc_rui_ro"], "cao")

    def test_dich_vu_du_moi_muc_rui_ro_thap(self):
        kq = rs.ra_soat(HD_DV_DU, "HĐ dịch vụ đầy đủ")
        loi = [m for m in kq["muc"] if m["trang_thai"] != "dat"]
        self.assertEqual(loi, [], loi)
        self.assertEqual(kq["tong_ket"]["muc_rui_ro"], "thap")
        self.assertEqual(_theo_ma(kq, "phat_vi_pham_pct")["gia_tri"], 8)
        self.assertEqual(_theo_ma(kq, "lai_cham_nam")["gia_tri"], 10)

    def test_phat_vuot_8_phan_tram(self):
        text = HD_DV_DU.replace("phạt vi phạm 8%", "phạt vi phạm 12%")
        m = _theo_ma(rs.ra_soat(text, loai="hop_dong_dich_vu"), "phat_vi_pham_pct")
        self.assertEqual(m["trang_thai"], "canh_bao")
        self.assertIn("Điều 301", m["can_cu"])

    def test_li_xang_dieu_khoan_han_che(self):
        text = ("HỢP ĐỒNG LI-XĂNG NHÃN HIỆU\nBên chuyển quyền và Bên nhận quyền sử dụng.\n"
                "Điều 1. Đối tượng\nNhãn hiệu ABC theo Giấy chứng nhận đăng ký số 12345.\n"
                "Điều 2. Phạm vi và thời hạn\nKhông độc quyền, lãnh thổ Việt Nam, thời hạn 5 năm.\n"
                "Điều 3. Phí\nPhí li-xăng 3% doanh thu, thanh toán hàng quý.\n"
                "Điều 4. Nghĩa vụ\nBên nhận quyền không được cải tiến sáng chế và phải mua "
                "nguyên liệu từ Bên chuyển quyền.\n")
        kq = rs.ra_soat(text, "HĐ li-xăng")
        self.assertEqual(kq["loai"], "hop_dong_li_xang")
        hc = _theo_ma(kq, "han_che_bat_hop_ly")
        self.assertEqual(hc["trang_thai"], "canh_bao")
        self.assertIn("Điều 144", hc["can_cu"])
        self.assertEqual(hc["dieu_khoan"], "Điều 4. Nghĩa vụ")
        kq2 = rs.ra_soat(text.replace("không được cải tiến sáng chế và phải mua nguyên liệu "
                                      "từ Bên chuyển quyền", "báo cáo doanh thu hàng quý"),
                         loai="hop_dong_li_xang")
        self.assertEqual(_theo_ma(kq2, "han_che_bat_hop_ly")["trang_thai"], "dat")

    def test_khac_chi_kiem_muc_chung(self):
        kq = rs.ra_soat("Biên bản họp ngày 01/02/2026. Các bên: Bên A, Bên B. Đại diện ký tên.",
                        "Biên bản")
        self.assertEqual(kq["loai"], "khac")
        self.assertEqual(sorted(m["ma"] for m in kq["muc"]),
                         sorted(c["ma"] for c in rs.MUC_CHUNG))
        self.assertEqual(kq["so_dieu_khoan"], 1)
        self.assertEqual(_theo_ma(kq, "tranh_chap")["trang_thai"], "canh_bao")
        self.assertEqual(_theo_ma(kq, "cac_ben")["dieu_khoan"], "Toàn văn")

    def test_van_ban_rong(self):
        kq = rs.ra_soat("", "")
        self.assertEqual(kq["loai"], "khac")
        self.assertEqual(kq["so_dieu_khoan"], 0)
        self.assertEqual(kq["tong_ket"]["canh_bao"], 4)
        self.assertEqual(kq["tong_ket"]["muc_rui_ro"], "cao")

    def test_loai_la_thi_tu_nhan_dien(self):
        kq = rs.ra_soat(HD_VAY, "HĐ vay", loai="khong_ton_tai")
        self.assertEqual(kq["loai"], "hop_dong_vay")


class DanhMucTests(unittest.TestCase):
    def test_du_loai_va_cau_truc(self):
        bat_buoc = {"hop_dong_lao_dong", "hop_dong_dich_vu", "hop_dong_mua_ban", "hop_dong_thue",
                    "hop_dong_vay", "hop_dong_hop_tac", "hop_dong_chuyen_nhuong_von", "nda",
                    "hop_dong_nhuong_quyen", "hop_dong_li_xang", "khac"}
        self.assertTrue(bat_buoc.issubset(rs.LOAI_HOP_DONG))
        for ma, cfg in rs.LOAI_HOP_DONG.items():
            with self.subTest(loai=ma):
                self.assertTrue(cfg["ten"])
                for dk in cfg["dieu_khoan"]:
                    for k in ("ma", "ten", "tu_khoa", "bat_buoc", "can_cu", "goi_y"):
                        self.assertIn(k, dk)
                    for p in dk["tu_khoa"]:
                        rs._bien_dich(p)
                for ng in cfg["nguong"]:
                    for k in ("ma", "ten", "regex", "kieu", "gia_tri", "don_vi", "can_cu",
                              "canh_bao", "goi_y"):
                        self.assertIn(k, ng)
                    self.assertIn(ng["kieu"], ("max", "min"))
                    self.assertGreaterEqual(rs._bien_dich(ng["regex"]).groups, 1)
                ma_muc = [d["ma"] for d in cfg["dieu_khoan"]] + [n["ma"] for n in cfg["nguong"]]
                self.assertEqual(len(ma_muc), len(set(ma_muc)), "mã mục trùng")

    def test_cau_hoi_tra_luat(self):
        c = rs.cau_hoi_tra_luat({"ten": "Thời gian thử việc (ngày)",
                                 "can_cu": "Điều 25 Bộ luật Lao động 2019"})
        self.assertIn("Điều 25 Bộ luật Lao động 2019", c)
        self.assertIn("thử việc", c)
        self.assertTrue(c.endswith("?"))
        self.assertIn("Pháp luật Việt Nam", rs.cau_hoi_tra_luat({"ten": "Bảo hành"}))


class XuatBaoCaoTests(unittest.TestCase):
    def test_docx_mo_duoc_co_bang(self):
        import docx
        kq = rs.ra_soat(HDLD, "HĐLĐ")
        payload = rs.xuat_bao_cao_docx(kq, "HĐLĐ Trần Thị B.docx", nguoi_lap="LS Nguyễn A",
                                       ngay=date(2026, 9, 15))
        self.assertIsInstance(payload, bytes)
        d = docx.Document(io.BytesIO(payload))
        van_ban = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("BÁO CÁO RÀ SOÁT RỦI RO", van_ban)
        self.assertIn("HĐLĐ Trần Thị B.docx", van_ban)
        self.assertIn("LS Nguyễn A", van_ban)
        self.assertIn("15/09/2026", van_ban)
        self.assertIn("luật sư", van_ban.lower())
        self.assertEqual(len(d.tables), 1)
        bang = d.tables[0]
        self.assertEqual(len(bang.rows), 1 + len(kq["muc"]))
        self.assertEqual(bang.rows[0].cells[2].text, "Trạng thái")
        self.assertIn(bang.rows[1].cells[2].text, ("CẢNH BÁO", "THIẾU"))
        self.assertEqual(d.styles["Normal"].font.name, "Times New Roman")

    def test_docx_ket_qua_rong(self):
        import docx
        payload = rs.xuat_bao_cao_docx({"ten_loai": "x", "do_tin_cay": 0, "muc": [],
                                        "tong_ket": {}}, "rỗng")
        d = docx.Document(io.BytesIO(payload))
        self.assertEqual(len(d.tables[0].rows), 1)


if __name__ == "__main__":
    unittest.main()
