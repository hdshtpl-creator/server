"""Văn bản chỉ có BẢN CHỮ .md (07/10/2026) — THUẦN, không chạm CSDL/LibreOffice.

Kệ luật nạp lô 17–19/09 có ~24.700 văn bản chỉ là .md. Trước đây nút Xem phát
text/plain (thấy nguyên "## Điều 1.", "</p>") và Tải về đưa tệp .md trần. Phải
đúng:
  · dàn trang KHÔNG sửa chữ: số khoản "1." và gạch "- " giữ nguyên;
  · phần đầu văn bản căn giữa, "Điều N." căn trái, cả hai in đậm;
  · thẻ HTML sót bị gỡ nhưng "<tên cơ quan>" trong biểu mẫu là chữ thật, giữ;
  · ký tự điều khiển của chữ OCR không làm hỏng cả tệp;
  · Xem → PDF; chuyển không được thì lùi về chữ thuần chứ không báo lỗi;
  · Tải về → .docx.
"""
import contextlib
import io
import os
import tempfile
import unittest
import unittest.mock as mock
from pathlib import Path

from app import xem_truoc

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:  # pragma: no cover
    Document = None

MAU = """## THỦ TƯỚNG CHÍNH PHỦ

Số: 36/2025/QĐ-TTg

## QUYẾT ĐỊNH

## Ban hành Hệ thống ngành kinh tế Việt Nam

Căn cứ Luật Thống kê số 89/2015/QH13;

## Điều 1. Phạm vi điều chỉnh

1. Phạm vi điều chỉnh
- Thông tư số 96/2002/TT-BTC ngày 24 tháng 10 năm 2002 hướng
dẫn thi hành</p>

Kính gửi: …… <tên cơ quan có thẩm quyền hoàn trả>

(**): Không sử dụng sản phẩm sau hội chợ\x0b

## Chương II

**Ghi chú cuối**
"""


@unittest.skipIf(Document is None, "Thiếu python-docx")
class DanTrang(unittest.TestCase):
    def setUp(self):
        d = Document(io.BytesIO(xem_truoc.docx_tu_markdown(MAU, "36/2025/QĐ-TTg")))
        self.doan = [(p.text, p.alignment, any(r.bold for r in p.runs)) for p in d.paragraphs]
        self.chu = "\n".join(t for t, _, _ in self.doan)

    def _doan(self, bat_dau):
        return next(x for x in self.doan if x[0].startswith(bat_dau))

    def test_phan_dau_can_giua_dieu_can_trai(self):
        for t in ("THỦ TƯỚNG CHÍNH PHỦ", "QUYẾT ĐỊNH", "Ban hành Hệ thống", "Chương II"):
            with self.subTest(t=t):
                _, can, dam = self._doan(t)
                self.assertEqual(can, WD_ALIGN_PARAGRAPH.CENTER)
                self.assertTrue(dam)
        _, can, dam = self._doan("Điều 1.")
        self.assertNotEqual(can, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(dam)

    def test_khong_sua_chu(self):
        self.assertNotIn("##", self.chu)
        self.assertIn("1. Phạm vi điều chỉnh", self.chu)
        self.assertIn("- Thông tư số 96/2002/TT-BTC", self.chu)
        # Hai dòng liền nhau là MỘT đoạn, giữ xuống dòng như tệp gốc.
        t, _, _ = self._doan("1. Phạm vi điều chỉnh")
        self.assertIn("hướng\ndẫn thi hành", t)

    def test_go_the_html_sot_giu_ngoac_nhon_that(self):
        self.assertNotIn("</p>", self.chu)
        self.assertIn("<tên cơ quan có thẩm quyền hoàn trả>", self.chu)

    def test_chu_thich_sao_khong_bi_in_dam_ca_dong(self):
        t, _, dam = self._doan("(**)")
        self.assertFalse(dam)
        self.assertNotIn("\x0b", t)
        _, _, dam = self._doan("Ghi chú cuối")
        self.assertTrue(dam)

    def test_van_ban_rong(self):
        Document(io.BytesIO(xem_truoc.docx_tu_markdown("")))


class PdfTuMarkdown(unittest.TestCase):
    def test_khoa_cache_rieng_va_dung_lai(self):
        with tempfile.TemporaryDirectory() as d, mock.patch.dict(os.environ, {"DATA_WORK": d}):
            md = Path(d) / "a.md"
            md.write_text("## QUYẾT ĐỊNH\n", encoding="utf-8")
            goi = []

            def gia_soffice(src, tmp, out):
                goi.append(src.read_bytes()[:2])
                out.write_bytes(b"%PDF")

            with mock.patch.object(xem_truoc, "_soffice_sang_pdf", gia_soffice):
                out = xem_truoc.pdf_tu_markdown(md, 42)
                # Khoá khác bản PDF của tệp Office cùng id, có số phiên bản
                # dàn trang để đổi cách dàn là bỏ cache cũ.
                self.assertEqual(out.name, f"md{xem_truoc.PHIEN_BAN_DAN_TRANG_MD}-42.pdf")
                self.assertEqual(goi, [b"PK"])  # nguồn đưa LibreOffice là .docx
                self.assertEqual(xem_truoc.pdf_tu_markdown(md, 42), out)
                self.assertEqual(len(goi), 1)

    def test_don_cache_crm_khong_dung_ban_web(self):
        self.assertFalse(f"md{xem_truoc.PHIEN_BAN_DAN_TRANG_MD}-1".startswith(
            xem_truoc.TIEN_TO_CACHE_TICH_HOP))


class CuaXemVaTai(unittest.TestCase):
    """Gọi thẳng hàm endpoint, vá quyền + CSDL (không ghi gì vào kho thật)."""

    def setUp(self):
        try:
            from app import api
        except (ImportError, RuntimeError) as e:  # pragma: no cover
            self.skipTest(f"Không nạp được app.api: {e}")
        self.api = api
        self.tmp = tempfile.TemporaryDirectory()
        self.md = Path(self.tmp.name) / "36-2025-QĐ-TTg_29092025.md"
        self.md.write_text(MAU, encoding="utf-8")
        self.vas = [
            mock.patch.object(api, "_can_tinh_nang", lambda *a, **k: None),
            mock.patch.object(api, "_original_file",
                              lambda doc_id, user: (self.md, "36-2025-QĐ-TTg_29092025")),
            mock.patch.object(api.db, "session", lambda **k: contextlib.nullcontext(None)),
            mock.patch.object(api.db, "audit", lambda *a, **k: None),
        ]
        for v in self.vas:
            v.start()

    def tearDown(self):
        for v in self.vas:
            v.stop()
        self.tmp.cleanup()

    def test_xem_ra_pdf(self):
        pdf = Path(self.tmp.name) / "x.pdf"
        pdf.write_bytes(b"%PDF")
        with mock.patch.object(self.api.xem_truoc, "pdf_tu_markdown", lambda p, k: pdf):
            r = self.api.files_preview(5, user={"id": 1})
        self.assertEqual(r.media_type, "application/pdf")
        self.assertEqual(Path(r.path), pdf)

    def test_chuyen_khong_duoc_thi_lui_ve_chu_thuan(self):
        def hong(p, k):
            raise xem_truoc.LoiXemTruoc("quá giờ")
        with mock.patch.object(self.api.xem_truoc, "pdf_tu_markdown", hong):
            r = self.api.files_preview(5, user={"id": 1})
        self.assertEqual(r.media_type, "text/plain; charset=utf-8")
        self.assertEqual(Path(r.path), self.md)

    @unittest.skipIf(Document is None, "Thiếu python-docx")
    def test_tai_ve_ra_word(self):
        r = self.api.files_download(5, user={"id": 1})
        self.assertIn("wordprocessingml", r.media_type)
        self.assertIn(".docx", r.headers["content-disposition"])
        self.assertNotIn(".md", r.headers["content-disposition"])
        Document(io.BytesIO(r.body))


if __name__ == "__main__":
    unittest.main()
