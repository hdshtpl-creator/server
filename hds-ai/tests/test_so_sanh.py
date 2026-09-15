"""
Test SO SÁNH VĂN BẢN + xuất .docx theo dõi thay đổi (app/so_sanh.py) — phần
THUẦN, không cần PostgreSQL/Ollama.

Chạy: python -m unittest tests.test_so_sanh -v

Ba thứ phải đứng vững:
  1. Tách token khả nghịch: nối lại ra đúng chuỗi gốc, nên diff mức từ không
     "chuẩn hoá" ngầm khoảng trắng thành thay đổi giả; hai vế diff nối lại ra
     đúng bản cũ / bản mới.
  2. Căn đoạn trước rồi so từ: thêm/xoá/sửa đoạn ra đúng op và đúng số đếm;
     bản rỗng, bản giống hệt, CRLF không làm vỡ.
  3. File .docx sinh ra là OOXML hợp lệ (python-docx mở được, ElementTree
     parse được), w:ins/w:del đúng chỗ, w:id không trùng, escape đúng, tiếng
     Việt nguyên vẹn, tiêu đề không nằm trong tracked change.
"""
import io
import re
import unittest
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timezone

import docx

from app import so_sanh as ss

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _noi(phan, ops):
    return "".join(p["text"] for p in phan if p["op"] in ops)


class TachTuTests(unittest.TestCase):
    def test_noi_lai_bang_chuoi_goc(self):
        for s in ["", "a", "a b", "  a\tb  \n c ", "Điều 1.  Phạm vi\tđiều chỉnh\n",
                  "x\r\ny", "   "]:
            with self.subTest(s=repr(s)):
                self.assertEqual("".join(ss.tach_tu(s)), s)

    def test_rong(self):
        self.assertEqual(ss.tach_tu(""), [])
        self.assertEqual(ss.tach_tu(None), [])

    def test_xen_ke_tu_va_khoang_trang(self):
        self.assertEqual(ss.tach_tu("a b"), ["a", " ", "b"])
        self.assertEqual(ss.tach_tu(" a  b"), [" ", "a", "  ", "b"])


class SoSanhDoanTests(unittest.TestCase):
    def test_doi_mot_tu(self):
        self.assertEqual(ss.so_sanh_doan("a b c", "a x c"), [
            {"op": "equal", "text": "a "},
            {"op": "delete", "text": "b"},
            {"op": "insert", "text": "x"},
            {"op": "equal", "text": " c"},
        ])

    def test_noi_lai_hai_ve(self):
        cap = [
            ("a b c", "a x c"),
            ("Bên A thanh toán trong 30 ngày.", "Bên A phải thanh toán trong vòng 45 ngày làm việc."),
            ("", "toàn bộ mới"),
            ("toàn bộ cũ", ""),
            ("giữ  hai  dấu cách", "giữ hai dấu cách"),
            ("x", "x"),
        ]
        for cu, moi in cap:
            with self.subTest(cu=cu, moi=moi):
                phan = ss.so_sanh_doan(cu, moi)
                self.assertEqual(_noi(phan, ("equal", "delete")), cu)
                self.assertEqual(_noi(phan, ("equal", "insert")), moi)

    def test_bang_nhau_mot_equal(self):
        self.assertEqual(ss.so_sanh_doan("a b c", "a b c"),
                         [{"op": "equal", "text": "a b c"}])

    def test_rong(self):
        self.assertEqual(ss.so_sanh_doan("", ""), [])
        self.assertEqual(ss.so_sanh_doan("", "a"), [{"op": "insert", "text": "a"}])
        self.assertEqual(ss.so_sanh_doan("a", ""), [{"op": "delete", "text": "a"}])

    def test_gop_token_lien_tiep(self):
        """Nhiều từ liền nhau bị xoá phải là MỘT mảnh, không phải từng chữ."""
        phan = ss.so_sanh_doan("a b c d e", "a e")
        self.assertEqual(phan, [
            {"op": "equal", "text": "a "},
            {"op": "delete", "text": "b c d "},
            {"op": "equal", "text": "e"},
        ])


class SoSanhVanBanTests(unittest.TestCase):
    def _ops(self, kq):
        return [d["op"] for d in kq["doan"]]

    def test_them_doan_giua(self):
        kq = ss.so_sanh_van_ban("Điều 1. A\nĐiều 2. B\nĐiều 3. C",
                                "Điều 1. A\nĐiều 2. B\nĐiều 2a. Bổ sung mới\nĐiều 3. C")
        self.assertEqual(self._ops(kq), ["equal", "equal", "insert", "equal"])
        d = kq["doan"][2]
        self.assertIsNone(d["cu"])
        self.assertEqual(d["moi"], "Điều 2a. Bổ sung mới")
        self.assertEqual(d["phan"], [{"op": "insert", "text": "Điều 2a. Bổ sung mới"}])
        tk = kq["thong_ke"]
        self.assertEqual((tk["them"], tk["xoa"]), (5, 0))  # Điều / 2a. / Bổ / sung / mới
        self.assertEqual((tk["doan_them"], tk["doan_xoa"], tk["doan_sua"]), (1, 0, 0))
        self.assertLess(tk["giong_nhau"], 1.0)

    def test_xoa_doan_cuoi(self):
        kq = ss.so_sanh_van_ban("Điều 1. A\nĐiều 2. B\nĐiều 3. Hiệu lực hợp đồng",
                                "Điều 1. A\nĐiều 2. B")
        self.assertEqual(self._ops(kq), ["equal", "equal", "delete"])
        d = kq["doan"][2]
        self.assertEqual(d["cu"], "Điều 3. Hiệu lực hợp đồng")
        self.assertIsNone(d["moi"])
        tk = kq["thong_ke"]
        self.assertEqual((tk["them"], tk["xoa"]), (0, 6))  # Điều / 3. / Hiệu / lực / hợp / đồng
        self.assertEqual((tk["doan_them"], tk["doan_xoa"], tk["doan_sua"]), (0, 1, 0))

    def test_sua_mot_tu_trong_doan(self):
        kq = ss.so_sanh_van_ban("Điều 1. A\nBên A thanh toán trong 30 ngày.",
                                "Điều 1. A\nBên A thanh toán trong 45 ngày.")
        self.assertEqual(self._ops(kq), ["equal", "replace"])
        d = kq["doan"][1]
        self.assertEqual(d["cu"], "Bên A thanh toán trong 30 ngày.")
        self.assertEqual(d["moi"], "Bên A thanh toán trong 45 ngày.")
        self.assertEqual(d["phan"], [
            {"op": "equal", "text": "Bên A thanh toán trong "},
            {"op": "delete", "text": "30"},
            {"op": "insert", "text": "45"},
            {"op": "equal", "text": " ngày."},
        ])
        tk = kq["thong_ke"]
        self.assertEqual((tk["them"], tk["xoa"]), (1, 1))
        self.assertEqual((tk["doan_them"], tk["doan_xoa"], tk["doan_sua"]), (0, 0, 1))

    def test_equal_mang_du_hai_ve(self):
        kq = ss.so_sanh_van_ban("A", "A")
        self.assertEqual(kq["doan"], [{"op": "equal", "cu": "A", "moi": "A",
                                       "phan": [{"op": "equal", "text": "A"}]}])

    def test_giong_het(self):
        vb = "# Hợp đồng\n\nĐiều 1. A\nĐiều 2. B"
        kq = ss.so_sanh_van_ban(vb, vb)
        self.assertTrue(all(d["op"] == "equal" for d in kq["doan"]))
        self.assertEqual(len(kq["doan"]), 4)  # đoạn rỗng vẫn là phần tử
        tk = kq["thong_ke"]
        self.assertEqual((tk["them"], tk["xoa"], tk["doan_them"], tk["doan_xoa"], tk["doan_sua"]),
                         (0, 0, 0, 0, 0))
        self.assertEqual(tk["giong_nhau"], 1.0)

    def test_rong_ca_hai(self):
        kq = ss.so_sanh_van_ban("", "")
        self.assertEqual(kq["doan"], [])
        tk = kq["thong_ke"]
        self.assertEqual((tk["them"], tk["xoa"]), (0, 0))
        self.assertEqual(tk["giong_nhau"], 1.0)
        self.assertIn("giống nhau 100%", ss.tom_tat_thay_doi(kq))

    def test_mot_ben_rong(self):
        kq = ss.so_sanh_van_ban("", "Điều 1. A\nĐiều 2. B")
        self.assertEqual(self._ops(kq), ["insert", "insert"])
        self.assertEqual(kq["thong_ke"]["doan_them"], 2)
        self.assertEqual(kq["thong_ke"]["them"], 6)
        self.assertEqual(kq["thong_ke"]["giong_nhau"], 0.0)

        kq = ss.so_sanh_van_ban("Điều 1. A\nĐiều 2. B", "")
        self.assertEqual(self._ops(kq), ["delete", "delete"])
        self.assertEqual(kq["thong_ke"]["doan_xoa"], 2)
        self.assertEqual(kq["thong_ke"]["xoa"], 6)

    def test_khoi_replace_ghep_cap_theo_thu_tu_va_phan_du(self):
        """3 đoạn cũ ↔ 1 đoạn mới: cặp đầu là replace, hai đoạn dư là delete."""
        kq = ss.so_sanh_van_ban("A\nB1 x\nB2 y\nB3 z\nC", "A\nB1 w\nC")
        self.assertEqual(self._ops(kq), ["equal", "replace", "delete", "delete", "equal"])
        self.assertEqual(kq["doan"][1]["cu"], "B1 x")
        self.assertEqual(kq["doan"][1]["moi"], "B1 w")
        tk = kq["thong_ke"]
        self.assertEqual((tk["doan_sua"], tk["doan_xoa"], tk["doan_them"]), (1, 2, 0))
        # Chiều ngược lại: phần dư là insert.
        kq = ss.so_sanh_van_ban("A\nB1 w\nC", "A\nB1 x\nB2 y\nB3 z\nC")
        self.assertEqual(self._ops(kq), ["equal", "replace", "insert", "insert", "equal"])

    def test_crlf_khong_thanh_khac_nhau_gia(self):
        kq = ss.so_sanh_van_ban("Điều 1. A\r\nĐiều 2. B", "Điều 1. A\nĐiều 2. B")
        self.assertEqual(self._ops(kq), ["equal", "equal"])
        self.assertEqual(kq["thong_ke"]["giong_nhau"], 1.0)


class XuatDocxTests(unittest.TestCase):
    CU = ("# HỢP ĐỒNG DỊCH VỤ PHÁP LÝ\n"
          "Điều 1. Bên A & Bên B <cùng> ký kết.\n"
          "Điều 2. Giá trị 100 triệu đồng.\n"
          "Điều 3. Điều khoản bị bỏ hẳn.\n"
          "Điều 5. Hiệu lực.")
    MOI = ("# HỢP ĐỒNG DỊCH VỤ PHÁP LÝ\n"
           "Điều 1. Bên A & Bên B <cùng> ký kết.\n"
           "Điều 2. Giá trị 120 triệu đồng.\n"
           "Điều 5. Hiệu lực.\n"
           "Điều 6. Điều khoản thêm mới.")
    TIEU_DE = "BẢN ĐỐI CHIẾU THAY ĐỔI"

    @classmethod
    def setUpClass(cls):
        cls.data = ss.xuat_docx_theo_doi(
            cls.CU, cls.MOI, tieu_de=cls.TIEU_DE,
            ngay=datetime(2026, 9, 15, 8, 30, 0, tzinfo=timezone.utc))
        with zipfile.ZipFile(io.BytesIO(cls.data)) as z:
            cls.ten_file = z.namelist()
            cls.xml = z.read("word/document.xml").decode("utf-8")
            cls.cac_phan = {n: z.read(n).decode("utf-8") for n in cls.ten_file}

    def test_mo_bang_python_docx(self):
        doc = docx.Document(io.BytesIO(self.data))
        self.assertEqual(doc.paragraphs[0].text, self.TIEU_DE)
        # Đoạn không đổi vẫn đọc ra chữ thường (python-docx bỏ run trong
        # w:ins/w:del, nên đọc được nghĩa là run nằm ngoài tracked change).
        self.assertIn("Điều 1. Bên A & Bên B <cùng> ký kết.",
                      [p.text for p in doc.paragraphs])

    def test_cac_phan_zip_va_xml_hop_le(self):
        for ten in ["[Content_Types].xml", "_rels/.rels", "word/document.xml",
                    "word/_rels/document.xml.rels", "word/styles.xml", "docProps/core.xml"]:
            self.assertIn(ten, self.ten_file)
            ET.fromstring(self.cac_phan[ten].encode("utf-8"))  # không ném lỗi
        self.assertIn(self.TIEU_DE, self.cac_phan["docProps/core.xml"])
        self.assertIn("<dc:creator>HDS AI</dc:creator>", self.cac_phan["docProps/core.xml"])
        self.assertIn('w:val="vi-VN"', self.cac_phan["word/styles.xml"])
        self.assertIn("Times New Roman", self.cac_phan["word/styles.xml"])

    def test_co_ins_del_va_tac_gia(self):
        self.assertIn("<w:ins ", self.xml)
        self.assertIn("<w:del ", self.xml)
        self.assertIn("<w:delText", self.xml)
        self.assertIn('w:author="HDS AI"', self.xml)
        self.assertIn('w:date="2026-09-15T08:30:00Z"', self.xml)

    def test_id_khong_trung(self):
        ids = re.findall(r'w:id="(\d+)"', self.xml)
        self.assertGreater(len(ids), 3)
        self.assertEqual(len(ids), len(set(ids)))

    def test_escape_ky_tu_dac_biet(self):
        self.assertIn("Bên A &amp; Bên B &lt;cùng&gt;", self.xml)
        root = ET.fromstring(self.xml.encode("utf-8"))
        chu = [t.text for t in root.iter("{%s}t" % NS["w"])]
        self.assertIn("Điều 1. Bên A & Bên B <cùng> ký kết.", chu)

    def test_sua_mot_tu_chi_danh_dau_mot_tu(self):
        self.assertIn('<w:delText xml:space="preserve">100</w:delText>', self.xml)
        self.assertIn('<w:t xml:space="preserve">120</w:t>', self.xml)

    def test_xoa_ca_doan_co_dau_doan_trong_ppr(self):
        # Dấu đoạn bị xoá: w:del rỗng đứng đầu trong w:pPr/w:rPr, rồi run
        # trong w:del với w:delText.
        m = re.search(r'<w:p><w:pPr><w:rPr><w:del w:id="\d+" w:author="HDS AI" w:date="[^"]+"/></w:rPr></w:pPr>'
                      r'<w:del [^>]+><w:r><w:delText xml:space="preserve">Điều 3\. Điều khoản bị bỏ hẳn\.</w:delText></w:r></w:del></w:p>',
                      self.xml)
        self.assertIsNotNone(m, self.xml)

    def test_them_ca_doan_co_dau_doan_trong_ppr(self):
        m = re.search(r'<w:p><w:pPr><w:rPr><w:ins w:id="\d+" w:author="HDS AI" w:date="[^"]+"/></w:rPr></w:pPr>'
                      r'<w:ins [^>]+><w:r><w:t xml:space="preserve">Điều 6\. Điều khoản thêm mới\.</w:t></w:r></w:ins></w:p>',
                      self.xml)
        self.assertIsNotNone(m, self.xml)

    def test_tieu_de_khong_nam_trong_tracked_change(self):
        root = ET.fromstring(self.xml.encode("utf-8"))
        p0 = root.find("w:body/w:p", NS)
        self.assertEqual("".join(t.text for t in p0.iter("{%s}t" % NS["w"])), self.TIEU_DE)
        self.assertIsNone(p0.find(".//w:ins", NS))
        self.assertIsNone(p0.find(".//w:del", NS))
        self.assertIsNotNone(p0.find("w:pPr/w:jc[@w:val='center']", NS))
        self.assertIsNotNone(p0.find("w:r/w:rPr/w:b", NS))

    def test_heading_markdown(self):
        self.assertNotIn("# HỢP ĐỒNG", self.xml)
        self.assertIn('<w:sz w:val="28"/>', self.xml)
        m = re.search(r'<w:p><w:pPr><w:keepNext/></w:pPr><w:r><w:rPr><w:b/><w:bCs/><w:sz w:val="28"/>'
                      r'<w:szCs w:val="28"/></w:rPr><w:t xml:space="preserve">HỢP ĐỒNG DỊCH VỤ PHÁP LÝ</w:t></w:r></w:p>',
                      self.xml)
        self.assertIsNotNone(m, self.xml)

    def test_tieng_viet_nguyen_ven(self):
        self.assertIn("Điều khoản bị bỏ hẳn", self.xml)
        self.assertIn("Điều khoản thêm mới", self.xml)
        self.assertIn("HỢP ĐỒNG DỊCH VỤ PHÁP LÝ", self.xml)

    def test_tac_gia_va_ngay_tuy_chon(self):
        data = ss.xuat_docx_theo_doi("a b", "a c", tac_gia="Luật sư Hồng",
                                     ngay=datetime(2026, 1, 2, 3, 4, 5))
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml").decode("utf-8")
            core = z.read("docProps/core.xml").decode("utf-8")
        self.assertIn('w:author="Luật sư Hồng"', xml)
        self.assertIn('w:date="2026-01-02T03:04:05Z"', xml)
        self.assertIn("<dc:creator>Luật sư Hồng</dc:creator>", core)
        docx.Document(io.BytesIO(data))

    def test_tac_gia_rong_ve_mac_dinh(self):
        data = ss.xuat_docx_theo_doi("a", "b", tac_gia="  ")
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            self.assertIn('w:author="HDS AI"', z.read("word/document.xml").decode("utf-8"))

    def test_giong_het_khong_co_tracked_change(self):
        data = ss.xuat_docx_theo_doi(self.CU, self.CU)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        self.assertNotIn("<w:ins", xml)
        self.assertNotIn("<w:del", xml)
        doc = docx.Document(io.BytesIO(data))
        self.assertEqual(doc.paragraphs[0].text, "HỢP ĐỒNG DỊCH VỤ PHÁP LÝ")

    def test_van_ban_rong_khong_loi(self):
        data = ss.xuat_docx_theo_doi("", "")
        doc = docx.Document(io.BytesIO(data))
        self.assertEqual(len(doc.paragraphs), 0)
        data = ss.xuat_docx_theo_doi("", "Điều 1. Mới", tieu_de="T")
        doc = docx.Document(io.BytesIO(data))
        self.assertEqual(doc.paragraphs[0].text, "T")

    def test_ky_tu_dieu_khien_bi_loai(self):
        data = ss.xuat_docx_theo_doi("a\x0cb", "a\x0cc")
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            ET.fromstring(z.read("word/document.xml"))


class TomTatTests(unittest.TestCase):
    def test_dang_cau(self):
        kq = ss.so_sanh_van_ban("A b c\nX y\nCuối", "A b d\nX y\nMới thêm\nCuối")
        s = ss.tom_tat_thay_doi(kq)
        self.assertRegex(s, r"^Thêm \d+ từ, xoá \d+ từ; \d+ đoạn sửa, \d+ đoạn thêm, "
                            r"\d+ đoạn xoá; giống nhau \d+%$")
        self.assertEqual(s, "Thêm 3 từ, xoá 1 từ; 1 đoạn sửa, 1 đoạn thêm, 0 đoạn xoá; "
                            f"giống nhau {round(kq['thong_ke']['giong_nhau'] * 100)}%")

    def test_gia_tri_dung(self):
        s = ss.tom_tat_thay_doi({"thong_ke": {"them": 12, "xoa": 4, "doan_sua": 2,
                                              "doan_them": 1, "doan_xoa": 0,
                                              "giong_nhau": 0.912}})
        self.assertEqual(s, "Thêm 12 từ, xoá 4 từ; 2 đoạn sửa, 1 đoạn thêm, 0 đoạn xoá; "
                            "giống nhau 91%")

    def test_thieu_thong_ke_khong_loi(self):
        self.assertIn("0 từ", ss.tom_tat_thay_doi({}))


if __name__ == "__main__":
    unittest.main()
