"""
ra_soat_rui_ro.py — RÀ SOÁT RỦI RO TÀI LIỆU: nhân viên tải một hợp đồng lên,
hệ thống đối chiếu với DANH MỤC ĐIỀU KHOẢN CHUẨN theo loại hợp đồng (điều
khoản nào bắt buộc / nên có, ngưỡng nào bất thường theo luật Việt Nam) và trả
về bảng từng mục: Đạt / Cảnh báo / Thiếu + giải thích + đề xuất sửa + căn cứ.

Mô-đun này là phần TẤT ĐỊNH: thuần Python, không CSDL, không mạng, không gọi
model. Tra thêm đoạn luật thật trong kho và nhờ model diễn giải là việc của
lớp bên ngoài (dùng cau_hoi_tra_luat() để đặt câu hỏi cho từng mục).

Nguyên tắc so khớp: mọi regex trong danh mục viết KHÔNG DẤU, chữ thường, và
được chạy trên bo_dau(text) — nhờ vậy văn bản gõ có dấu, không dấu, viết hoa,
OCR mất dấu đều khớp như nhau. Dấu cách trong regex tự nở thành ``\\s+`` để
không lệ thuộc vào xuống dòng / nhiều khoảng trắng của bản OCR.

Căn cứ pháp lý chỉ ghi số điều ĐÃ CHẮC; chỗ không chắc ghi tên luật không
kèm số điều, hoặc ghi "thông lệ". Ngưỡng là con số cứng của luật hiện hành
tại thời điểm viết (BLLĐ 2019, BLDS 2015, LTM 2005, Luật SHTT, Luật DN 2020).
"""
import io
import re
import unicodedata
from datetime import date

# Độ dài đoạn trích quanh chỗ khớp (ký tự) — đủ để luật sư thấy ngữ cảnh mà
# không chép cả điều khoản vào bảng.
TRICH_MAX = 200
# Chỉ đọc phần đầu văn bản khi nhận diện loại: tên loại hợp đồng nằm ở tiêu đề
# và phần mở đầu; đọc cả văn bản thì phụ lục / điều khoản chung gây nhiễu.
NHAN_DIEN_DAU = 3000


def bo_dau(text: str) -> str:
    """Bỏ dấu tiếng Việt + chữ thường. Giữ NGUYÊN độ dài chuỗi (mỗi ký tự có
    dấu → đúng một ký tự không dấu) để vị trí khớp trên bản không dấu dùng lại
    được trên bản gốc khi cắt đoạn trích."""
    s = unicodedata.normalize("NFD", (text or "").lower())
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return s.replace("đ", "d")


def _bien_dich(mau: str) -> re.Pattern:
    """Regex danh mục viết với dấu cách thường; nở thành ``\\s+`` để khớp qua
    xuống dòng và nhiều khoảng trắng."""
    return re.compile(mau.replace(" ", r"\s+"), re.I)


def _so(chuoi: str) -> float | None:
    """'5.000.000' → 5000000; '8,5' → 8.5; '12' → 12. Trả None khi không đọc
    được — không đoán mò số để rồi cảnh báo sai."""
    s = (chuoi or "").strip()
    if not s:
        return None
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# Bốn mục CHUNG cho mọi loại (kể cả "khac"): thiếu → canh_bao.
# ---------------------------------------------------------------------------
MUC_CHUNG = [
    {"ma": "cac_ben", "ten": "Thông tin các bên ký kết",
     "tu_khoa": [r"\bben (a|b|c)\b",
                 r"\bben (ban|mua|thue|cho thue|vay|cho vay|cung cap|cung ung|su dung|"
                 r"nhuong quyen|nhan quyen|chuyen nhuong|nhan chuyen nhuong|giao|nhan|"
                 r"tiet lo|chuyen quyen|nhan quyen su dung|dat hang|thue dich vu)\b",
                 r"nguoi lao dong|nguoi su dung lao dong"],
     "can_cu": "Điều 385, 398 Bộ luật Dân sự 2015",
     "giai_thich": "Không thấy tên/vai trò các bên (Bên A/Bên B hoặc tên vai như Bên "
                   "bán, Bên thuê…) — không xác định được chủ thể chịu ràng buộc.",
     "goi_y": "Ghi đủ tên, địa chỉ, mã số doanh nghiệp/CCCD, người đại diện và "
              "chức vụ của từng bên ngay phần mở đầu."},
    {"ma": "ngay_ky_hieu_luc", "ten": "Ngày ký / thời điểm có hiệu lực",
     "tu_khoa": [r"hieu luc", r"ngay ky",
                 r"ngay \d{1,2} thang \d{1,2} nam \d{4}",
                 r"\b\d{1,2}[/.-]\d{1,2}[/.-](19|20)\d{2}\b"],
     "can_cu": "Điều 401 Bộ luật Dân sự 2015",
     "giai_thich": "Không thấy ngày ký hoặc điều khoản hiệu lực — khó xác định "
                   "thời điểm phát sinh quyền, nghĩa vụ và thời hiệu.",
     "goi_y": "Ghi ngày ký ở phần mở đầu và một điều 'Hiệu lực hợp đồng' nêu rõ "
              "thời điểm bắt đầu."},
    {"ma": "tranh_chap", "ten": "Giải quyết tranh chấp",
     "tu_khoa": [r"tranh chap", r"toa an", r"trong tai"],
     "can_cu": "Bộ luật Tố tụng dân sự 2015; Luật Trọng tài thương mại 2010",
     "giai_thich": "Không có điều khoản giải quyết tranh chấp — khi có mâu thuẫn "
                   "phải mất thời gian xác định cơ quan/phương thức giải quyết.",
     "goi_y": "Thêm điều khoản: thương lượng → hoà giải → Toà án có thẩm quyền "
              "(hoặc Trọng tài, ghi rõ trung tâm và quy tắc)."},
    {"ma": "chu_ky_dai_dien", "ten": "Chữ ký / người đại diện",
     "tu_khoa": [r"dai dien", r"chu ky", r"ky ten", r"ky, ghi ro", r"ghi ro ho ten",
                 r"chuc vu"],
     "can_cu": "Điều 134–138 Bộ luật Dân sự 2015; Điều 12 Luật Doanh nghiệp 2020",
     "giai_thich": "Không thấy phần ký/đại diện — hợp đồng do người không có "
                   "thẩm quyền ký có thể bị vô hiệu hoặc không ràng buộc pháp nhân.",
     "goi_y": "Thêm phần ký của người đại diện theo pháp luật (hoặc theo uỷ quyền, "
              "kèm giấy uỷ quyền) của mỗi bên, ghi rõ họ tên, chức vụ."},
]


# ---------------------------------------------------------------------------
# Ngưỡng dùng chung cho hợp đồng thương mại / dân sự (phạt, lãi chậm)
# ---------------------------------------------------------------------------
def _nguong_phat_va_lai():
    return [
        {"ma": "phat_vi_pham_pct", "ten": "Mức phạt vi phạm",
         "regex": r"phat[^.\n]{0,80}?(\d{1,3}(?:[.,]\d+)?)\s*%",
         "kieu": "max", "gia_tri": 8, "don_vi": "%",
         "can_cu": "Điều 301 Luật Thương mại 2005",
         "canh_bao": "Giữa các thương nhân, mức phạt vi phạm tối đa 8% giá trị phần "
                     "nghĩa vụ bị vi phạm (Điều 301 LTM 2005); phần vượt có thể không "
                     "được Toà án chấp nhận. Hợp đồng dân sự thuần tuý thì các bên tự "
                     "thoả thuận (Điều 418 BLDS 2015).",
         "goi_y": "Hạ mức phạt về tối đa 8% giá trị phần nghĩa vụ bị vi phạm, hoặc tách "
                  "riêng phạt và bồi thường thiệt hại."},
        {"ma": "lai_cham_nam", "ten": "Lãi chậm thanh toán (theo năm)",
         "regex": r"(?:(?:lai|lai suat)[^.\n]{0,40}?cham|cham[^.\n]{0,40}?lai)[^.\n]{0,60}?"
                  r"(\d{1,3}(?:[.,]\d+)?)\s*%\s*(?:/|moi|mot|tren|trong mot)?\s*nam",
         "kieu": "max", "gia_tri": 20, "don_vi": "%/năm",
         "can_cu": "Điều 357, 468 Bộ luật Dân sự 2015; Điều 306 Luật Thương mại 2005",
         "canh_bao": "Lãi do chậm thanh toán theo thoả thuận không được vượt 20%/năm "
                     "(Điều 357 dẫn chiếu Điều 468 BLDS 2015); phần vượt vô hiệu.",
         "goi_y": "Đưa lãi chậm thanh toán về ≤ 20%/năm hoặc theo lãi suất nợ quá hạn "
                  "trung bình trên thị trường (Điều 306 LTM 2005)."},
        {"ma": "lai_cham_thang", "ten": "Lãi chậm thanh toán (theo tháng)",
         "regex": r"(?:(?:lai|lai suat)[^.\n]{0,40}?cham|cham[^.\n]{0,40}?lai)[^.\n]{0,60}?"
                  r"(\d{1,2}(?:[.,]\d+)?)\s*%\s*(?:/|moi|mot|tren|trong mot)\s*thang",
         "kieu": "max", "gia_tri": 1.67, "don_vi": "%/tháng",
         "can_cu": "Điều 357, 468 Bộ luật Dân sự 2015",
         "canh_bao": "Quy đổi theo năm vượt trần 20%/năm (≈1,67%/tháng) của Điều 468 "
                     "BLDS 2015.",
         "goi_y": "Đưa lãi chậm thanh toán về ≤ 1,67%/tháng (20%/năm)."},
    ]


def _dk(ma, ten, tu_khoa, bat_buoc, can_cu, goi_y, **them):
    d = {"ma": ma, "ten": ten, "tu_khoa": list(tu_khoa), "bat_buoc": bool(bat_buoc),
         "can_cu": can_cu, "goi_y": goi_y}
    d.update(them)
    return d


_TRANH_CHAP = _dk("tranh_chap", "Giải quyết tranh chấp",
                  [r"tranh chap", r"toa an", r"trong tai"], False,
                  "Bộ luật Tố tụng dân sự 2015; Luật Trọng tài thương mại 2010",
                  "Thêm điều khoản thương lượng → hoà giải → Toà án/Trọng tài "
                  "(ghi rõ trung tâm, quy tắc, ngôn ngữ nếu chọn trọng tài).")
_PHAT_BOI_THUONG = _dk("phat_boi_thuong", "Phạt vi phạm và bồi thường thiệt hại",
                       [r"phat vi pham|phat hop dong|muc phat|tien phat", r"boi thuong"],
                       False,
                       "Điều 418 Bộ luật Dân sự 2015; Điều 300–302 Luật Thương mại 2005",
                       "Thêm điều khoản phạt vi phạm (ghi rõ mức, cách tính) và bồi "
                       "thường thiệt hại thực tế; lưu ý trần 8% giữa thương nhân.")
_BAT_KHA_KHANG = _dk("bat_kha_khang", "Bất khả kháng",
                     [r"bat kha khang", r"force majeure"], False,
                     "Điều 156, 351 Bộ luật Dân sự 2015; Điều 294 Luật Thương mại 2005",
                     "Thêm điều khoản bất khả kháng: định nghĩa, nghĩa vụ thông báo, "
                     "hệ quả (miễn trách, gia hạn, chấm dứt).")
_QUYEN_NGHIA_VU = _dk("quyen_nghia_vu", "Quyền và nghĩa vụ của các bên",
                      [r"quyen va nghia vu|quyen, nghia vu|quyen (va )?trach nhiem",
                       r"nghia vu (cua )?ben|trach nhiem (cua )?ben"], False,
                      "Điều 398 Bộ luật Dân sự 2015",
                      "Tách riêng quyền và nghĩa vụ của từng bên thành điều khoản "
                      "để dễ đối chiếu khi có vi phạm.")
# Dấu hiệu điều khoản MỘT CHIỀU / bất lợi cho một bên — CÓ mới là rủi ro. Máy
# chỉ nhận diện theo cụm chữ, kết luận cuối vẫn là của luật sư; nhưng không có
# nhóm này thì "Bảo mật: bên B được dùng thông tin bên A không cần chấp thuận"
# được chấm ĐẠT chỉ vì có chữ "bảo mật" (kiểm thử 18/09/2026).
_MOT_CHIEU_CHAM_DUT = _dk(
    "mot_chieu_cham_dut", "Quyền chấm dứt một chiều",
    [r"khong (duoc|co quyen) (don phuong )?cham dut[^.\n]{0,60}(trong )?moi truong hop",
     r"cham dut[^.\n]{0,60}(bat ky|bat cu) (luc|thoi diem) nao[^.\n]{0,60}khong can (bao|thong bao) truoc",
     r"khong can (bao|thong bao) truoc[^.\n]{0,40}cham dut"],
    False, "Điều 3 (bình đẳng, thiện chí), Điều 428 Bộ luật Dân sự 2015",
    "Cân bằng quyền chấm dứt: cả hai bên đều được đơn phương chấm dứt khi bên kia vi "
    "phạm nghiêm trọng, có thời hạn báo trước và cách xử lý phần đã thực hiện.",
    nguoc=True,
    giai_thich_co="Điều khoản chấm dứt chỉ trói MỘT bên (bên kia chấm dứt lúc nào cũng "
                  "được / bên này không được chấm dứt trong mọi trường hợp) — bất lợi "
                  "rõ cho bên bị trói, dễ bị coi là trái nguyên tắc bình đẳng, thiện chí.")
_DUNG_THONG_TIN_KHONG_CHAP_THUAN = _dk(
    "dung_thong_tin_khong_chap_thuan", "Dùng / tiết lộ thông tin bên kia không cần chấp thuận",
    [r"(su dung|dung|cong bo|tiet lo|chia se|khai thac)[^.\n]{0,80}thong tin[^.\n]{0,80}"
     r"(khong can|ma khong can|khong phai)[^.\n]{0,20}(su )?(chap thuan|dong y|xin phep)"],
    False, "Điều 38 Bộ luật Dân sự 2015; Luật Bảo vệ dữ liệu cá nhân; Điều 517 BLDS 2015",
    "Bỏ quyền tự ý dùng/tiết lộ thông tin; quy định rõ mục đích, phạm vi, và phải có "
    "chấp thuận bằng văn bản của bên có thông tin.",
    nguoc=True,
    giai_thich_co="Cho phép một bên dùng/tiết lộ thông tin của bên kia mà KHÔNG cần "
                  "chấp thuận — ngược với mục đích của điều khoản bảo mật và có thể vi "
                  "phạm quy định về dữ liệu cá nhân, bí mật kinh doanh.")
_HIEU_LUC_KHONG_CHU_KY = _dk(
    "hieu_luc_khong_chu_ky", "Hiệu lực không cần chữ ký người đại diện",
    [r"(co )?hieu luc[^.\n]{0,60}khong can (co )?chu ky",
     r"khong can (co )?chu ky[^.\n]{0,60}(nguoi dai dien|dai dien theo phap luat)"],
    False, "Điều 117, 401 Bộ luật Dân sự 2015; Điều 12 Luật Doanh nghiệp 2020",
    "Hợp đồng chỉ có hiệu lực khi được người có thẩm quyền (người đại diện theo pháp "
    "luật hoặc người được uỷ quyền hợp lệ) ký; bỏ câu này và bổ sung uỷ quyền nếu "
    "người ký không phải đại diện theo pháp luật.",
    nguoc=True,
    giai_thich_co="Tuyên bố hợp đồng có hiệu lực dù không có chữ ký của người đại diện "
                  "theo pháp luật — dễ dẫn tới hợp đồng vô hiệu do người ký không có "
                  "thẩm quyền.")

_CHAM_DUT = _dk("cham_dut", "Chấm dứt / đơn phương chấm dứt hợp đồng",
                [r"cham dut", r"huy bo (hop dong)?|huy hop dong", r"don phuong"], False,
                "Điều 422–428 Bộ luật Dân sự 2015",
                "Nêu các trường hợp chấm dứt, quyền đơn phương chấm dứt, thời hạn "
                "báo trước và hậu quả (thanh toán phần đã thực hiện, hoàn trả).")


# ---------------------------------------------------------------------------
# DANH MỤC ĐIỀU KHOẢN CHUẨN theo loại hợp đồng
# nhan_dien: regex (không dấu) để đoán loại từ tiêu đề + phần đầu văn bản.
# dieu_khoan: tu_khoa là danh sách regex — khớp BẤT KỲ regex nào trong một mục
#             đã tách là coi như có điều khoản đó.
# nguong: regex có nhóm bắt số; kieu max/min; không tìm thấy số thì bỏ qua.
# ---------------------------------------------------------------------------
LOAI_HOP_DONG: dict[str, dict] = {}

LOAI_HOP_DONG["hop_dong_lao_dong"] = {
    "ten": "Hợp đồng lao động",
    "nhan_dien": [r"hop dong lao dong", r"\bhdld\b", r"nguoi lao dong",
                  r"nguoi su dung lao dong", r"thu viec", r"bao hiem xa hoi|bhxh"],
    "dieu_khoan": [
        _dk("nsdld", "Tên, địa chỉ người sử dụng lao động và người giao kết",
            [r"nguoi su dung lao dong", r"(cong ty|doanh nghiep)[^\n]{0,200}dia chi",
             r"ma so (doanh nghiep|thue)"], True,
            "Điểm a khoản 1 Điều 21 Bộ luật Lao động 2019",
            "Ghi tên, địa chỉ trụ sở của người sử dụng lao động và họ tên, chức danh "
            "người giao kết bên phía người sử dụng lao động."),
        _dk("nld", "Họ tên, ngày sinh, giới tính, nơi cư trú, CCCD người lao động",
            [r"nguoi lao dong[^\n]{0,200}(sinh|cccd|can cuoc|cmnd)",
             r"can cuoc cong dan|\bcccd\b|\bcmnd\b|chung minh nhan dan|ho chieu",
             r"ngay sinh|sinh ngay|nam sinh"], True,
            "Điểm b khoản 1 Điều 21 Bộ luật Lao động 2019",
            "Ghi đủ họ tên, ngày tháng năm sinh, giới tính, nơi cư trú, số CCCD/hộ "
            "chiếu của người lao động."),
        _dk("cong_viec", "Công việc và địa điểm làm việc",
            [r"cong viec", r"chuc danh|vi tri (cong viec|lam viec)", r"dia diem lam viec"],
            True, "Điểm c khoản 1 Điều 21 Bộ luật Lao động 2019",
            "Ghi rõ chức danh, mô tả công việc, địa điểm làm việc (hoặc phạm vi địa "
            "điểm nếu làm ở nhiều nơi)."),
        _dk("thoi_han", "Thời hạn hợp đồng",
            [r"thoi han (cua )?hop dong", r"loai hop dong",
             r"(xac dinh|khong xac dinh) thoi han", r"hop dong co thoi han"], True,
            "Điểm d khoản 1 Điều 21; Điều 20 Bộ luật Lao động 2019",
            "Nêu loại hợp đồng (xác định / không xác định thời hạn) và thời điểm bắt "
            "đầu, kết thúc; xác định thời hạn thì không quá 36 tháng."),
        _dk("luong", "Mức lương, hình thức trả lương, thời hạn trả lương, phụ cấp",
            [r"muc luong|tien luong|luong (co ban|chinh|thang|thoa thuan)",
             r"hinh thuc tra luong|tra luong", r"phu cap"], True,
            "Điểm đ khoản 1 Điều 21; Điều 90, 94–96 Bộ luật Lao động 2019",
            "Ghi mức lương theo công việc/chức danh, hình thức trả (tiền mặt/chuyển "
            "khoản), kỳ hạn trả và các khoản phụ cấp, bổ sung."),
        _dk("nang_luong", "Chế độ nâng bậc, nâng lương",
            [r"nang (bac|luong)", r"tang luong", r"xet (tang|nang) luong|dieu chinh luong"],
            True, "Điểm e khoản 1 Điều 21 Bộ luật Lao động 2019",
            "Nêu điều kiện, thời điểm, mức nâng bậc/nâng lương hoặc dẫn chiếu quy "
            "chế lương, thoả ước lao động tập thể."),
        _dk("thoi_gio", "Thời giờ làm việc, thời giờ nghỉ ngơi",
            [r"thoi gio( lam viec)?", r"gio lam viec|gio\s*/\s*ngay|gio moi ngay",
             r"nghi ngoi|nghi hang tuan|nghi phep|nghi le"], True,
            "Điểm g khoản 1 Điều 21; Điều 105–115 Bộ luật Lao động 2019",
            "Ghi số giờ làm việc mỗi ngày/tuần, ngày nghỉ hằng tuần, nghỉ lễ, nghỉ "
            "hằng năm hoặc dẫn chiếu nội quy lao động."),
        _dk("bao_ho", "Trang bị bảo hộ lao động",
            [r"bao ho lao dong", r"trang bi (bao ho|phuong tien)",
             r"an toan(,| va)? ve sinh lao dong|an toan lao dong"], True,
            "Điểm h khoản 1 Điều 21 Bộ luật Lao động 2019; Luật An toàn, vệ sinh "
            "lao động 2015",
            "Nêu trang bị bảo hộ được cấp hoặc dẫn chiếu quy định về an toàn, vệ "
            "sinh lao động của đơn vị."),
        _dk("bao_hiem", "Bảo hiểm xã hội, bảo hiểm y tế, bảo hiểm thất nghiệp",
            [r"bao hiem xa hoi|\bbhxh\b", r"bao hiem y te|\bbhyt\b",
             r"bao hiem that nghiep|\bbhtn\b"], True,
            "Điểm i khoản 1 Điều 21; Điều 168 Bộ luật Lao động 2019",
            "Ghi rõ tham gia BHXH, BHYT, BHTN theo quy định; tỷ lệ đóng của mỗi bên."),
        _dk("dao_tao", "Đào tạo, bồi dưỡng, nâng cao trình độ, kỹ năng nghề",
            [r"dao tao", r"boi duong|nang cao (trinh do|tay nghe|ky nang)"], True,
            "Điểm k khoản 1 Điều 21; Điều 62 Bộ luật Lao động 2019",
            "Nêu quyền, nghĩa vụ về đào tạo; nếu có cam kết đào tạo thì lập hợp đồng "
            "đào tạo riêng (Điều 62)."),
        _dk("cham_dut", "Chấm dứt hợp đồng lao động",
            [r"cham dut", r"don phuong", r"nghi viec|thoi viec"], False,
            "Điều 34–36 Bộ luật Lao động 2019",
            "Dẫn chiếu các trường hợp chấm dứt, thời hạn báo trước khi đơn phương "
            "chấm dứt (Điều 35, 36) và trợ cấp thôi việc."),
    ],
    "nguong": [
        {"ma": "thu_viec_ngay", "ten": "Thời gian thử việc (ngày)",
         "regex": r"thu viec[^.\n]{0,80}?(\d{1,3}) ngay", "kieu": "max",
         "gia_tri": 60, "don_vi": "ngày", "can_cu": "Điều 25 Bộ luật Lao động 2019",
         "canh_bao": "Thử việc tối đa 60 ngày với công việc cần trình độ cao đẳng trở "
                     "lên (30 ngày trung cấp, 6 ngày công việc khác; 180 ngày chỉ với "
                     "người quản lý doanh nghiệp).",
         "goi_y": "Rút thời gian thử việc về đúng trần theo tính chất công việc; chỉ "
                  "thử việc một lần cho một công việc."},
        {"ma": "thu_viec_thang", "ten": "Thời gian thử việc (tháng)",
         "regex": r"thu viec[^.\n]{0,80}?(\d{1,2}) thang", "kieu": "max",
         "gia_tri": 2, "don_vi": "tháng", "can_cu": "Điều 25 Bộ luật Lao động 2019",
         "canh_bao": "Thử việc quá 2 tháng (60 ngày) vượt trần Điều 25 BLLĐ 2019, trừ "
                     "người quản lý doanh nghiệp (180 ngày).",
         "goi_y": "Rút thời gian thử việc về tối đa 60 ngày (hoặc 30/6 ngày tuỳ công "
                  "việc)."},
        {"ma": "luong_thu_viec_pct", "ten": "Tiền lương thử việc (% lương chính thức)",
         "regex": r"thu viec[^.\n]{0,100}?(\d{2,3})\s*%|(\d{2,3})\s*%[^.\n]{0,60}?thu viec",
         "kieu": "min", "gia_tri": 85, "don_vi": "%",
         "can_cu": "Điều 26 Bộ luật Lao động 2019",
         "canh_bao": "Tiền lương thử việc phải ít nhất bằng 85% mức lương của công việc "
                     "đó.",
         "goi_y": "Nâng lương thử việc lên tối thiểu 85% mức lương chính thức."},
        {"ma": "hd_xac_dinh_thang", "ten": "Thời hạn hợp đồng xác định thời hạn",
         "regex": r"thoi han[^.\n]{0,60}?(\d{1,3}) thang", "kieu": "max",
         "gia_tri": 36, "don_vi": "tháng", "can_cu": "Điều 20 Bộ luật Lao động 2019",
         "canh_bao": "Hợp đồng xác định thời hạn không quá 36 tháng kể từ thời điểm có "
                     "hiệu lực.",
         "goi_y": "Rút về ≤ 36 tháng hoặc chuyển thành hợp đồng không xác định thời hạn."},
        {"ma": "gio_ngay", "ten": "Thời giờ làm việc mỗi ngày",
         "regex": r"(\d{1,2})\s*gio\s*(?:/|moi|mot|trong mot|trong)\s*ngay", "kieu": "max",
         "gia_tri": 8, "don_vi": "giờ/ngày", "can_cu": "Điều 105 Bộ luật Lao động 2019",
         "canh_bao": "Thời giờ làm việc bình thường không quá 8 giờ/ngày (kể cả làm thêm "
                     "không quá 12 giờ/ngày — Điều 107). Kiểm tra lại con số này có phải "
                     "giờ làm việc bình thường không.",
         "goi_y": "Đưa thời giờ làm việc bình thường về ≤ 8 giờ/ngày; giờ làm thêm ghi "
                  "riêng."},
        {"ma": "gio_tuan", "ten": "Thời giờ làm việc mỗi tuần",
         "regex": r"(\d{1,2})\s*gio\s*(?:/|moi|mot|trong mot|trong)\s*tuan", "kieu": "max",
         "gia_tri": 48, "don_vi": "giờ/tuần", "can_cu": "Điều 105 Bộ luật Lao động 2019",
         "canh_bao": "Thời giờ làm việc bình thường không quá 48 giờ/tuần.",
         "goi_y": "Đưa về ≤ 48 giờ/tuần (Nhà nước khuyến khích 40 giờ/tuần)."},
        {"ma": "lam_them_thang", "ten": "Giờ làm thêm mỗi tháng",
         "regex": r"lam them[^.\n]{0,80}?(\d{1,3}) gio\s*(?:/|moi|mot|trong mot)\s*thang",
         "kieu": "max", "gia_tri": 40, "don_vi": "giờ/tháng",
         "can_cu": "Điều 107 Bộ luật Lao động 2019",
         "canh_bao": "Làm thêm không quá 40 giờ/tháng.",
         "goi_y": "Đưa giờ làm thêm về ≤ 40 giờ/tháng và phải được người lao động đồng ý."},
        {"ma": "lam_them_nam", "ten": "Giờ làm thêm mỗi năm",
         "regex": r"lam them[^.\n]{0,120}?(\d{1,3}) gio\s*(?:/|moi|mot|trong mot)\s*nam",
         "kieu": "max", "gia_tri": 200, "don_vi": "giờ/năm",
         "can_cu": "Điều 107 Bộ luật Lao động 2019",
         "canh_bao": "Làm thêm không quá 200 giờ/năm (300 giờ/năm chỉ với ngành nghề, "
                     "trường hợp luật cho phép).",
         "goi_y": "Đưa về ≤ 200 giờ/năm, hoặc nêu rõ thuộc trường hợp được làm tới "
                  "300 giờ/năm."},
    ],
}

LOAI_HOP_DONG["hop_dong_dich_vu"] = {
    "ten": "Hợp đồng dịch vụ",
    "nhan_dien": [r"hop dong dich vu", r"cung (cap|ung) dich vu", r"ben (cung cap|cung ung)",
                  r"ben su dung dich vu|ben thue dich vu", r"phi dich vu|gia dich vu",
                  r"dich vu (tu van|phap ly|ke toan|bao ve|ve sinh|van chuyen|thiet ke|"
                  r"quang cao|bao tri)"],
    "dieu_khoan": [
        _dk("doi_tuong", "Đối tượng / phạm vi dịch vụ",
            [r"doi tuong", r"pham vi (cong viec|dich vu)", r"noi dung (cong viec|dich vu)",
             r"dich vu"], True, "Điều 513, 514 Bộ luật Dân sự 2015",
            "Mô tả rõ công việc phải làm, kết quả mong đợi, tiêu chuẩn chất lượng; "
            "dịch vụ phải là công việc có thể thực hiện và không bị cấm."),
        _dk("gia_thanh_toan", "Giá dịch vụ và thanh toán",
            [r"gia (dich vu|tri hop dong)|phi dich vu|gia ca|don gia|tong gia tri",
             r"thanh toan|tra tien"], True, "Điều 519 Bộ luật Dân sự 2015",
            "Ghi giá (có/không gồm VAT), tiến độ và phương thức thanh toán, chứng từ."),
        _dk("thoi_han", "Thời hạn thực hiện",
            [r"thoi han", r"tien do", r"thoi gian thuc hien"], True,
            "Điều 517 Bộ luật Dân sự 2015",
            "Ghi thời hạn hoàn thành từng hạng mục và toàn bộ dịch vụ."),
        _QUYEN_NGHIA_VU,
        _dk("nghiem_thu", "Nghiệm thu / bàn giao kết quả",
            [r"nghiem thu", r"ban giao", r"ket qua (cong viec|dich vu)"], False,
            "Điều 517, 518 Bộ luật Dân sự 2015",
            "Nêu cách nghiệm thu, thời hạn phản hồi, biên bản nghiệm thu là căn cứ "
            "thanh toán."),
        _dk("bao_mat", "Bảo mật thông tin",
            [r"bao mat", r"bi mat"], False, "Điều 517 Bộ luật Dân sự 2015",
            "Thêm nghĩa vụ giữ bí mật thông tin biết được trong quá trình thực hiện "
            "dịch vụ."),
        _CHAM_DUT, _PHAT_BOI_THUONG, _BAT_KHA_KHANG, _TRANH_CHAP,
        _MOT_CHIEU_CHAM_DUT, _DUNG_THONG_TIN_KHONG_CHAP_THUAN, _HIEU_LUC_KHONG_CHU_KY,
    ],
    "nguong": _nguong_phat_va_lai(),
}

LOAI_HOP_DONG["hop_dong_mua_ban"] = {
    "ten": "Hợp đồng mua bán hàng hoá",
    "nhan_dien": [r"hop dong mua ban", r"hop dong (cung cap|cung ung) hang", r"\bben ban\b",
                  r"\bben mua\b", r"hang hoa", r"giao hang", r"don dat hang|\bpo\b"],
    "dieu_khoan": [
        _dk("hang_hoa", "Hàng hoá (tên, số lượng, chất lượng, quy cách)",
            [r"hang hoa|san pham", r"so luong", r"chat luong|quy cach|tieu chuan|thong so"],
            True, "Điều 430, 432 Bộ luật Dân sự 2015; Điều 39 Luật Thương mại 2005",
            "Ghi tên hàng, mã, số lượng, đơn vị, chất lượng/quy cách/tiêu chuẩn áp "
            "dụng, xuất xứ."),
        _dk("gia", "Giá",
            [r"\bgia\b|don gia|gia tri hop dong|tong gia tri|gia ban|gia mua"], True,
            "Điều 433 Bộ luật Dân sự 2015; Điều 52 Luật Thương mại 2005",
            "Ghi đơn giá, tổng giá trị, có/không gồm VAT, chi phí vận chuyển, bốc dỡ."),
        _dk("thanh_toan", "Phương thức thanh toán",
            [r"thanh toan|tra tien"], True,
            "Điều 440 Bộ luật Dân sự 2015; Điều 50, 55 Luật Thương mại 2005",
            "Ghi phương thức, đợt thanh toán, thời hạn, tài khoản, chứng từ."),
        _dk("giao_hang", "Giao hàng (thời gian, địa điểm)",
            [r"giao hang|giao nhan|thoi gian giao|dia diem giao|thoi han giao"], True,
            "Điều 434, 435 Bộ luật Dân sự 2015; Điều 35, 37 Luật Thương mại 2005",
            "Ghi thời hạn, địa điểm giao, bên chịu chi phí vận chuyển, chứng từ giao "
            "nhận."),
        _dk("chuyen_rui_ro", "Chuyển rủi ro và quyền sở hữu",
            [r"chuyen (giao )?rui ro|rui ro", r"quyen so huu"], False,
            "Điều 441 Bộ luật Dân sự 2015; Điều 57–62 Luật Thương mại 2005",
            "Nêu thời điểm chuyển rủi ro (thường khi giao hàng) và thời điểm chuyển "
            "quyền sở hữu (thường khi thanh toán đủ)."),
        _dk("bao_hanh", "Bảo hành",
            [r"bao hanh"], False, "Điều 446–449 Bộ luật Dân sự 2015; Điều 49 Luật Thương mại 2005",
            "Ghi thời hạn, phạm vi bảo hành, cách xử lý hàng lỗi (sửa, đổi, trả)."),
        _PHAT_BOI_THUONG, _BAT_KHA_KHANG, _TRANH_CHAP,
    ],
    "nguong": _nguong_phat_va_lai(),
}

LOAI_HOP_DONG["hop_dong_thue"] = {
    "ten": "Hợp đồng thuê (nhà / mặt bằng / tài sản)",
    "nhan_dien": [r"hop dong (cho )?thue", r"\bben cho thue\b", r"\bben thue\b",
                  r"tien thue|gia thue", r"mat bang", r"thue nha|thue can ho|thue van phong"],
    "dieu_khoan": [
        _dk("doi_tuong_thue", "Đối tượng thuê (địa chỉ, diện tích, hiện trạng)",
            [r"doi tuong (cho )?thue|tai san (cho )?thue|nha (cho )?thue|mat bang|can ho",
             r"dien tich", r"dia chi"], True,
            "Điều 472 Bộ luật Dân sự 2015; Điều 163 Luật Nhà ở 2023",
            "Ghi địa chỉ, diện tích, hiện trạng, giấy tờ pháp lý của tài sản thuê."),
        _dk("gia_thue", "Giá thuê và kỳ thanh toán",
            [r"gia thue|tien thue|gia cho thue", r"ky (han )?thanh toan|thanh toan"], True,
            "Điều 473 Bộ luật Dân sự 2015; Điều 163 Luật Nhà ở 2023",
            "Ghi giá thuê, có/không gồm VAT và phí dịch vụ, kỳ thanh toán, cách điều "
            "chỉnh giá."),
        _dk("thoi_han_thue", "Thời hạn thuê",
            [r"thoi han (thue|cho thue|hop dong)"], True,
            "Điều 474 Bộ luật Dân sự 2015; Điều 163 Luật Nhà ở 2023",
            "Ghi thời hạn thuê, ngày bàn giao, điều kiện gia hạn."),
        _dk("dat_coc", "Đặt cọc / ký quỹ",
            [r"dat coc|tien coc|ky quy"], False, "Điều 328 Bộ luật Dân sự 2015",
            "Ghi số tiền cọc, mục đích, điều kiện hoàn trả / mất cọc / phạt cọc."),
        _dk("muc_dich", "Mục đích sử dụng",
            [r"muc dich (su dung|thue)"], False, "Điều 480 Bộ luật Dân sự 2015",
            "Ghi mục đích thuê (ở, văn phòng, kinh doanh…) — sai mục đích là căn cứ "
            "chấm dứt."),
        _dk("sua_chua", "Sửa chữa, bảo trì, cải tạo",
            [r"sua chua|bao tri|bao duong|cai tao"], False,
            "Điều 477, 479 Bộ luật Dân sự 2015",
            "Phân định bên chịu sửa chữa lớn / nhỏ, điều kiện cải tạo và xử lý khi "
            "trả nhà."),
        _dk("cham_dut_ban_giao", "Chấm dứt và bàn giao lại tài sản",
            [r"cham dut", r"ban giao|tra lai (nha|mat bang|tai san)|hoan tra"], False,
            "Điều 481, 482 Bộ luật Dân sự 2015; Điều 171, 172 Luật Nhà ở 2023",
            "Nêu trường hợp chấm dứt, thời hạn báo trước, tình trạng tài sản khi trả, "
            "xử lý tài sản gắn liền."),
        _QUYEN_NGHIA_VU, _TRANH_CHAP,
    ],
    "nguong": [
        {"ma": "dat_coc_thang", "ten": "Tiền đặt cọc (số tháng tiền thuê)",
         "regex": r"(?:dat coc|tien coc|ky quy)[^.\n]{0,80}?(\d{1,2}) thang", "kieu": "max",
         "gia_tri": 3, "don_vi": "tháng tiền thuê",
         "can_cu": "Thông lệ thị trường (luật không ấn định trần); Điều 328 Bộ luật "
                   "Dân sự 2015",
         "canh_bao": "Đặt cọc trên 3 tháng tiền thuê là bất thường so với thông lệ; rủi "
                     "ro mất cọc lớn cho bên thuê.",
         "goi_y": "Thương lượng về 1–3 tháng tiền thuê, ghi rõ điều kiện hoàn cọc."},
    ],
}

LOAI_HOP_DONG["hop_dong_vay"] = {
    "ten": "Hợp đồng vay tài sản",
    "nhan_dien": [r"hop dong (cho )?vay", r"\bben cho vay\b", r"\bben vay\b", r"khoan vay",
                  r"tien vay|so tien vay", r"lai suat"],
    "dieu_khoan": [
        _dk("so_tien", "Số tiền / tài sản vay",
            [r"so tien (vay|cho vay)|khoan vay|tien vay|tai san vay|so tien"], True,
            "Điều 463 Bộ luật Dân sự 2015",
            "Ghi số tiền bằng số và bằng chữ, đồng tiền, cách giải ngân."),
        _dk("lai_suat", "Lãi suất",
            [r"lai suat|khong tinh lai|khong lai|vay khong lai"], False,
            "Điều 468 Bộ luật Dân sự 2015",
            "Ghi rõ lãi suất %/năm (tối đa 20%/năm) hoặc ghi rõ vay không lãi."),
        _dk("thoi_han", "Thời hạn vay",
            [r"thoi han (vay|cho vay|tra no|hop dong)"], False,
            "Điều 469, 470 Bộ luật Dân sự 2015",
            "Ghi thời hạn vay và ngày đến hạn; vay không kỳ hạn thì nêu cách đòi nợ."),
        _dk("tra_no", "Phương thức trả nợ",
            [r"phuong thuc (tra|thanh toan)|tra no|hoan tra|tra goc|ky tra|lich tra"], False,
            "Điều 466 Bộ luật Dân sự 2015",
            "Ghi lịch trả gốc, lãi; tài khoản nhận; xử lý khi chậm trả."),
        _dk("bao_dam", "Biện pháp bảo đảm",
            [r"bao dam|the chap|cam co|bao lanh|tin chap"], False,
            "Điều 292 Bộ luật Dân sự 2015",
            "Nếu có tài sản bảo đảm thì ghi rõ và lập hợp đồng thế chấp/cầm cố, đăng "
            "ký biện pháp bảo đảm khi luật yêu cầu."),
        _dk("tra_truoc_han", "Trả nợ trước hạn",
            [r"tra (no )?truoc (thoi )?han|thanh toan truoc han|tat toan truoc"], False,
            "Điều 470 Bộ luật Dân sự 2015",
            "Nêu bên vay được trả trước hạn hay không và có phải trả lãi cho phần thời "
            "hạn còn lại không."),
        _TRANH_CHAP,
    ],
    "nguong": [
        {"ma": "lai_suat_nam", "ten": "Lãi suất vay (theo năm)",
         "regex": r"lai suat[^.\n]{0,60}?(\d{1,3}(?:[.,]\d+)?)\s*%\s*"
                  r"(?:/|moi|mot|tren|trong mot)?\s*nam",
         "kieu": "max", "gia_tri": 20, "don_vi": "%/năm",
         "can_cu": "Điều 468 Bộ luật Dân sự 2015",
         "canh_bao": "Lãi suất thoả thuận không được vượt 20%/năm; phần vượt vô hiệu. "
                     "Lãi gấp 5 lần mức trần có thể cấu thành tội cho vay lãi nặng "
                     "(Điều 201 Bộ luật Hình sự 2015).",
         "goi_y": "Đưa lãi suất về ≤ 20%/năm."},
        {"ma": "lai_suat_thang", "ten": "Lãi suất vay (theo tháng)",
         "regex": r"lai suat[^.\n]{0,60}?(\d{1,2}(?:[.,]\d+)?)\s*%\s*"
                  r"(?:/|moi|mot|tren|trong mot)\s*thang",
         "kieu": "max", "gia_tri": 1.67, "don_vi": "%/tháng",
         "can_cu": "Điều 468 Bộ luật Dân sự 2015",
         "canh_bao": "Quy đổi theo năm vượt trần 20%/năm (≈1,67%/tháng) của Điều 468 "
                     "BLDS 2015.",
         "goi_y": "Đưa lãi suất về ≤ 1,67%/tháng (20%/năm)."},
    ],
}

LOAI_HOP_DONG["hop_dong_hop_tac"] = {
    "ten": "Hợp đồng hợp tác kinh doanh (BCC)",
    "nhan_dien": [r"hop dong hop tac", r"hop tac kinh doanh", r"\bbcc\b",
                  r"phan chia loi nhuan", r"gop von hop tac|dong gop", r"ben hop tac"],
    "dieu_khoan": [
        _dk("muc_dich", "Mục đích hợp tác",
            [r"muc dich|muc tieu", r"noi dung hop tac|pham vi hop tac"], True,
            "Điều 504, 505 Bộ luật Dân sự 2015; Điều 27, 28 Luật Đầu tư 2020",
            "Ghi rõ dự án/hoạt động hợp tác, phạm vi, địa bàn."),
        _dk("dong_gop", "Đóng góp của các bên",
            [r"dong gop|gop von|phan gop|tai san gop|von gop"], True,
            "Điều 505, 506 Bộ luật Dân sự 2015",
            "Ghi từng bên góp gì (tiền, tài sản, công sức), giá trị, thời hạn góp, "
            "xử lý khi góp chậm."),
        _dk("phan_chia", "Phân chia lợi nhuận / lỗ",
            [r"phan chia (loi nhuan|ket qua|lai)|chia loi nhuan|phan phoi loi nhuan",
             r"(chiu|ganh|phan chia) (lo|rui ro)|phan chia loi nhuan va lo"], True,
            "Điều 505, 509 Bộ luật Dân sự 2015; Điều 28 Luật Đầu tư 2020",
            "Ghi tỷ lệ, thời điểm, cách xác định lợi nhuận và phân chia lỗ."),
        _dk("quan_ly", "Quản lý, đại diện các bên",
            [r"quan ly|dieu hanh|dai dien|ban dieu phoi|ban dieu hanh"], False,
            "Điều 508 Bộ luật Dân sự 2015",
            "Nêu ai đại diện các bên trong giao dịch với bên thứ ba, cơ chế biểu quyết."),
        _dk("thoi_han", "Thời hạn hợp tác",
            [r"thoi han"], False, "Điều 505 Bộ luật Dân sự 2015; Điều 28 Luật Đầu tư 2020",
            "Ghi thời hạn hợp tác và điều kiện gia hạn."),
        _dk("cham_dut_thanh_ly", "Chấm dứt, rút khỏi, thanh lý hợp tác",
            [r"cham dut|thanh ly|giai the|rut khoi"], False,
            "Điều 510–512 Bộ luật Dân sự 2015",
            "Nêu điều kiện rút khỏi, chấm dứt, cách thanh lý tài sản chung."),
        _TRANH_CHAP,
    ],
    "nguong": [],
}

LOAI_HOP_DONG["hop_dong_chuyen_nhuong_von"] = {
    "ten": "Hợp đồng chuyển nhượng phần vốn góp / cổ phần",
    "nhan_dien": [r"chuyen nhuong (phan )?von( gop)?", r"chuyen nhuong co phan",
                  r"\bben chuyen nhuong\b", r"\bben nhan chuyen nhuong\b",
                  r"phan von gop", r"so co phan chuyen nhuong"],
    "dieu_khoan": [
        _dk("cac_ben_cn", "Bên chuyển nhượng và bên nhận chuyển nhượng",
            [r"ben chuyen nhuong", r"ben nhan chuyen nhuong"], True,
            "Điều 52, 127 Luật Doanh nghiệp 2020",
            "Ghi đủ thông tin hai bên và công ty có vốn được chuyển nhượng."),
        _dk("phan_von", "Phần vốn góp / cổ phần và tỷ lệ",
            [r"phan von gop|co phan|ty le|so co phan|gia tri von gop|menh gia"], True,
            "Điều 52, 127 Luật Doanh nghiệp 2020",
            "Ghi giá trị vốn góp/số cổ phần, loại cổ phần, tỷ lệ trên vốn điều lệ."),
        _dk("gia_thanh_toan", "Giá chuyển nhượng và thanh toán",
            [r"gia chuyen nhuong|gia tri chuyen nhuong", r"thanh toan"], True,
            "Điều 430 Bộ luật Dân sự 2015",
            "Ghi giá, đợt thanh toán, điều kiện thanh toán gắn với hoàn tất thủ tục."),
        _dk("dieu_kien", "Điều kiện chuyển nhượng (quyền ưu tiên, cổ đông sáng lập)",
            [r"quyen uu tien|chao ban|thanh vien con lai|co dong sang lap|"
             r"dieu kien chuyen nhuong|chap thuan"], False,
            "Điều 52, 120, 127 Luật Doanh nghiệp 2020",
            "Xác nhận đã chào bán cho thành viên còn lại (TNHH 2TV) hoặc được ĐHĐCĐ "
            "chấp thuận (cổ phần phổ thông của cổ đông sáng lập trong 3 năm đầu)."),
        _dk("thue_chi_phi", "Thuế và chi phí",
            [r"thue (thu nhap|chuyen nhuong)|le phi|chi phi|nghia vu thue"], False,
            "Luật Thuế thu nhập cá nhân; Luật Thuế thu nhập doanh nghiệp",
            "Ghi bên chịu thuế thu nhập từ chuyển nhượng vốn và lệ phí đăng ký."),
        _dk("thu_tuc", "Thủ tục đăng ký thay đổi",
            [r"dang ky thay doi|thong bao thay doi|so dang ky (thanh vien|co dong)|"
             r"giay chung nhan (dang ky doanh nghiep|phan von gop|co phan)|"
             r"co quan dang ky kinh doanh"], False,
            "Điều 52, 127 Luật Doanh nghiệp 2020; Nghị định 01/2021/NĐ-CP",
            "Nêu bên chịu trách nhiệm nộp hồ sơ thay đổi thành viên/cổ đông và thời "
            "hạn; quyền của bên nhận chỉ phát sinh khi ghi vào sổ."),
        _TRANH_CHAP,
    ],
    "nguong": [],
}

LOAI_HOP_DONG["nda"] = {
    "ten": "Thoả thuận bảo mật thông tin (NDA)",
    "nhan_dien": [r"thoa thuan bao mat", r"hop dong bao mat", r"\bnda\b", r"non.?disclosure",
                  r"thong tin (bao )?mat|thong tin bi mat", r"ben tiet lo|ben nhan thong tin",
                  r"cam ket bao mat"],
    "dieu_khoan": [
        _dk("dinh_nghia", "Định nghĩa thông tin mật",
            [r"thong tin (bao )?mat|thong tin bi mat", r"dinh nghia", r"bi mat kinh doanh"],
            True, "Điều 84 Luật Sở hữu trí tuệ; Điều 387 Bộ luật Dân sự 2015",
            "Định nghĩa rõ phạm vi thông tin mật (hình thức, cách đánh dấu, thông tin "
            "nói miệng)."),
        _dk("nghia_vu_bao_mat", "Nghĩa vụ bảo mật",
            [r"nghia vu bao mat|cam ket bao mat|khong (duoc )?tiet lo|bao mat"], True,
            "Điều 387 Bộ luật Dân sự 2015; Điều 21 Bộ luật Lao động 2019",
            "Nêu nghĩa vụ không tiết lộ, không sử dụng ngoài mục đích, giới hạn người "
            "được tiếp cận."),
        _dk("ngoai_le", "Ngoại lệ bảo mật",
            [r"ngoai le|khong (thuoc|bao gom|ap dung)|da (duoc )?cong khai|da biet|"
             r"theo yeu cau cua co quan|co quan nha nuoc co tham quyen"], False,
            "Thông lệ soạn thảo NDA",
            "Liệt kê ngoại lệ: thông tin đã công khai, đã biết trước, do bên thứ ba "
            "cung cấp hợp pháp, phải cung cấp theo yêu cầu cơ quan có thẩm quyền."),
        _dk("thoi_han", "Thời hạn bảo mật",
            [r"thoi han (bao mat|hieu luc|thoa thuan)|trong (thoi han|vong)[^\n]{0,20}nam|"
             r"ke ca sau khi|sau khi (cham dut|ket thuc)"], False,
            "Thông lệ soạn thảo NDA",
            "Ghi thời hạn bảo mật (kể cả sau khi chấm dứt quan hệ), vô thời hạn với "
            "bí mật kinh doanh."),
        _dk("che_tai", "Chế tài khi vi phạm",
            [r"che tai|phat|boi thuong|vi pham"], False,
            "Điều 418 Bộ luật Dân sự 2015",
            "Nêu phạt vi phạm, bồi thường thiệt hại, quyền yêu cầu biện pháp khẩn cấp."),
        _dk("hoan_tra", "Hoàn trả / tiêu huỷ thông tin",
            [r"hoan tra|tra lai|tieu huy|xoa bo"], False, "Thông lệ soạn thảo NDA",
            "Nêu nghĩa vụ hoàn trả hoặc tiêu huỷ tài liệu khi kết thúc, kèm xác nhận "
            "bằng văn bản."),
        _TRANH_CHAP,
    ],
    "nguong": [],
}

LOAI_HOP_DONG["hop_dong_nhuong_quyen"] = {
    "ten": "Hợp đồng nhượng quyền thương mại",
    "nhan_dien": [r"nhuong quyen thuong mai", r"hop dong nhuong quyen", r"\bben nhuong quyen\b",
                  r"ben nhan quyen(?! su dung)", r"franchis", r"phi nhuong quyen",
                  r"quyen thuong mai"],
    "dieu_khoan": [
        _dk("quyen_thuong_mai", "Nội dung quyền thương mại",
            [r"quyen thuong mai|noi dung nhuong quyen|he thong (nhuong quyen|kinh doanh)",
             r"nhan hieu|bi quyet|mo hinh kinh doanh|khau hieu|bieu tuong"], True,
            "Điều 284 Luật Thương mại 2005; Điều 11 Nghị định 35/2006/NĐ-CP",
            "Mô tả quyền được cấp: nhãn hiệu, tên thương mại, bí quyết, cách bố trí, "
            "phạm vi lãnh thổ, độc quyền hay không."),
        _dk("quyen_nghia_vu_nhuong", "Quyền và nghĩa vụ bên nhượng quyền",
            [r"ben nhuong quyen"], True,
            "Điều 286, 287 Luật Thương mại 2005",
            "Nêu nghĩa vụ cung cấp tài liệu, đào tạo, trợ giúp kỹ thuật, bảo đảm quyền "
            "SHTT; quyền kiểm soát, nhận phí."),
        _dk("quyen_nghia_vu_nhan", "Quyền và nghĩa vụ bên nhận quyền",
            [r"ben nhan quyen|ben nhan nhuong quyen"], True,
            "Điều 288, 289 Luật Thương mại 2005",
            "Nêu nghĩa vụ trả phí, tuân thủ hệ thống, giữ bí mật, ngừng dùng nhãn "
            "hiệu khi chấm dứt."),
        _dk("phi", "Giá / phí nhượng quyền",
            [r"phi nhuong quyen|phi (ban dau|dinh ky|thuong hieu|quan ly|gia nhap)|"
             r"gia nhuong quyen|royalt|phi ban quyen"], True,
            "Điều 11 Nghị định 35/2006/NĐ-CP",
            "Ghi phí ban đầu, phí định kỳ (tỷ lệ doanh thu), phí đào tạo/quảng cáo, "
            "kỳ thanh toán."),
        _dk("thoi_han", "Thời hạn hợp đồng",
            [r"thoi han"], True, "Điều 11 Nghị định 35/2006/NĐ-CP",
            "Ghi thời hạn, điều kiện gia hạn."),
        _dk("gia_han_cham_dut", "Gia hạn, chấm dứt hợp đồng",
            [r"gia han|cham dut|tai ky"], False,
            "Điều 11 Nghị định 35/2006/NĐ-CP; Điều 16 Nghị định 35/2006/NĐ-CP",
            "Nêu điều kiện gia hạn, trường hợp đơn phương chấm dứt và nghĩa vụ hậu "
            "chấm dứt (ngừng dùng nhãn hiệu, hoàn trả tài liệu)."),
        _dk("dang_ky", "Đăng ký hoạt động nhượng quyền",
            [r"dang ky (hoat dong )?nhuong quyen|bo cong thuong|so cong thuong"], False,
            "Điều 291 Luật Thương mại 2005; Điều 17 Nghị định 35/2006/NĐ-CP",
            "Xác nhận bên nhượng quyền đã đăng ký hoạt động nhượng quyền với Bộ Công "
            "Thương (nhượng quyền từ nước ngoài vào Việt Nam) trước khi ký."),
        _TRANH_CHAP,
    ],
    "nguong": [],
}

LOAI_HOP_DONG["hop_dong_li_xang"] = {
    "ten": "Hợp đồng chuyển quyền sử dụng đối tượng sở hữu công nghiệp (li-xăng)",
    "nhan_dien": [r"li.?xang", r"licen[cs]e", r"chuyen quyen su dung",
                  r"hop dong su dung (nhan hieu|sang che|kieu dang|doi tuong)",
                  r"\bben chuyen quyen\b", r"\bben nhan quyen su dung\b|ben duoc chuyen quyen",
                  r"so huu (cong nghiep|tri tue)"],
    "dieu_khoan": [
        _dk("doi_tuong", "Đối tượng chuyển giao",
            [r"doi tuong (chuyen giao|chuyen quyen|li.?xang|licen|hop dong)",
             r"nhan hieu|sang che|kieu dang cong nghiep|giai phap huu ich|thiet ke bo tri|"
             r"bi quyet|van bang bao ho|giay chung nhan dang ky|bang doc quyen"], True,
            "Điều 141, 144 Luật Sở hữu trí tuệ",
            "Ghi tên đối tượng, số văn bằng bảo hộ, chủ sở hữu, thời hạn bảo hộ còn lại."),
        _dk("pham_vi", "Phạm vi chuyển quyền (dạng, lãnh thổ)",
            [r"pham vi|lanh tho|doc quyen|khong doc quyen|thu cap"], True,
            "Điều 143, 144 Luật Sở hữu trí tuệ",
            "Ghi dạng hợp đồng (độc quyền / không độc quyền / thứ cấp), lãnh thổ, "
            "lĩnh vực sử dụng."),
        _dk("thoi_han", "Thời hạn hợp đồng",
            [r"thoi han"], True, "Điều 144 Luật Sở hữu trí tuệ",
            "Ghi thời hạn (không vượt thời hạn bảo hộ còn lại của văn bằng)."),
        _dk("gia", "Giá chuyển giao / phí và thanh toán",
            [r"gia (chuyen giao|chuyen quyen|li.?xang|licen)|phi (li.?xang|licen|su dung|"
             r"ban quyen)|thanh toan|royalt|gia ca"], True,
            "Điều 144 Luật Sở hữu trí tuệ",
            "Ghi giá/phí (trọn gói hay theo doanh thu), kỳ thanh toán, thuế."),
        _QUYEN_NGHIA_VU,
        _dk("han_che_bat_hop_ly", "Điều khoản hạn chế bất hợp lý quyền của bên nhận",
            [r"(cam|khong duoc)[^\n.]{0,40}cai tien",
             r"chuyen giao (mien phi|khong boi hoan)[^\n.]{0,40}cai tien",
             r"(han che|cam|khong duoc)[^\n.]{0,40}xuat khau",
             r"(phai|buoc|bat buoc|chi duoc) mua[^\n.]{0,80}(nguyen lieu|vat lieu|thiet bi|"
             r"linh kien)",
             r"(cam|khong duoc)[^\n.]{0,40}(khieu kien|khieu nai|phan doi)[^\n.]{0,40}"
             r"(hieu luc|van bang|quyen so huu)"], False,
            "Khoản 2 Điều 144 Luật Sở hữu trí tuệ",
            "Bỏ các điều khoản: cấm bên nhận cải tiến; buộc chuyển giao miễn phí cải "
            "tiến; hạn chế xuất khẩu bất hợp lý; buộc mua nguyên liệu/thiết bị từ bên "
            "giao; cấm khiếu kiện hiệu lực quyền — các điều khoản này mặc nhiên vô hiệu.",
            nguoc=True),
        _TRANH_CHAP,
    ],
    "nguong": [],
}

# Không nhận diện được → chỉ kiểm 4 mục chung.
LOAI_HOP_DONG["khac"] = {"ten": "Tài liệu khác / chưa nhận diện", "nhan_dien": [],
                         "dieu_khoan": [], "nguong": []}


# ---------------------------------------------------------------------------
# Nhận diện loại hợp đồng
# ---------------------------------------------------------------------------
# Tiêu đề (tên file / dòng người dùng gõ) nặng nhất; vài dòng đầu văn bản
# (quốc hiệu, tên hợp đồng) nặng vừa; 3000 ký tự đầu nhẹ nhất. Điểm 5 = chắc.
_TRONG_SO_TIEU_DE = 3
_TRONG_SO_DAU = 2
_TRONG_SO_THAN = 1
_DIEM_CHAC = 5


def _dau_van_ban(text: str, so_dong: int = 5, max_chars: int = 400) -> str:
    dong = [d.strip() for d in (text or "").splitlines() if d.strip()]
    return "\n".join(dong[:so_dong])[:max_chars]


def nhan_dien_loai(tieu_de: str, text: str) -> tuple[str, float]:
    """(mã loại, độ tin cậy 0..1). Không khớp regex nào → ("khac", 0.0).
    Hai loại đồng điểm → giữ loại đứng trước trong danh mục nhưng hạ tin cậy
    một nửa, để lớp ngoài biết mà hỏi lại người dùng."""
    td = bo_dau(tieu_de or "")
    dau = bo_dau(_dau_van_ban(text))
    than = bo_dau((text or "")[:NHAN_DIEN_DAU])
    diem: dict[str, int] = {}
    for ma, cfg in LOAI_HOP_DONG.items():
        d = 0
        for mau in cfg.get("nhan_dien") or []:
            rx = _bien_dich(mau)
            if td and rx.search(td):
                d += _TRONG_SO_TIEU_DE
            if dau and rx.search(dau):
                d += _TRONG_SO_DAU
            if than and rx.search(than):
                d += _TRONG_SO_THAN
        if d:
            diem[ma] = d
    if not diem:
        return "khac", 0.0
    thu_tu = list(LOAI_HOP_DONG)
    xep = sorted(diem.items(), key=lambda kv: (-kv[1], thu_tu.index(kv[0])))
    ma, cao_nhat = xep[0]
    tin = min(1.0, cao_nhat / _DIEM_CHAC)
    if len(xep) > 1 and xep[1][1] == cao_nhat:
        tin *= 0.5
    return ma, round(tin, 2)


# ---------------------------------------------------------------------------
# Tách điều khoản
# ---------------------------------------------------------------------------
# Chạy trên DÒNG đã bỏ dấu (cùng độ dài dòng gốc) nên "Điều"/"ĐIỀU"/"Dieu"
# đều bắt được. Lookahead âm loại câu văn bắt đầu bằng "Điều 3 của hợp đồng
# này…" — đó là trích dẫn, không phải tiêu đề.
_RX_DIEU = re.compile(
    r"^[ \t]*dieu\s*(\d{1,3})(?![.,]\d)\s*[.:\-–—)]?[ \t]*"
    r"(?!(?:cua|nay|neu|tren|va|la|duoc|khoan|noi)\b)(.*)$")
_RX_SO = re.compile(r"^[ \t]*(\d{1,2})[.)](?!\d)[ \t]+(\S.*)$")


def _lam_gon_tieu_de(s: str) -> str:
    s = " ".join((s or "").split()).strip(" .:;-–—\t")
    if len(s) > 120:
        cat = re.search(r"[.:;]", s)
        s = s[:cat.start()] if cat and cat.start() < 120 else s[:120]
    return s.strip(" .:;-–—")


def _tach(text: str):
    """(vị trí bắt đầu mục đầu tiên, danh sách mục). Mỗi mục có thêm
    'bat_dau' (offset trong text) và 'kieu' ('dieu' | 'so' | 'toan_van') để
    ra_soat gắn nhãn 'Điều 5. …' / 'Mục 5. …' và định vị chỗ khớp ngưỡng."""
    text = text or ""
    dong = text.splitlines(keepends=True)
    dau_dong = []
    pos = 0
    for d in dong:
        dau_dong.append((pos, d))
        pos += len(d)

    def _quet(rx):
        ket = []
        for i, (bat_dau, d) in enumerate(dau_dong):
            noi_dung_dong = d.rstrip("\r\n")
            f = bo_dau(noi_dung_dong)
            m = rx.match(f)
            if not m:
                continue
            goc = noi_dung_dong if len(f) == len(noi_dung_dong) else f
            tieu_de = _lam_gon_tieu_de(goc[m.start(2):m.end(2)])
            # "ĐIỀU 3 -" rồi xuống dòng mới ghi tên: lấy dòng kế tiếp nếu ngắn.
            if not tieu_de and i + 1 < len(dau_dong):
                ke = dau_dong[i + 1][1].strip()
                if 0 < len(ke) <= 100:
                    tieu_de = _lam_gon_tieu_de(ke)
            ket.append({"so": m.group(1), "tieu_de": tieu_de, "bat_dau": bat_dau})
        return ket

    muc = _quet(_RX_DIEU)
    kieu = "dieu"
    if not muc:
        muc = _quet(_RX_SO)
        kieu = "so"
    if not muc:
        return 0, [{"so": "", "tieu_de": "", "noi_dung": text, "bat_dau": 0,
                    "kieu": "toan_van"}]
    for i, m in enumerate(muc):
        ket_thuc = muc[i + 1]["bat_dau"] if i + 1 < len(muc) else len(text)
        m["noi_dung"] = text[m["bat_dau"]:ket_thuc].rstrip()
        m["kieu"] = kieu
    return muc[0]["bat_dau"], muc


def tach_dieu_khoan(text: str) -> list[dict]:
    """[{"so": "5", "tieu_de": "Thanh toán", "noi_dung": "..."}] — cắt theo
    dòng 'Điều 5.' / 'ĐIỀU 5:' / 'Điều 5 -'; không có 'Điều' thì theo '5.' /
    '5)' đầu dòng; không có gì → một mục duy nhất. Phần mở đầu (trước Điều 1:
    tên các bên, ngày ký) không nằm trong danh sách — ra_soat vẫn quét nó."""
    return _tach(text)[1]


def _nhan_muc(m: dict) -> str:
    if m.get("kieu") == "toan_van":
        return "Toàn văn"
    tien_to = "Điều" if m.get("kieu") == "dieu" else "Mục"
    return f"{tien_to} {m['so']}. {m['tieu_de']}".rstrip(". ")


# ---------------------------------------------------------------------------
# Rà soát
# ---------------------------------------------------------------------------
NHAN_TRANG_THAI = {"dat": "ĐẠT", "canh_bao": "CẢNH BÁO", "thieu": "THIẾU"}
NHAN_RUI_RO = {"thap": "THẤP", "trung_binh": "TRUNG BÌNH", "cao": "CAO"}
_THU_TU = {"thieu": 0, "canh_bao": 1, "dat": 2}


def _trich(goc: str, fold: str, m: re.Match) -> str:
    """≤ TRICH_MAX ký tự quanh chỗ khớp, lấy từ bản gốc có dấu khi độ dài hai
    bản khớp nhau (bo_dau giữ độ dài), gộp khoảng trắng để vào một ô bảng."""
    nguon = goc if len(goc) == len(fold) else fold
    s = max(0, m.start() - 70)
    e = min(len(nguon), m.end() + 130)
    doan = " ".join(nguon[s:e].split())
    if s > 0:
        doan = "…" + doan
    if len(doan) > TRICH_MAX:
        doan = doan[:TRICH_MAX - 1] + "…"
    return doan


def _vung(text: str, mo_dau_toi: int, muc: list[dict]):
    """Các vùng quét theo thứ tự văn bản: (nhãn, offset, gốc, bản bỏ dấu)."""
    vung = []
    if muc and muc[0].get("kieu") != "toan_van" and text[:mo_dau_toi].strip():
        vung.append(("Phần mở đầu", 0, text[:mo_dau_toi]))
    for m in muc:
        vung.append((_nhan_muc(m), m["bat_dau"], m["noi_dung"]))
    return [(nhan, off, goc, bo_dau(goc)) for nhan, off, goc in vung]


def _tim_dieu_khoan(tu_khoa: list[str], vung) -> tuple[str, str] | None:
    """(nhãn mục, đoạn trích) của mục ĐẦU TIÊN theo thứ tự văn bản khớp bất
    kỳ regex nào; None nếu không mục nào khớp."""
    rxs = [_bien_dich(p) for p in tu_khoa]
    for nhan, _off, goc, fold in vung:
        for rx in rxs:
            m = rx.search(fold)
            if m:
                return nhan, _trich(goc, fold, m)
    return None


def _nhan_tai(vung, offset: int) -> str | None:
    nhan = None
    for ten, off, _goc, _fold in vung:
        if off <= offset:
            nhan = ten
    return nhan


def _fmt_so(v: float) -> str:
    return str(int(v)) if float(v).is_integer() else f"{v:g}".replace(".", ",")


def _kiem_nguong(ng: dict, text: str, fold: str, vung) -> dict | None:
    """Quét ngưỡng trên toàn văn bản. Nhiều con số thì lấy con số XẤU NHẤT
    (vi phạm trước, rồi xa ngưỡng nhất); không có số → None (không báo thiếu,
    vì hợp đồng không nêu là chuyện khác với hợp đồng nêu sai)."""
    rx = _bien_dich(ng["regex"])
    kieu, tran = ng["kieu"], float(ng["gia_tri"])
    ung_vien = []
    for m in rx.finditer(fold):
        raw = next((g for g in m.groups() if g), None)
        v = _so(raw)
        if v is None:
            continue
        vi_pham = (kieu == "max" and v > tran) or (kieu == "min" and v < tran)
        ung_vien.append((vi_pham, v if kieu == "max" else -v, v, m))
    if not ung_vien:
        return None
    ung_vien.sort(key=lambda t: (not t[0], -t[1]))
    vi_pham, _k, v, m = ung_vien[0]
    don_vi = ng.get("don_vi", "")
    gioi_han = "tối đa" if kieu == "max" else "tối thiểu"
    if vi_pham:
        so_sanh = "vượt trần" if kieu == "max" else "thấp hơn mức tối thiểu"
        giai_thich = (f"{ng['ten']} = {_fmt_so(v)} {don_vi}, {so_sanh} "
                      f"{_fmt_so(tran)} {don_vi}. {ng['canh_bao']}")
        trang_thai, de_xuat = "canh_bao", ng["goi_y"]
    else:
        giai_thich = (f"{ng['ten']} = {_fmt_so(v)} {don_vi}, trong giới hạn {gioi_han} "
                      f"{_fmt_so(tran)} {don_vi}.")
        trang_thai, de_xuat = "dat", ""
    return {"ma": ng["ma"], "ten": ng["ten"], "trang_thai": trang_thai,
            "dieu_khoan": _nhan_tai(vung, m.start()), "trich": _trich(text, fold, m),
            "gia_tri": v, "giai_thich": giai_thich, "de_xuat": de_xuat,
            "can_cu": ng["can_cu"]}


def _muc_dieu_khoan(dk: dict, hit) -> dict:
    ten = dk["ten"]
    if dk.get("nguoc"):
        # Điều khoản mà CÓ mới là rủi ro (hạn chế bất hợp lý — Điều 144 SHTT,
        # hoặc điều khoản một chiều). Lời giải thích riêng nếu quy tắc có.
        if hit:
            giai_thich = dk.get("giai_thich_co") or (
                f"Phát hiện nội dung có dấu hiệu '{ten.lower()}' tại {hit[0]} — loại "
                f"điều khoản này mặc nhiên vô hiệu theo {dk['can_cu']}.")
            return {"ma": dk["ma"], "ten": ten, "trang_thai": "canh_bao",
                    "dieu_khoan": hit[0], "trich": hit[1],
                    "giai_thich": giai_thich,
                    "de_xuat": dk["goi_y"], "can_cu": dk["can_cu"]}
        return {"ma": dk["ma"], "ten": ten, "trang_thai": "dat", "dieu_khoan": None,
                "trich": None,
                "giai_thich": f"Không thấy nội dung thuộc nhóm '{ten.lower()}'.",
                "de_xuat": "", "can_cu": dk["can_cu"]}
    if hit:
        # CHỈ KIỂM SỰ CÓ MẶT. Nói thẳng ra — kiểm thử 18/09/2026: "Bảo mật —
        # ĐẠT" cho một điều khoản bảo mật cho phép bên kia dùng thông tin
        # không cần chấp thuận, người soát tưởng máy đã đọc và duyệt nội dung.
        return {"ma": dk["ma"], "ten": ten, "trang_thai": "dat", "dieu_khoan": hit[0],
                "trich": hit[1], "chi_co_mat": True,
                "giai_thich": f"Có điều khoản về {ten.lower()} tại {hit[0]} — máy mới "
                              "kiểm là CÓ, chưa đánh giá nội dung có lợi hay bất lợi; "
                              "đọc phần trích và phần rà soát bằng AI để kết luận.",
                "de_xuat": "", "can_cu": dk["can_cu"]}
    if dk["bat_buoc"]:
        return {"ma": dk["ma"], "ten": ten, "trang_thai": "thieu", "dieu_khoan": None,
                "trich": None,
                "giai_thich": f"Không tìm thấy điều khoản về {ten.lower()} — đây là nội "
                              f"dung BẮT BUỘC theo {dk['can_cu']}.",
                "de_xuat": dk["goi_y"], "can_cu": dk["can_cu"]}
    return {"ma": dk["ma"], "ten": ten, "trang_thai": "canh_bao", "dieu_khoan": None,
            "trich": None,
            "giai_thich": f"Không tìm thấy điều khoản về {ten.lower()} — không bắt buộc "
                          f"nhưng thiếu sẽ khó bảo vệ quyền lợi khi có tranh chấp.",
            "de_xuat": dk["goi_y"], "can_cu": dk["can_cu"]}


def _muc_chung(chung: dict, hit) -> dict:
    if hit:
        return {"ma": chung["ma"], "ten": chung["ten"], "trang_thai": "dat",
                "dieu_khoan": hit[0], "trich": hit[1], "chi_co_mat": True,
                "giai_thich": f"Có {chung['ten'].lower()} tại {hit[0]} — máy mới kiểm "
                              "là CÓ, chưa đánh giá nội dung.",
                "de_xuat": "", "can_cu": chung["can_cu"]}
    return {"ma": chung["ma"], "ten": chung["ten"], "trang_thai": "canh_bao",
            "dieu_khoan": None, "trich": None, "giai_thich": chung["giai_thich"],
            "de_xuat": chung["goi_y"], "can_cu": chung["can_cu"]}


def _muc_rui_ro(dat: int, canh_bao: int, thieu: int) -> str:
    if thieu >= 1 or canh_bao >= 3:
        return "cao"
    if canh_bao >= 1:
        return "trung_binh"
    return "thap"


def ra_soat(text: str, tieu_de: str = "", loai: str | None = None) -> dict:
    """Đối chiếu văn bản với danh mục điều khoản chuẩn của loại hợp đồng.
    loai=None (hoặc mã lạ) → tự nhận diện. Kết quả tất định, không gọi model."""
    text = text or ""
    if loai and loai in LOAI_HOP_DONG:
        ma, tin = loai, 1.0
    else:
        ma, tin = nhan_dien_loai(tieu_de, text)
    cfg = LOAI_HOP_DONG[ma]
    mo_dau_toi, ds_muc = _tach(text)
    vung = _vung(text, mo_dau_toi, ds_muc)
    fold = bo_dau(text)

    muc = []
    for dk in cfg["dieu_khoan"]:
        muc.append(_muc_dieu_khoan(dk, _tim_dieu_khoan(dk["tu_khoa"], vung)))
    for ng in cfg["nguong"]:
        kq = _kiem_nguong(ng, text, fold, vung)
        if kq:
            muc.append(kq)
    # Mục chung: bỏ qua nếu danh mục loại đã có mục cùng mã (tranh_chap) để
    # bảng không lặp hai dòng cho một chuyện.
    da_co = {m["ma"] for m in muc}
    for chung in MUC_CHUNG:
        if chung["ma"] in da_co:
            continue
        muc.append(_muc_chung(chung, _tim_dieu_khoan(chung["tu_khoa"], vung)))

    muc.sort(key=lambda m: _THU_TU[m["trang_thai"]])
    dem = {k: sum(1 for m in muc if m["trang_thai"] == k) for k in _THU_TU}
    so_dieu_khoan = 0 if (ds_muc and ds_muc[0].get("kieu") == "toan_van" and not text.strip()) \
        else len(ds_muc)
    return {"loai": ma, "ten_loai": cfg["ten"], "do_tin_cay": tin,
            "so_dieu_khoan": so_dieu_khoan, "muc": muc,
            "tong_ket": {"dat": dem["dat"], "canh_bao": dem["canh_bao"],
                         "thieu": dem["thieu"],
                         "muc_rui_ro": _muc_rui_ro(dem["dat"], dem["canh_bao"],
                                                   dem["thieu"])}}


def cau_hoi_tra_luat(muc: dict) -> str:
    """Câu hỏi ngắn để lớp ngoài tra kho luật cho một mục — ghép tên mục với
    căn cứ để retrieval ghim đúng điều luật thay vì lan man."""
    ten = " ".join((muc.get("ten") or "").split()).rstrip(".")
    can_cu = " ".join((muc.get("can_cu") or "").split())
    if can_cu:
        return f"{can_cu} quy định thế nào về {ten.lower()}?"
    return f"Pháp luật Việt Nam quy định thế nào về {ten.lower()} trong hợp đồng?"


# ---------------------------------------------------------------------------
# Xuất báo cáo .docx
# ---------------------------------------------------------------------------
_FONT = "Times New Roman"


def _dat_font(run, size=12, bold=None, italic=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = _FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), _FONT)


def _ngay_chuoi(ngay) -> str:
    if ngay is None:
        ngay = date.today()
    if hasattr(ngay, "strftime"):
        return ngay.strftime("%d/%m/%Y")
    return str(ngay)


def xuat_bao_cao_docx(ket_qua: dict, tieu_de: str, nguoi_lap: str = "", ngay=None) -> bytes:
    """Bản .docx cho luật sư sửa tay trước khi gửi khách: thông tin tài liệu,
    tóm tắt, bảng từng mục, dòng lưu ý trách nhiệm. Toàn bộ Times New Roman."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), _FONT)

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _dat_font(p.add_run("BÁO CÁO RÀ SOÁT RỦI RO"), size=14, bold=True)

    tk = ket_qua.get("tong_ket") or {}
    tin = ket_qua.get("do_tin_cay") or 0.0
    dong = [
        f"Tài liệu: {tieu_de or '—'}",
        f"Loại hợp đồng: {ket_qua.get('ten_loai', '—')} "
        f"(độ tin cậy nhận diện: {int(round(tin * 100))}%)",
        f"Ngày lập: {_ngay_chuoi(ngay)}",
        f"Người lập: {nguoi_lap or '—'}",
        f"Tóm tắt: {tk.get('dat', 0)} đạt, {tk.get('canh_bao', 0)} cảnh báo, "
        f"{tk.get('thieu', 0)} thiếu — Mức rủi ro: "
        f"{NHAN_RUI_RO.get(tk.get('muc_rui_ro', ''), '—')}",
    ]
    for s in dong:
        _dat_font(d.add_paragraph().add_run(s))

    cot = ["#", "Mục", "Trạng thái", "Điều khoản liên quan", "Giải thích", "Đề xuất sửa",
           "Căn cứ"]
    bang = d.add_table(rows=1, cols=len(cot))
    bang.style = "Table Grid"
    for o, ten in zip(bang.rows[0].cells, cot):
        o.text = ""
        _dat_font(o.paragraphs[0].add_run(ten), size=11, bold=True)
    for i, m in enumerate(ket_qua.get("muc") or [], 1):
        hang = bang.add_row().cells
        gia_tri = [str(i), m.get("ten", ""), NHAN_TRANG_THAI.get(m.get("trang_thai"), ""),
                   m.get("dieu_khoan") or "—", m.get("giai_thich", ""),
                   m.get("de_xuat") or "—", m.get("can_cu", "")]
        for o, v in zip(hang, gia_tri):
            o.text = ""
            _dat_font(o.paragraphs[0].add_run(v), size=11,
                      bold=(v in ("CẢNH BÁO", "THIẾU")) or None)

    d.add_paragraph()
    _dat_font(d.add_paragraph().add_run(
        "Lưu ý: báo cáo do AI hỗ trợ lập trên cơ sở đối chiếu từ khoá với danh mục "
        "điều khoản chuẩn; luật sư phụ trách rà soát lại toàn văn và chịu trách nhiệm "
        "cuối cùng về nội dung tư vấn."), italic=True)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
