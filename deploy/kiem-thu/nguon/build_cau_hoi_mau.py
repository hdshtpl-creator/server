# -*- coding: utf-8 -*-
"""Sinh BO_CAU_HOI_MAU.md + BO_CAU_HOI_TAI_LIEU_TIENG_ANH.md + PHIEU_CAU_HOI_MAU.xlsx
từ cau_hoi_mau.py (câu hỏi) + do_nguon_0510.json (máy đo TRƯỚC sửa) + do_nguon_0510_sau.json
(máy đo SAU khi sửa và cập nhật máy chủ cùng ngày).
Chạy: python deploy/kiem-thu/nguon/build_cau_hoi_mau.py deploy/kiem-thu"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import cau_hoi_mau as Q  # noqa: E402
import kb_cases  # noqa: E402
from build_docs import bang, plain  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(HERE)
DO_TRUOC = json.load(open(os.path.join(HERE, "do_nguon_0510.json"), encoding="utf-8"))
DO = json.load(open(os.path.join(HERE, "do_nguon_0510_sau.json"), encoding="utf-8"))
TEN_VN = {m: (t, b) for m, t, b in Q.NHOM_VN}
TEN_EN = {m: (t, g) for m, t, g in Q.NHOM_EN}
KB = [k for k in kb_cases.load() if k["id"] not in Q.KB_DA_DUNG]


def ghi_do(ma, kiem):
    """Kết quả máy đo SAU sửa, kèm kết quả trước sửa khi khác."""
    sau = _ghi_do(DO, ma, kiem)
    truoc = _ghi_do(DO_TRUOC, ma, kiem) if ma in DO_TRUOC else None
    if truoc and truoc != sau and not sau.startswith("—"):
        return f"{sau} (trước sửa: {truoc})"
    return sau


def _ghi_do(DO, ma, kiem):
    """Một dòng mô tả kết quả máy đo phần tìm nguồn cho câu `ma`."""
    if not kiem:
        return "— (cần file đính kèm / thao tác, máy không đo)"
    r = DO.get(ma)
    if not r:
        return "— (chưa đo)"
    if kiem.get("truc_tiep"):
        return "Trả thẳng từ CSDL" if r["ok"] else "CHƯA — không trả thẳng từ CSDL"
    vt = r.get("vi_tri") or []
    if "doc_id" in kiem:
        return f"Có — đúng tài liệu ở [Nguồn {vt[0]}]" if r["ok"] else "CHƯA — không lấy được tài liệu đã chọn"
    if "ngan" in kiem:
        if not r["ok"]:
            return "CHƯA — không có tài liệu SEC nào trong nguồn"
        if vt[0] <= 10:
            return f"Có — mẫu SEC ở [Nguồn {vt[0]}]"
        return f"Có nhưng xa — mẫu SEC từ [Nguồn {vt[0]}], sau {vt[0] - 1} đoạn luật VN (dễ bị bỏ qua)"
    d = kiem["dieu"]
    if not r["ok"]:
        dau = (r.get("dau") or ["?"])[0].split(": ", 1)[-1]
        return f"CHƯA — 8 nguồn đầu không có Điều {d} (đứng đầu: {dau})"
    if vt[0] > 12:
        return f"Có nhưng xa — Điều {d} ở [Nguồn {vt[0]}]"
    return f"Có — Điều {d} ở [Nguồn {vt[0]}]"


def thong_ke(rows, idx_kiem, DO=DO):
    do = [r for r in rows if r[idx_kiem] and r[0] in DO]
    co = sum(1 for r in do if DO.get(r[0], {}).get("ok"))
    return co, len(do)


CHAM_TXT = ("Mỗi câu: nhìn **huy hiệu kiểm chứng** → bấm một chip **[Nguồn n] → Xem trước** xem đúng văn bản / đúng "
            "Điều / đúng tài liệu chưa → rồi mới đọc nội dung. Ghi kết quả vào phiếu **PHIEU_CAU_HOI_MAU.xlsx** "
            "(ô nền vàng). Câu Sai / Bịa: bấm 👎 ngay dưới câu trả lời, ghi sai ở đâu + căn cứ đúng.")

MAY_DO_TXT = ("Cột **Máy đo 05/10** là kết quả chạy thử phần TÌM NGUỒN trên máy chủ thật (không gọi model) SAU KHI sửa "
              "các phát hiện ở mục dưới và cập nhật máy chủ cùng ngày; ngoặc *(trước sửa: …)* là kết quả lần đo đầu. "
              "Máy chỉ xác nhận bot ĐÃ CÓ đúng nguồn trong tay — câu trả lời đúng hay sai vẫn phải người chấm. Câu nào "
              "máy báo có nguồn mà bot vẫn trả lời sai → lỗi ở phần viết câu trả lời, ghi rõ vào phiếu.")


def md_vn():
    co, tong = thong_ke(Q.VN, 8)
    co0, tong0 = thong_ke(Q.VN, 8, DO_TRUOC)
    L = ["# BỘ CÂU HỎI MẪU — NHÂN VIÊN THỬ THÊM", "",
         f"*Lập ngày {Q.NGAY} · {len(Q.VN)} câu theo {len(Q.NHOM_VN)} loại bài + {len(KB)} tình huống dài · "
         f"máy đo nguồn sau sửa: {co}/{tong} câu đã có đúng căn cứ (trước sửa {co0}/{tong0})*", "",
         "Bộ này **mở rộng từng loại bài** trong [KIEM_THU_CHO_NHAN_VIEN.md](KIEM_THU_CHO_NHAN_VIEN.md) (B1…D12): mỗi loại "
         "thêm vài câu ở lĩnh vực khác để thử rộng hơn. Làm xong bộ 43 bài gốc rồi mới làm bộ này; mỗi người chọn các câu "
         "đúng phòng mình (cột *Ai làm*), mỗi câu một **Cuộc trò chuyện mới** trừ nhóm 6 (hỏi nối tiếp).", "",
         "Mọi đáp án ở cột *Đạt khi* đã **đối chiếu với văn bản đang có trong kho** ngày 05/10/2026 (đúng số Điều, đúng con "
         "số) — không lấy từ trí nhớ. Kho dùng bản luật/văn bản hợp nhất mới nhất; nếu công ty chốt bản khác thì chấm theo "
         "bản công ty chốt.", "",
         "Bộ câu riêng cho **tài liệu tiếng Anh mới học** (2.304 hợp đồng Hoa Kỳ – SEC): "
         "[BO_CAU_HOI_TAI_LIEU_TIENG_ANH.md](BO_CAU_HOI_TAI_LIEU_TIENG_ANH.md).", "",
         "## 1. Cách làm & cách chấm", "", CHAM_TXT, "",
         bang(["Kết quả", "Khi nào"], Q.CHAM), "", MAY_DO_TXT, "",
         "## 2. Đọc trước — điều đã thấy khi soạn bộ câu (đã sửa 05/10)", "",
         bang(["Mã", "Mức", "Phát hiện", "Chi tiết", "Câu liên quan", "Đã xử lý"],
              [p for p in Q.PHAT_HIEN if not p[4].startswith("EN-")]), ""]
    so = 3
    for ma, ten, bai in Q.NHOM_VN:
        rows = [r for r in Q.VN if r[1] == ma]
        L += [f"## {so}. {ten} *(mở rộng bài {bai})*", ""]
        so += 1
        for r in rows:
            L += [f"### {r[0]}", "",
                  f"- **Làm:** {r[3]}", "",
                  f"> {r[4]}", "",
                  f"- **Ai làm:** {r[2]}",
                  f"- **Đạt khi:** {r[5]}",
                  f"- **Căn cứ:** {r[6]}"]
            if r[7]:
                L.append(f"- **Lưu ý:** {r[7]}")
            L += [f"- **Máy đo 05/10:** {ghi_do(r[0], r[8])}",
                  "- [ ] Đúng  [ ] Một phần  [ ] Sai  [ ] Bịa — ghi chú: ……", ""]
    L += [f"## {so}. Tình huống dài bổ sung *(mở rộng bài B4)*", "",
          f"{len(KB)} tình huống còn lại của bộ 40 kịch bản HDS (*Kịch bản test AI*) — phụ lục bài B4 mới dùng "
          f"{len(Q.KB_DA_DUNG)} câu. Dán **nguyên văn** cột *Câu gửi bot* vào cuộc trò chuyện mới. Cột *Máy chấm 03/10* là "
          "lượt chấm sơ bộ tự động (chỉ đếm Điều được dẫn): **BỎ QUA** = tiêu chí không nêu số Điều, luật sư phải chấm; "
          "**SAI** = bot không dẫn Điều nào trong tiêu chí — ưu tiên thử lại các câu này.", "",
          bang(["Mã", "Lĩnh vực · Chủ đề", "Phòng", "Câu gửi bot", "Phải dẫn", "Máy chấm 03/10"],
               [(k["id"], f"{k['linh_vuc']} · {k['tieu_de']}", Q.KB_PHONG[k["id"].split(".")[0]], k["cau_hoi"],
                 k["dieu"], Q.KB_MAY_CHAM_0310.get(k["id"], "—")) for k in KB]), ""]
    return "\n".join(L)


def md_en():
    co, tong = thong_ke(Q.EN, 6)
    co0, tong0 = thong_ke(Q.EN, 6, DO_TRUOC)
    L = ["# BỘ CÂU HỎI MẪU — TÀI LIỆU TIẾNG ANH MỚI HỌC (HỢP ĐỒNG HOA KỲ – SEC)", "",
         f"*Lập ngày {Q.NGAY} · {len(Q.EN)} câu · máy đo nguồn sau sửa: {co}/{tong} câu có đúng tài liệu "
         f"(trước sửa {co0}/{tong0})*", "",
         "## 1. Lô tài liệu này là gì", "",
         "- **2.304 tài liệu tiếng Anh** trong ngăn *3. HỢP ĐỒNG MẪU / Hợp đồng Hoa Kỳ – SEC*, học và duyệt 28–30/09/2026: "
         "phụ lục hợp đồng (EX-10) các công ty Mỹ nộp lên SEC EDGAR — tài liệu công khai, luật Hoa Kỳ.",
         "- **Dieu_khoan/** — 2.067 điều khoản tách riêng theo 25 loại: Bất khả kháng, Bảo mật, Bồi thường, Cam đoan – bảo "
         "đảm, Cam kết, Chuyển nhượng, Điều kiện tiên quyết, Định nghĩa, Dữ liệu, Giải quyết tranh chấp, Giao hàng – nghiệm "
         "thu, Giới hạn trách nhiệm, Hạn chế cạnh tranh, Kiểm toán, Luật áp dụng, Quản trị, Sở hữu trí tuệ, Điều khoản chung, "
         "Thanh toán, Thời hạn – chấm dứt, Thông báo, Thuế, Tuân thủ, Vi phạm – chế tài, Bảo hiểm.",
         "- **Hop_dong_day_du/** — 237 hợp đồng đầy đủ theo 12 loại: NDA, Li-xăng, Cung ứng, Đầu tư – cổ đông, Dịch vụ, Liên "
         "doanh, Lao động – quản lý, M&A, Phân phối – đại lý, Chuyển giao công nghệ, Thuê bất động sản, Vay – tín dụng.",
         "- Mỗi điều khoản có: **bản gốc tiếng Anh**, **bản dịch tham khảo tiếng Việt** (máy dịch — có chỗ sai thuật ngữ), "
         "tên hợp đồng nguồn, ngày nộp SEC và **đường dẫn bản gốc trên sec.gov**.", "",
         "## 2. Ai làm", "",
         "Từ 05/10 ngăn *Hợp đồng Hoa Kỳ – SEC* (tài liệu công khai trên sec.gov) mở cho **mọi tài khoản nội bộ**, kể cả "
         "trợ lý và chuyên viên DN-ĐT / SHTT. Mẫu hợp đồng **của HDS** (tiếng Việt) vẫn theo ma trận quyền cũ: admin / Ban "
         "QT, trưởng bộ phận, chuyên viên HTPL-TVTX và Tranh tụng. Câu EN-F1 dành cho trợ lý để xác nhận cả hai điều này.", "",
         "## 3. Cách làm & cách chấm", "", CHAM_TXT, "",
         bang(["Kết quả", "Khi nào"], Q.CHAM), "", MAY_DO_TXT, "",
         "## 4. Đọc trước — điều đã thấy khi soạn bộ câu (đã sửa 05/10)", "",
         bang(["Mã", "Mức", "Phát hiện", "Chi tiết", "Câu liên quan", "Đã xử lý"],
              [p for p in Q.PHAT_HIEN if "EN-" in p[4]]), ""]
    so = 5
    for ma, ten, goi_y in Q.NHOM_EN:
        rows = [r for r in Q.EN if r[1] == ma]
        L += [f"## {so}. Nhóm {ma} — {ten}", ""]
        so += 1
        if goi_y:
            L += [goi_y, ""]
        for r in rows:
            L += [f"### {r[0]}", ""]
            if r[2]:
                L += [f"- **Chọn nguồn:** gõ tìm *{Q.TIM_NGUON.get(r[0], r[2])}* → tick `{r[2]}` → Dùng 1 nguồn"]
            L += [f"- **Gõ:** {r[3]}",
                  f"- **Đạt khi:** {r[4]}"]
            if r[5]:
                L.append(f"- **Lưu ý:** {r[5]}")
            L += [f"- **Máy đo 05/10:** {ghi_do(r[0], r[6])}",
                  "- [ ] Đúng  [ ] Một phần  [ ] Sai  [ ] Bịa — ghi chú: ……", ""]
    return "\n".join(L)


def build_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    navy = "1F3A5F"
    fn, fb = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    fh, ft = Font(name=F, size=10, bold=True, color="FFFFFF"), Font(name=F, size=14, bold=True, color=navy)
    fnote = Font(name=F, size=9, italic=True, color="595959")
    fdo_ok, fdo_ko = Font(name=F, size=10, color="006100"), Font(name=F, size=10, bold=True, color="9C0006")
    head = PatternFill("solid", fgColor=navy)
    vang = PatternFill("solid", fgColor="FFF2CC")
    nhom_fill = PatternFill("solid", fgColor="DDEBF7")
    s = Side(style="thin", color="BFBFBF")
    bd = Border(left=s, right=s, top=s, bottom=s)
    wrap = Alignment(wrap_text=True, vertical="top")
    ctr = Alignment(horizontal="center", vertical="top", wrap_text=True)

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    def header(ws, row, cols, widths):
        for j, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(row=row, column=j, value=h)
            c.font, c.fill, c.alignment, c.border = fh, head, ctr, bd
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[row].height = 30

    def chon_ket_qua(ws, col, r0, r1):
        dv = DataValidation(type="list", formula1='"Đúng,Một phần,Sai,Bịa,Bỏ qua"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}{r0}:{col}{r1}")
        for val, color in (("Đúng", "C6EFCE"), ("Một phần", "FFEB9C"), ("Sai", "FFC7CE"), ("Bịa", "FF7C80")):
            ws.conditional_formatting.add(f"{col}{r0}:{col}{r1}", CellIsRule(
                operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=color)))

    def tong_hop(ws, col, r0, r1, at_row, label_col=3):
        ws.cell(row=at_row, column=label_col, value="TỔNG HỢP").font = fb
        for k, kq in enumerate(("Đúng", "Một phần", "Sai", "Bịa", "Bỏ qua"), 1):
            ws.cell(row=at_row + k, column=label_col, value=kq).font = fn
            ws.cell(row=at_row + k, column=label_col + 1,
                    value=f'=COUNTIF(${col}${r0}:${col}${r1},"{kq}")').font = fn
        c = ws.cell(row=at_row + 6, column=label_col, value="Tỉ lệ Đúng + Một phần (trên câu đã chấm)")
        c.font = fb
        a, b = get_column_letter(label_col + 1), at_row
        c = ws.cell(row=at_row + 6, column=label_col + 1,
                    value=f'=IFERROR(({a}{b + 1}+{a}{b + 2})/({a}{b + 1}+{a}{b + 2}+{a}{b + 3}+{a}{b + 4}),"—")')
        c.font, c.number_format = fb, "0%"

    def phieu_dau(ws, tieu_de, ghi):
        ws["A1"] = tieu_de
        ws["A1"].font = ft
        ws["A2"] = ghi
        ws["A2"].font = fnote
        for i, k in enumerate(("Họ tên", "Phòng", "Vai", "Ngày thử", "Mã bản (chân trang)"), 3):
            ws.cell(row=i, column=2, value=k).font = fb
            c = ws.cell(row=i, column=3)
            c.fill, c.border = vang, bd

    def o_do(c, txt):
        c.value = txt
        c.font = fdo_ko if txt.startswith("CHƯA") else (fdo_ok if txt.startswith(("Có —", "Trả thẳng")) else fn)

    # ---------------- Câu hỏi mẫu (tiếng Việt)
    ws = wb.active
    ws.title = "Câu hỏi mẫu"
    phieu_dau(ws, "BỘ CÂU HỎI MẪU — NHÂN VIÊN THỬ THÊM (mở rộng bộ 43 bài)",
              "Điền ô nền vàng. Mỗi câu một cuộc trò chuyện mới (trừ nhóm 6). Hướng dẫn: BO_CAU_HOI_MAU.md. "
              "Cột 'Máy đo 05/10': kết quả sau khi sửa (ngoặc = trước sửa). Có nguồn mà vẫn trả lời sai → ghi rõ.")
    cols = ["Mã", "Nhóm (bài gốc)", "Ai làm", "Làm gì", "Câu gửi bot", "Đạt khi", "Căn cứ", "Lưu ý", "Máy đo 05/10",
            "Kết quả", "Thời gian (giây)", "Ghi chú (sai ở đâu, căn cứ đúng)"]
    header(ws, 9, cols, [7, 20, 16, 26, 50, 60, 20, 30, 30, 12, 10, 36])
    r = 10
    r0 = r
    for ma, ten, bai in Q.NHOM_VN:
        c = ws.cell(row=r, column=1, value=f"Nhóm {ma} — {ten} (mở rộng bài {bai})")
        c.font, c.fill = fb, nhom_fill
        for j in range(2, len(cols) + 1):
            ws.cell(row=r, column=j).fill = nhom_fill
        r += 1
        for row in [x for x in Q.VN if x[1] == ma]:
            vals = [row[0], f"{ten} ({bai})", row[2], plain(row[3]), row[4], plain(row[5]), row[6], plain(row[7])]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.font, c.alignment, c.border = fn, wrap, bd
            c = ws.cell(row=r, column=9)
            c.alignment, c.border = wrap, bd
            o_do(c, ghi_do(row[0], row[8]))
            for j in (10, 11, 12):
                c = ws.cell(row=r, column=j)
                c.fill, c.border, c.alignment = vang, bd, wrap
            r += 1
    r1 = r - 1
    chon_ket_qua(ws, "J", r0, r1)
    ws.freeze_panes = f"F{r0}"
    ws.auto_filter.ref = f"A9:L{r1}"
    tong_hop(ws, "J", r0, r1, r1 + 2)

    # ---------------- Tình huống bổ sung (24 KB)
    wk = wb.create_sheet("Tình huống bổ sung")
    phieu_dau(wk, f"{len(KB)} TÌNH HUỐNG DÀI BỔ SUNG (bộ 40 kịch bản HDS — phần chưa có ở bài B4)",
              "Dán nguyên văn cột D vào một cuộc trò chuyện mới. 'Máy chấm 03/10' chỉ đếm Điều được dẫn: "
              "BỎ QUA = luật sư phải chấm; SAI = ưu tiên thử lại.")
    cols = ["Mã", "Lĩnh vực · Chủ đề", "Phòng", "Câu gửi bot (nguyên văn)", "Tiêu chí của HDS", "Phải dẫn",
            "Máy chấm 03/10", "Kết quả", "Điều/văn bản bot dẫn", "Thời gian (giây)", "Ghi chú"]
    header(wk, 9, cols, [7, 30, 18, 70, 50, 32, 12, 12, 30, 10, 30])
    for i, k in enumerate(KB, 10):
        vals = [k["id"], f"{k['linh_vuc']} · {k['tieu_de']}", Q.KB_PHONG[k["id"].split(".")[0]], k["cau_hoi"],
                k["tieu_chi"], k["dieu"], Q.KB_MAY_CHAM_0310.get(k["id"], "—")]
        for j, v in enumerate(vals, 1):
            c = wk.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = fn, wrap, bd
            if j == 7 and v == "SAI":
                c.font = fdo_ko
        for j in range(8, 12):
            c = wk.cell(row=i, column=j)
            c.fill, c.border, c.alignment = vang, bd, wrap
    k1 = 9 + len(KB)
    chon_ket_qua(wk, "H", 10, k1)
    wk.freeze_panes = "D10"
    wk.auto_filter.ref = f"A9:K{k1}"
    tong_hop(wk, "H", 10, k1, k1 + 2)

    # ---------------- Tài liệu tiếng Anh
    we = wb.create_sheet("Tài liệu tiếng Anh")
    phieu_dau(we, "BỘ CÂU HỎI — TÀI LIỆU TIẾNG ANH MỚI HỌC (2.304 hợp đồng Hoa Kỳ – SEC)",
              "Mọi tài khoản nội bộ (ngăn SEC mở cho tất cả từ 05/10); EN-F1 dành cho trợ lý. "
              "Nhóm A: bấm Chọn nguồn, gõ tên ở cột C. Hướng dẫn: BO_CAU_HOI_TAI_LIEU_TIENG_ANH.md.")
    cols = ["Mã", "Nhóm", "Chọn nguồn (tên tài liệu)", "Câu gửi bot", "Đạt khi", "Lưu ý", "Máy đo 05/10",
            "Kết quả", "Thời gian (giây)", "Ghi chú"]
    header(we, 9, cols, [8, 24, 36, 50, 62, 30, 30, 12, 10, 36])
    r = 10
    for ma, ten, _ in Q.NHOM_EN:
        c = we.cell(row=r, column=1, value=f"Nhóm {ma} — {ten}")
        c.font, c.fill = fb, nhom_fill
        for j in range(2, len(cols) + 1):
            we.cell(row=r, column=j).fill = nhom_fill
        r += 1
        for row in [x for x in Q.EN if x[1] == ma]:
            chon = (f"Gõ tìm: {Q.TIM_NGUON.get(row[0], row[2])}\nTick: {row[2]}") if row[2] else "—"
            vals = [row[0], ten, chon, row[3], plain(row[4]), plain(row[5])]
            for j, v in enumerate(vals, 1):
                c = we.cell(row=r, column=j, value=v)
                c.font, c.alignment, c.border = fn, wrap, bd
            c = we.cell(row=r, column=7)
            c.alignment, c.border = wrap, bd
            o_do(c, ghi_do(row[0], row[6]))
            for j in (8, 9, 10):
                c = we.cell(row=r, column=j)
                c.fill, c.border, c.alignment = vang, bd, wrap
            r += 1
    e1 = r - 1
    chon_ket_qua(we, "H", 10, e1)
    we.freeze_panes = "E10"
    we.auto_filter.ref = f"A9:J{e1}"
    tong_hop(we, "H", 10, e1, e1 + 2)

    # ---------------- Cách chấm & phát hiện
    wc = wb.create_sheet("Cách chấm & phát hiện")
    wc["A1"] = "CÁCH CHẤM"
    wc["A1"].font = ft
    header(wc, 2, ["Kết quả", "Khi nào"], [14, 110])
    for i, (a, b) in enumerate(Q.CHAM, 3):
        for j, v in enumerate((a, b), 1):
            c = wc.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = fn, wrap, bd
    rr = 3 + len(Q.CHAM) + 1
    wc.cell(row=rr, column=1, value=f"PHÁT HIỆN KHI SOẠN BỘ CÂU ({Q.NGAY}) — đã sửa và cập nhật máy chủ cùng ngày").font = ft
    header(wc, rr + 1, ["Mã", "Mức", "Phát hiện", "Chi tiết", "Câu liên quan", "Đã xử lý"], [6, 8, 40, 80, 24, 80])
    for i, p in enumerate(Q.PHAT_HIEN, rr + 2):
        for j, v in enumerate(p, 1):
            c = wc.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = fn, wrap, bd
    wc.column_dimensions["B"].width = 110
    for w in wb.worksheets:
        w.sheet_view.zoomScale = 90
    wb.save(path)


if __name__ == "__main__":
    for ten, noi_dung in (("BO_CAU_HOI_MAU.md", md_vn()), ("BO_CAU_HOI_TAI_LIEU_TIENG_ANH.md", md_en())):
        with open(os.path.join(OUT, ten), "w", encoding="utf-8", newline="\n") as f:
            f.write(noi_dung + "\n")
    build_xlsx(os.path.join(OUT, "PHIEU_CAU_HOI_MAU.xlsx"))
    print("đã sinh", OUT)
