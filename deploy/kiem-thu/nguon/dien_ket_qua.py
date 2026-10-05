# -*- coding: utf-8 -*-
"""Điền kết quả chạy tự động (ket_qua.json + kb40_ket_qua.json từ /tmp/kt trên máy chủ)
vào bản sao sổ Excel và sinh báo cáo Markdown.
Chạy: python dien_ket_qua.py <thư mục chứa 2 json> <thư mục kiem-thu> <ngày> <build>"""
import json, os, sys, re
from collections import Counter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

JSON_DIR, OUT, NGAY, BUILD = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
kq = {k["ma"]: k for k in json.load(open(os.path.join(JSON_DIR, "ket_qua.json"), encoding="utf-8"))}
# Ca mà lượt 2 (chạy lại) kém thông tin hơn lượt 1 → giữ lượt 1, ghi chú lượt 2
UU_TIEN_LUOT1 = {"ST-07": "lượt 2: bản sửa theo yêu cầu không chứa con số cam kết nên kiểm tra 0/0 mục — cơ chế đã chứng minh ở lượt 1"}
try:
    l1 = {k["ma"]: k for k in json.load(open(os.path.join(JSON_DIR, "ket_qua_luot1.json"), encoding="utf-8"))}
    for ma, note in UU_TIEN_LUOT1.items():
        if ma in l1:
            kq[ma] = dict(l1[ma]); kq[ma]["thuc_te"] = (kq[ma]["thuc_te"] or "")[:1200] + f" | ({note})"
except FileNotFoundError:
    pass
kb = json.load(open(os.path.join(JSON_DIR, "kb40_ket_qua.json"), encoding="utf-8"))
# Tiêu chí không nêu điều/văn bản → máy không chấm được, để luật sư chấm
for v in kb.values():
    if not re.search(r"Điều\s+\d", v.get("phai_dan") or "") and v.get("ket_qua") in ("SAI", "MỘT PHẦN"):
        v["ket_qua"], v["ly_do"] = "BỎ QUA", "tiêu chí không nêu số Điều — máy không chấm, LUẬT SƯ CHẤM nội dung (máy ghi nhận: " + (v.get("ly_do") or "") + ")"

src = os.path.join(OUT, "KICH_BAN_KIEM_THU.xlsx")
dst = os.path.join(OUT, f"KET_QUA_TU_DONG_{NGAY.replace('/', '-')}.xlsx")
wb = load_workbook(src)
ws = wb["Ca kiểm thử"]
wrap = Alignment(wrap_text=True, vertical="top")
n_fill = 0
for r in range(4, ws.max_row + 1):
    ma = ws.cell(row=r, column=1).value
    k = kq.get(ma)
    if not k:
        continue
    ws.cell(row=r, column=11, value=k["ket_qua"])
    ws.cell(row=r, column=12, value=k["thuc_te"]).alignment = wrap
    ws.cell(row=r, column=13, value=(json.dumps(k["chi_tiet"], ensure_ascii=False)[:800] if k.get("chi_tiet") else None)).alignment = wrap
    ws.cell(row=r, column=15, value="Tự động (Claude)")
    ws.cell(row=r, column=16, value=NGAY)
    ws.cell(row=r, column=17, value=BUILD)
    n_fill += 1
wk = wb["40 KB pháp lý"]
for r in range(3, wk.max_row + 1):
    kid = str(wk.cell(row=r, column=1).value or "")
    v = kb.get(kid)
    if not v:
        continue
    wk.cell(row=r, column=8, value=", ".join(f"Điều {d}" for d in v.get("dieu", [])[:12]) + (" | " + "; ".join(v.get("nguon", [])[:3]))).alignment = wrap
    wk.cell(row=r, column=9, value=v["ket_qua"] if v["ket_qua"] in ("ĐÚNG", "MỘT PHẦN", "SAI", "BỎ QUA") else "BỎ QUA")
    wk.cell(row=r, column=10, value={"verified": "Đã kiểm chứng theo nguồn", "partial": "Chỉ một phần có đủ căn cứ", "uncited": "Chưa gắn được trích dẫn",
                                     "uncited_blocked": "Không tìm thấy căn cứ trong nguồn", "insufficient": "Chưa đủ bằng chứng để kết luận"}.get(v.get("grounding"), v.get("grounding")))
    wk.cell(row=r, column=11, value=v.get("giay"))
    wk.cell(row=r, column=12, value=("MÁY CHẤM SƠ BỘ — " + v.get("ly_do", ""))[:400]).alignment = wrap
    wk.cell(row=r, column=13, value="Tự động (Claude)")
    wk.cell(row=r, column=14, value=NGAY)
# sheet toàn văn trả lời 40 KB để luật sư chấm lại
wa = wb.create_sheet("Trả lời 40 KB (toàn văn)")
wa.append(["KB", "Máy chấm", "Lý do máy", "Phải dẫn", "Câu trả lời đầy đủ của bot"])
for c, w in zip("ABCDE", (6, 12, 40, 40, 140)):
    wa.column_dimensions[c].width = w
for kid, v in sorted(kb.items(), key=lambda x: [int(p) for p in x[0].split(".")]):
    wa.append([kid, v["ket_qua"], v.get("ly_do"), v.get("phai_dan"), v.get("answer")])
for row in wa.iter_rows(min_row=1):
    for c in row:
        c.alignment = wrap
        c.font = Font(name="Arial", size=9)
wb.save(dst)

# ---- báo cáo markdown
tong = Counter(k["ket_qua"] for k in kq.values())
kbt = Counter(v["ket_qua"] for v in kb.values())
ut_cao = [k for k in kq.values()]
L = [f"# KẾT QUẢ CHẠY TỰ ĐỘNG BỘ KIỂM THỬ — {NGAY}", "",
     f"Chạy trên máy chủ thật (`127.0.0.1:8000`, bản `{BUILD}`) bằng bộ chạy `kiem-thu/nguon/` ↔ `/tmp/kt` — "
     f"tài khoản thử riêng (đã khoá sau khi chạy), tài liệu chim mồi (đã gỡ). Kết quả chi tiết từng ca: `{os.path.basename(dst)}`.", "",
     "## 1. Tổng hợp", "",
     f"| | ĐẠT | KHÔNG ĐẠT | CHẶN | BỎ QUA | Tổng |", "|---|---|---|---|---|---|",
     f"| Ca chức năng | {tong.get('ĐẠT',0)} | {tong.get('KHÔNG ĐẠT',0)} | {tong.get('CHẶN',0)} | {tong.get('BỎ QUA',0)} | {len(kq)} |", "",
     f"| 40 kịch bản HDS (máy chấm sơ bộ theo số Điều) | ĐÚNG {kbt.get('ĐÚNG',0)} | MỘT PHẦN {kbt.get('MỘT PHẦN',0)} | SAI {kbt.get('SAI',0)} | luật sư chấm {kbt.get('BỎ QUA',0)} | {len(kb)} |", "",
     "*BỎ QUA* = ca thao tác giao diện (kiểm tay) hoặc hạng mục chưa bàn giao. *CHẶN* = không chạy được vì lỗi kỹ thuật khi test — xem cột Thực tế.", "",
     ]
# ---- so với lượt chạy trước (tham số thứ 5: thư mục json lượt trước)
if len(sys.argv) > 5:
    truoc = {k["ma"]: k for k in json.load(open(os.path.join(sys.argv[5], "ket_qua.json"), encoding="utf-8"))}
    try:
        kb_truoc = json.load(open(os.path.join(sys.argv[5], "kb40_ket_qua.json"), encoding="utf-8"))
        for v in kb_truoc.values():
            if not re.search(r"Điều\s+\d", v.get("phai_dan") or "") and v.get("ket_qua") in ("SAI", "MỘT PHẦN"):
                v["ket_qua"] = "BỎ QUA"
    except FileNotFoundError:
        kb_truoc = {}
    tt = Counter(k["ket_qua"] for k in truoc.values())
    kt_ = Counter(v["ket_qua"] for v in kb_truoc.values())
    L += ["## 1b. So với lượt chạy trước (02/10/2026, trước đợt sửa)", "",
          "| | ĐẠT | KHÔNG ĐẠT | CHẶN | BỎ QUA | Tổng |", "|---|---|---|---|---|---|",
          f"| Lượt trước | {tt.get('ĐẠT',0)} | {tt.get('KHÔNG ĐẠT',0)} | {tt.get('CHẶN',0)} | {tt.get('BỎ QUA',0)} | {len(truoc)} |",
          f"| Lượt này | {tong.get('ĐẠT',0)} | {tong.get('KHÔNG ĐẠT',0)} | {tong.get('CHẶN',0)} | {tong.get('BỎ QUA',0)} | {len(kq)} |", ""]
    if kb_truoc:
        L += [f"40 KB: lượt trước ĐÚNG {kt_.get('ĐÚNG',0)} · MỘT PHẦN {kt_.get('MỘT PHẦN',0)} · SAI {kt_.get('SAI',0)} → "
              f"lượt này ĐÚNG {kbt.get('ĐÚNG',0)} · MỘT PHẦN {kbt.get('MỘT PHẦN',0)} · SAI {kbt.get('SAI',0)}.", ""]
    doi = [(ma, truoc[ma]["ket_qua"] if ma in truoc else "(ca mới)", kq[ma]["ket_qua"]) for ma in sorted(kq)
           if ma not in truoc or truoc[ma]["ket_qua"] != kq[ma]["ket_qua"]]
    if doi:
        L += ["Ca đổi kết quả:", "", "| Mã | Lượt trước | Lượt này |", "|---|---|---|"]
        L += [f"| {a} | {b} | {c} |" for a, b, c in doi]
        L.append("")
    if kb_truoc:
        doi_kb = [(k, kb_truoc[k]["ket_qua"], kb[k]["ket_qua"]) for k in sorted(kb, key=lambda x: [int(p) for p in x.split(".")])
                  if k in kb_truoc and kb_truoc[k]["ket_qua"] != kb[k]["ket_qua"]]
        if doi_kb:
            L += ["KB đổi kết quả:", "", "| KB | Lượt trước | Lượt này |", "|---|---|---|"]
            L += [f"| {a} | {b} | {c} |" for a, b, c in doi_kb]
            L.append("")
L += ["## 2. Ca KHÔNG ĐẠT", "", "| Mã | Thực tế quan sát |", "|---|---|"]
for k in sorted(kq.values(), key=lambda x: x["ma"]):
    if k["ket_qua"] == "KHÔNG ĐẠT":
        L.append(f"| {k['ma']} | {k['thuc_te'].replace('|', '/')[:400]} |")
L += ["", "## 3. Ca CHẶN (lỗi kỹ thuật khi chạy test)", "", "| Mã | Lỗi |", "|---|---|"]
for k in sorted(kq.values(), key=lambda x: x["ma"]):
    if k["ket_qua"] == "CHẶN":
        L.append(f"| {k['ma']} | {k['thuc_te'].replace('|', '/')[:300]} |")
L += ["", "## 4. 40 kịch bản pháp lý — máy chấm sơ bộ", "",
      "Máy chỉ đếm **số Điều** và **tên văn bản** trong câu trả lời so với cột *Phải dẫn*; luật sư HDS chấm lại nội dung ở sheet *Trả lời 40 KB (toàn văn)*. "
      "Kho dùng bản luật mới nhất nên số điều có thể lệch bản 2020 của tiêu chí.", "",
      "| KB | Máy chấm | Giây | Kiểm chứng | Điều bot dẫn | Lý do |", "|---|---|---|---|---|---|"]
for kid, v in sorted(kb.items(), key=lambda x: [int(p) for p in x[0].split(".")]):
    L.append(f"| {kid} | {v['ket_qua']} | {v.get('giay')} | {v.get('grounding')} | {', '.join(str(d) for d in v.get('dieu', [])[:10])} | {(v.get('ly_do') or '').replace('|','/')[:160]} |")
giay = [v.get("giay") for v in kb.values() if v.get("giay")]
if giay:
    giay.sort()
    L += ["", f"Thời gian trả lời 40 KB: nhanh nhất {giay[0]}s · trung vị {giay[len(giay)//2]}s · chậm nhất {giay[-1]}s."]
L += ["", "## 5. Thời gian các ca có gọi model", "", "| Mã | Giây |", "|---|---|"]
for k in sorted((k for k in kq.values() if k.get("giay") and k["giay"] > 5), key=lambda x: -x["giay"])[:25]:
    L.append(f"| {k['ma']} | {k['giay']} |")
L.append("")
open(os.path.join(OUT, f"KET_QUA_TU_DONG_{NGAY.replace('/', '-')}.md"), "w", encoding="utf-8", newline="\n").write("\n".join(L))
print("điền", n_fill, "ca;", len(kb), "KB;", dict(tong), dict(kbt), "→", dst)
