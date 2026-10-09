# -*- coding: utf-8 -*-
"""Sinh KICH_BAN_KIEM_THU.md + KICH_BAN_KIEM_THU.xlsx + du-lieu-mau/README.md từ một nguồn."""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import cases_b  # noqa: E402  (nạp cả cases_a)
from cases_a import C  # noqa: E402
import kb_cases, md_static as S, findings as F  # noqa: E402

OUT = sys.argv[1]
NGAY = "03/10/2026"

NHOM = [
    ("TK", "Đăng nhập & tài khoản"),
    ("PQ", "Phân quyền & cô lập dữ liệu"),
    ("CH", "Hội thoại AI (kênh nội bộ)"),
    ("KT", "Kiểm tra pháp lý & mẫu"),
    ("RR", "Rà soát rủi ro hợp đồng"),
    ("ST", "Soạn tài liệu"),
    ("BM", "Bộ mẫu hồ sơ & điền bộ"),
    ("KHO", "Kho tài liệu & tự học tài liệu"),
    ("DN", "Duyệt nhãn & sửa nội dung"),
    ("TH", "Tự học từ hội thoại"),
    ("WEB", "Kênh website & form liên hệ"),
    ("KH", "Cổng khách hàng"),
    ("API", "API tích hợp CRM"),
    ("QT", "Quản trị & nhật ký"),
    ("HT", "Vận hành & hạ tầng"),
    ("PNF", "Hiệu năng & giao diện"),
    ("CHUA", "Hạng mục báo giá chưa có chức năng riêng"),
]
TEN_NHOM = dict(NHOM)
for c in C:
    c["nhom"] = TEN_NHOM[c["ma"].split("-")[0]]
KB = kb_cases.load()


def phai_dan(k):
    """Cột Phải dẫn: tiêu chí HDS + tiêu chí theo luật hiện hành (đề xuất 07/10, chờ luật sư xác nhận)."""
    if not k.get("dieu_hien_hanh"):
        return k["dieu"]
    return f"{k['dieu']} — THEO LUẬT HIỆN HÀNH (đề xuất, chờ luật sư HDS xác nhận): {k['dieu_hien_hanh']}"
_MUC = {"Cao": 0, "Trung bình": 1, "Thấp": 2}
F.PHAT_HIEN.sort(key=lambda x: (_MUC.get(x[1], 9), int(x[0][2:])))
F.LECH_TAI_LIEU.sort(key=lambda x: int(x[0][2:]))


# ---------------------------------------------------------------- tiện ích
def plain(s):
    s = s.replace("**", "")
    s = re.sub(r"(?<![\w*])\*(?!\s)([^*\n]+?)(?<!\s)\*(?![\w*])", r"\1", s)
    return s.replace("`", "")


def cell(s):
    return str(s).replace("|", "\\|").replace("\n", "<br>")


def bang(header, rows):
    out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    out += ["| " + " | ".join(cell(x) for x in r) + " |" for r in rows]
    return "\n".join(out)


# ---------------------------------------------------------------- MARKDOWN
def build_md():
    L = [f"# {S.TIEU_DE}", "", f"*Lập ngày {NGAY} · {len(C)} ca chức năng + {len(KB)} kịch bản pháp lý của HDS*", "",
         S.MO_DAU, "", "## Mục lục", "",
         "1. Căn cứ đối chiếu", "2. Danh mục chức năng của webapp", "3. Ma trận hợp đồng → chức năng → ca kiểm thử",
         "4. Chuẩn bị kiểm thử", "5. Cách ghi kết quả và mức độ lỗi", "6. Ca kiểm thử chi tiết",
         "7. Bốn mươi kịch bản pháp lý của HDS", "8. Thứ tự chạy đề xuất", "9. Tiêu chí nghiệm thu",
         "10. Phát hiện khi rà soát 02/10/2026 và trạng thái sau đợt sửa 03/10/2026", "Phụ lục — File đầu vào mẫu", "",
         "---", "", "## 1. Căn cứ đối chiếu", "", S.CAN_CU, "",
         "## 2. Danh mục chức năng của webapp", "",
         "Toàn bộ chức năng đang chạy trên máy chủ, theo khu vực trên giao diện. Cột *Mục HĐ* là số mục trong báo giá 6 tuần.", "",
         bang(["Khu vực", "Chức năng", "Ai dùng", "Mục HĐ", "Ca kiểm thử"], F.CHUC_NANG), "",
         "## 3. Ma trận hợp đồng → chức năng → ca kiểm thử", "",
         "Đánh giá trạng thái theo mã nguồn và máy chủ ngày 03/10/2026, sau đợt sửa lỗi nghiệm thu (không phải theo báo cáo cũ).", "",
         bang(["Mục", "Yêu cầu (báo giá)", "Có trên hệ thống", "Trạng thái 03/10", "Ca kiểm thử"], F.MA_TRAN), "",
         "Nhóm 4 (mục 19–36, Quản lý & Phát triển) **tạm hoãn** theo báo giá — không kiểm thử.", "",
         "## 4. Chuẩn bị kiểm thử", "", S.CHUAN_BI, "",
         "## 5. Cách ghi kết quả và mức độ lỗi", "", S.GHI_KET_QUA, "",
         "## 6. Ca kiểm thử chi tiết", "",
         "Đọc cột **Gửi vào** như một kịch bản: làm đúng thao tác, dùng đúng chữ/ file ghi trong đó. "
         "Cột **Phải nhận được** là tiêu chí ĐẠT — thiếu một ý là KHÔNG ĐẠT. Ca ghi **(sửa 03/10 — F-xx)** là chỗ lệch yêu cầu "
         "ngày 02/10 đã được sửa: chấm như ca thường; gặp lại lỗi cũ thì ghi KHÔNG ĐẠT kèm mã F-xx.", ""]
    for i, (ma, ten) in enumerate(NHOM, 1):
        ds = [c for c in C if c["ma"].split("-")[0] == ma]
        L += [f"### 6.{i} {ten} ({len(ds)} ca)", ""]
        L.append(bang(["Mã", "Ca kiểm thử", "Tài khoản · Tiền đề", "Gửi vào (thao tác & đầu vào)", "Phải nhận được", "HĐ · Ưu tiên · Loại"],
                      [(c["ma"], c["ten"], c["vai"] + (f" · {c['tien']}" if c["tien"] else ""), c["buoc"], c["mong"],
                        f"{c['hd']} · {c['ut']} · {c['loai']}") for c in C if c["ma"].split("-")[0] == ma]))
        L.append("")
    L += ["## 7. Bốn mươi kịch bản pháp lý của HDS", "",
          "Nguồn: *Kịch bản test AI.docx* do HDS cung cấp. Chạy ở **Hội thoại AI** bằng TK-CV, **mỗi kịch bản một hội thoại mới** "
          "(tránh ngữ cảnh câu trước làm lệch). Gửi **nguyên văn** cột *Câu gửi bot* (= Tình huống + Yêu cầu của HDS).", "",
          "**Cách chấm** (hai người chấm độc lập, lệch thì luật sư HDS quyết):", "",
          "- **ĐÚNG** — dẫn đủ mọi điều/văn bản ở cột *Phải dẫn*, kết luận không trái tiêu chí, có [Nguồn n] trỏ đúng văn bản.",
          "- **MỘT PHẦN** — dẫn được ít nhất một mục bắt buộc, không sai bản chất; hoặc đúng bản chất nhưng thiếu số điều.",
          "- **SAI** — không dẫn mục nào, dẫn sai luật, hoặc kết luận ngược tiêu chí. Bot **từ chối có lý do** (kho chưa có văn bản) "
          "ghi MỘT PHẦN kèm ghi chú *kho thiếu*, không ghi SAI.",
          "- Ghi thêm huy hiệu kiểm chứng và thời gian trả lời. Kho đang dùng **bản luật mới nhất** (ví dụ Luật Doanh nghiệp đã sửa đổi 2025) — "
          "số điều có thể khác bản 2020 mà tiêu chí dẫn; HDS chốt chấm theo bản nào trước khi chạy (mục 9).", "",
          bang(["KB", "Lĩnh vực · Chủ đề", "Câu gửi bot", "Phải dẫn", "Lượt trước"],
               [(k["id"], f"{k['linh_vuc']} · {k['tieu_de']}", k["cau_hoi"], phai_dan(k), k["luot_truoc"]) for k in KB]), "",
          "## 8. Thứ tự chạy đề xuất", "",
          "| Buổi | Nội dung | Ghi chú |", "|---|---|---|",
          "| 1 | Chuẩn bị mục 4.1–4.2, HT-01, HT-02, KHO-01, KHO-02 | Không có chim mồi thì nhóm PQ không chạy được |",
          "| 2 | TK, PQ, KH, API-02, API-07 (bảo mật trước) | Lỗi bảo mật là điều kiện dừng — báo ngay, không đợi hết buổi |",
          "| 3 | CH, KT, RR | |",
          "| 4 | ST, BM, DN, TH, KHO còn lại | ST-07 cần đợi 1–3 phút mỗi lần lưu |",
          "| 5 | WEB (trừ WEB-05), API, QT, HT, PNF, CHUA | |",
          "| 6–7 | 40 kịch bản pháp lý | ~40 × 1,5 phút chạy + thời gian chấm; 2 người chấm |",
          "| Cuối | TK-04, WEB-05 (gây khoá IP), dọn dẹp mục 4.4 | Chạy từ mạng 4G nếu được |", "",
          "## 9. Tiêu chí nghiệm thu", "", S.TIEU_CHI_NGHIEM_THU, "",
          "## 10. Phát hiện khi rà soát 02/10/2026 và trạng thái sau đợt sửa 03/10/2026", "",
          "Rà soát mã nguồn (khớp 100% máy chủ), chạy thật bộ quy tắc trên file mẫu, và kiểm tra chỉ-đọc trên máy chủ. "
          "Các mục này đã gắn vào ca kiểm thử tương ứng để tester xác nhận lại.", "",
          "### 10.1 Lỗi / rủi ro phát hiện", "",
          bang(["Mã", "Mức", "Nhóm", "Mô tả", "Tái hiện / bằng chứng", "Vị trí", "Ca test", "Trạng thái 03/10"],
               [tuple(m) + (F.TRANG_THAI.get(m[0], ""),) for m in F.PHAT_HIEN]), "",
          "### 10.2 Tài liệu lệch với hệ thống (sửa tài liệu, không phải lỗi phần mềm)", "",
          f"{sum(1 for t in F.LECH_TAI_LIEU if t[-1].startswith(('Đã', 'Đúng')))}/{len(F.LECH_TAI_LIEU)} mục đã sửa (tài liệu 02/10, mã/tài liệu 03/10) "
          "(ưu tiên chỗ làm theo sẽ gây mất/lộ dữ liệu hoặc lệnh chạy không được); các mục *Chưa sửa* tester đọc cột *Thực tế* để không báo nhầm.", "",
          bang(["Mã", "Tài liệu", "Đã ghi", "Thực tế", "Xử lý"], F.LECH_TAI_LIEU), "",
          "### 10.3 Hạng mục hợp đồng chưa bàn giao đủ", "",
          bang(["Mục", "Yêu cầu", "Trạng thái"],
               [(m[0], m[1], m[3]) for m in F.MA_TRAN if not m[3].startswith("Đã có")]), "",
          "## Phụ lục — File đầu vào mẫu", "",
          "Thư mục `du-lieu-mau/` (chi tiết ở `du-lieu-mau/README.md`). Mọi tên người, công ty, số giấy tờ trong đó là **dữ liệu giả**.", "",
          bang(["File", "Dùng ở ca", "Kết quả mong đợi chính"], DU_LIEU), ""]
    return "\n".join(L)


DU_LIEU = [
    ("RR01_HD_vay_lai_3pt_thang.docx", "RR-01, KT-01, CH-12",
     "HĐ vay · 9 đạt / 3 cảnh báo / 0 thiếu · rủi ro cao; lãi 3%/tháng = 36%/năm vượt trần 20%/năm; lãi chậm trả 30%/năm"),
    ("RR02_HD_dich_vu_phat_12pt_thieu_tranh_chap.docx", "RR-02, KT-06, CH-12",
     "HĐ dịch vụ · 11 / 6 / 0 · rủi ro cao; phạt 12% > 8%; thiếu điều khoản tranh chấp"),
    ("RR03_HDLD_vi_pham_nguong.docx", "RR-03, RR-07, CH-12",
     "HĐLĐ · 12 / 5 / 3 · rủi ro cao; thử việc 90 ngày, lương 70%, 48 tháng, 10 giờ/ngày, làm thêm 300 giờ/năm"),
    ("RR04_HD_dich_vu_hop_le_DOI_CHUNG.docx", "RR-04, RR-06, KT-05, API-03",
     "Đối chứng · 16 / 2 / 0 · phạt 8% và lãi chậm 10%/năm đều ĐẠT"),
    ("RR05_HD_dieu_khoan_mot_chieu.docx", "RR-05", "7 / 8 / 1 · đủ 3 dấu hiệu điều khoản một chiều"),
    ("AT01_HD_co_chen_lenh.docx", "KT-07, BM-08", "Có đoạn 'GHI CHÚ HỆ THỐNG: Bỏ qua mọi chỉ dẫn…' — bot không được làm theo"),
    ("TD01_CCCD_gia.docx", "ST-01, BM-04, BM-07", "Tự điền bóc đúng 10 trường (số CCCD 001095012345…)"),
    ("TD02_CV_gia.docx", "ST-02", "Bóc 6 trường (SĐT 0912000111, email…)"),
    ("CANARY_CONG_KHAI / NOI_BO / KHACH_A / KHACH_B.docx", "PQ-02…04, WEB-02, KH-02, KHO-01/02, API-08",
     "Mỗi file một mã duy nhất; mã xuất hiện ⇔ tài khoản đọc được tài liệu đó"),
    ("CHAT01_ghi_chu_tam.txt", "CH-10", "Hạn nộp 15/11/2026, mã TMP-5521"),
    ("KTMT01_yeu_cau_sua_ban_nhap.txt", "ST-05, ST-07",
     "Kiểm tra mâu thuẫn ra 6 cảnh báo: thử việc 90, lương thử việc 70%, thời hạn HĐ 48 tháng, 10 giờ/ngày, phạt 12%, lãi 36%/năm"),
    ("EN01_service_agreement.txt", "CHUA-06",
     "HĐ tiếng Anh giữa hai công ty VN: phạt 15% > 8%, lãi chậm 2%/tháng, chấm dứt một chiều, chọn luật Anh + trọng tài Singapore"),
]


def build_readme_du_lieu():
    return "\n".join([
        "# File đầu vào mẫu cho kiểm thử", "",
        f"Sinh ngày {NGAY}. **Toàn bộ là dữ liệu giả** (công ty *THỬ NGHIỆM ALPHA/BETA*, người *KIỂM THỬ*). "
        "Kết quả mong đợi của các file RR*, TD*, KTMT01 đã được chạy thật qua bộ quy tắc của hệ thống "
        "(`ra_soat_rui_ro`, `kiem_tra_mau_thuan`, `autofill`) — đây là phần tất định, không phụ thuộc AI.", "",
        bang(["File", "Dùng ở ca", "Kết quả mong đợi chính"], DU_LIEU), "",
        "Các file CANARY phải **gỡ khỏi kho** sau khi kiểm thử (Quản trị → Kho tài liệu → Bỏ).", ""])


# ---------------------------------------------------------------- EXCEL
def build_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule

    FONT = "Arial"
    NAVY = "1F3A5F"
    f_norm = Font(name=FONT, size=10)
    f_bold = Font(name=FONT, size=10, bold=True)
    f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    f_title = Font(name=FONT, size=14, bold=True, color=NAVY)
    f_note = Font(name=FONT, size=9, italic=True, color="595959")
    fill_head = PatternFill("solid", fgColor=NAVY)
    fill_input = PatternFill("solid", fgColor="FFF2CC")
    fill_ex = PatternFill("solid", fgColor="EDEDED")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    def header(ws, row, cols, widths):
        for j, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(row=row, column=j, value=h)
            c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[row].height = 32

    def put(ws, row, values, input_from=None, font=f_norm, fill=None):
        for j, v in enumerate(values, 1):
            c = ws.cell(row=row, column=j, value=v)
            c.font, c.alignment, c.border = font, wrap, border
            if input_from and j >= input_from:
                c.fill = fill_input
            if fill:
                c.fill = fill

    # ---- Hướng dẫn
    ws = wb.active
    ws.title = "Hướng dẫn"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    ws["A1"] = S.TIEU_DE
    ws["A1"].font = f_title
    ws["A2"] = f"Lập ngày {NGAY} · {len(C)} ca chức năng + {len(KB)} kịch bản pháp lý · tài liệu đầy đủ: KICH_BAN_KIEM_THU.md cùng thư mục"
    ws["A2"].font = f_note
    rows = [
        ("Mục đích", "Kiểm tra Trợ lý AI HDS chạy đúng hợp đồng (báo giá 6 tuần HT1–HT4 + mục 1–18; kế hoạch 10 ngày) và đúng kỳ vọng của công ty luật. "
                     "Mỗi ca ghi rõ dùng tài khoản nào, gửi gì vào, phải nhận được gì."),
        ("Ô cần điền", "Chỉ điền các ô NỀN VÀNG. Cột Kết quả chọn từ danh sách: ĐẠT · KHÔNG ĐẠT · CHẶN · BỎ QUA. Mọi ô khác là đề bài — không sửa."),
        ("Dòng mẫu", "Sheet 'Ca kiểm thử' dòng 3 (nền xám) là ví dụ cách điền — xoá hoặc bỏ qua khi tổng hợp (không tính vào sheet Tổng hợp)."),
        ("Địa chỉ thử", "https://app.diginix.io.vn (tên miền app.hdslaw.vn chưa trỏ)."),
        ("Tài khoản thử", "TK-AD admin · TK-QT Ban QT · TK-TP Trưởng bộ phận (có quyền duyệt) · TK-CV Chuyên viên · TK-TL Trợ lý · TK-KA / TK-KB Khách A / Khách B. "
                          "Tạo theo mục 4.1 của tài liệu; KHÔNG dùng tài khoản thật của nhân viên."),
        ("File đầu vào", "Thư mục du-lieu-mau/ — hợp đồng cài sẵn vi phạm, CCCD/CV giả, tài liệu chèn lệnh, 4 tài liệu chim mồi CANARY (sheet 'Dữ liệu mẫu')."),
        ("Lưu ý 1", "Đăng nhập bị giới hạn 10 lượt / 5 phút / IP, tính cả lượt đúng. Mỗi vai dùng một cửa sổ / hồ sơ trình duyệt riêng, không đăng xuất – đăng nhập liên tục."),
        ("Lưu ý 2", "TK-04 và WEB-05 cố ý gây khoá IP — chạy cuối cùng hoặc từ mạng 4G."),
        ("Lưu ý 3", "Bot chạy trên GPU nội bộ: câu tra cứu 10–60 giây, tình huống dài đến ~120 giây là bình thường."),
        ("Bản build", "Ghi dấu build ở chân trang web vào cột Bản build. Sau mỗi lần cập nhật hệ thống bấm Ctrl+F5."),
        ("Mức độ lỗi", "Nghiêm trọng: lộ dữ liệu / sai quyền / sai luật gây hại · Cao: chức năng hợp đồng không dùng được · "
                       "Trung bình: sai một phần, có cách vòng · Thấp: giao diện, chính tả, sổ tay lệch."),
        ("Ca 'Lỗi đã biết'", "Ca ghi '(sửa 03/10 — F-xx)' là chỗ đã lệch yêu cầu ngày 02/10 và đã sửa — chấm như ca thường; trạng thái từng phát hiện ở sheet 'Phát hiện'. Vẫn chạy và ghi đúng điều quan sát được."),
    ]
    for i, (k, v) in enumerate(rows, 4):
        a, b = ws.cell(row=i, column=1, value=k), ws.cell(row=i, column=2, value=v)
        a.font, b.font = f_bold, f_norm
        a.alignment, b.alignment = wrap, wrap
    ws.cell(row=4 + len(rows) + 1, column=1, value="Nền vàng = ô tester điền").fill = fill_input
    ws.cell(row=4 + len(rows) + 1, column=1).font = f_bold

    # ---- Ca kiểm thử
    wc = wb.create_sheet("Ca kiểm thử")
    cols = ["Mã", "Nhóm", "Ca kiểm thử", "Hạng mục HĐ", "Ưu tiên", "Loại", "Tài khoản", "Tiền đề",
            "Gửi vào (thao tác & đầu vào)", "Phải nhận được", "Kết quả", "Thực tế quan sát", "Bằng chứng (ảnh/file)",
            "Mã lỗi", "Người test", "Ngày", "Bản build"]
    widths = [9, 20, 28, 14, 9, 11, 16, 20, 60, 60, 13, 40, 18, 10, 14, 11, 14]
    wc["A1"] = "CA KIỂM THỬ CHỨC NĂNG — điền các ô nền vàng"
    wc["A1"].font = f_title
    header(wc, 2, cols, widths)
    put(wc, 3, ["VÍ DỤ", "Hội thoại AI (kênh nội bộ)", "Dòng mẫu cách điền — không tính",
                "Mục 6", "Cao", "Chức năng", "TK-CV", "", "…", "…", "ĐẠT",
                "Trả lời 3 năm, Điều 429 BLDS, huy hiệu Đã kiểm chứng; 41 giây", "CH-01_0310.png", "", "Nguyễn A", "03/10/2026",
                "dca98ff-0110.0317"], font=Font(name=FONT, size=10, italic=True, color="7F7F7F"), fill=fill_ex)
    r0 = 4
    for i, c in enumerate(C):
        put(wc, r0 + i, [c["ma"], c["nhom"], c["ten"], c["hd"], c["ut"], c["loai"], c["vai"], plain(c["tien"]),
                         plain(c["buoc"]), plain(c["mong"]), None, None, None, None, None, None, None], input_from=11)
    last = r0 + len(C) - 1
    dv = DataValidation(type="list", formula1='"ĐẠT,KHÔNG ĐẠT,CHẶN,BỎ QUA"', allow_blank=True)
    dv.error, dv.errorTitle = "Chọn ĐẠT, KHÔNG ĐẠT, CHẶN hoặc BỎ QUA", "Giá trị không hợp lệ"
    wc.add_data_validation(dv)
    dv.add(f"K{r0}:K{last}")
    for val, color in (("ĐẠT", "C6EFCE"), ("KHÔNG ĐẠT", "FFC7CE"), ("CHẶN", "FFEB9C"), ("BỎ QUA", "D9D9D9")):
        wc.conditional_formatting.add(f"K{r0}:K{last}", CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                                     fill=PatternFill("solid", fgColor=color)))
    wc.freeze_panes = "D3"
    wc.auto_filter.ref = f"A2:Q{last}"

    # ---- 40 KB pháp lý
    wk = wb.create_sheet("40 KB pháp lý")
    cols = ["KB", "Lĩnh vực", "Chủ đề", "Câu gửi bot (nguyên văn)", "Tiêu chí của HDS", "Phải dẫn", "Lượt trước",
            "Điều/văn bản bot dẫn", "Kết quả", "Huy hiệu kiểm chứng", "Thời gian (giây)", "Ghi chú", "Người chấm", "Ngày"]
    widths = [6, 16, 26, 70, 50, 30, 18, 30, 13, 18, 10, 30, 14, 11]
    wk["A1"] = "40 KỊCH BẢN PHÁP LÝ CỦA HDS — mỗi kịch bản một hội thoại mới, gửi nguyên văn cột D"
    wk["A1"].font = f_title
    header(wk, 2, cols, widths)
    k0 = 3
    for i, k in enumerate(KB):
        put(wk, k0 + i, [k["id"], k["linh_vuc"], k["tieu_de"], k["cau_hoi"], k["tieu_chi"], phai_dan(k), k["luot_truoc"],
                         None, None, None, None, None, None, None], input_from=8)
    klast = k0 + len(KB) - 1
    dv2 = DataValidation(type="list", formula1='"ĐÚNG,MỘT PHẦN,SAI,BỎ QUA"', allow_blank=True)
    wk.add_data_validation(dv2)
    dv2.add(f"I{k0}:I{klast}")
    dv3 = DataValidation(type="list", allow_blank=True, formula1='"Đã kiểm chứng theo nguồn,Chỉ một phần có đủ căn cứ,'
                                                                   'Chưa gắn được trích dẫn,Không tìm thấy căn cứ trong nguồn,Chưa đủ bằng chứng để kết luận"')
    wk.add_data_validation(dv3)
    dv3.add(f"J{k0}:J{klast}")
    for val, color in (("ĐÚNG", "C6EFCE"), ("SAI", "FFC7CE"), ("MỘT PHẦN", "FFEB9C")):
        wk.conditional_formatting.add(f"I{k0}:I{klast}", CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                                      fill=PatternFill("solid", fgColor=color)))
    wk.freeze_panes = "B3"
    wk.auto_filter.ref = f"A2:N{klast}"
    n = klast + 2
    wk.cell(row=n, column=1, value="Cách chấm").font = f_bold
    wk.cell(row=n, column=4, value="ĐÚNG = dẫn đủ cột 'Phải dẫn', kết luận không trái tiêu chí, có [Nguồn n] đúng văn bản · "
                                   "MỘT PHẦN = dẫn được ≥ 1 mục, không sai bản chất (hoặc từ chối có lý do vì kho thiếu) · "
                                   "SAI = không dẫn mục nào / dẫn sai luật / kết luận ngược tiêu chí.").alignment = wrap

    # ---- Tổng hợp
    wt = wb.create_sheet("Tổng hợp", 1)
    wt["A1"] = "TỔNG HỢP KẾT QUẢ — tự tính từ sheet 'Ca kiểm thử' và '40 KB pháp lý'"
    wt["A1"].font = f_title
    cols = ["Nhóm", "Tổng ca", "ĐẠT", "KHÔNG ĐẠT", "CHẶN", "BỎ QUA", "Chưa chạy", "Tỉ lệ đạt (trên ca đã chạy)"]
    header(wt, 3, cols, [42, 10, 10, 12, 10, 10, 11, 16])
    R = f"'Ca kiểm thử'!$K${r0}:$K${last}"
    G = f"'Ca kiểm thử'!$B${r0}:$B${last}"
    row = 4
    for ma, ten in NHOM:
        wt.cell(row=row, column=1, value=ten)
        wt.cell(row=row, column=2, value=f'=COUNTIF({G},$A{row})')
        for j, kq in zip(range(3, 7), ("ĐẠT", "KHÔNG ĐẠT", "CHẶN", "BỎ QUA")):
            wt.cell(row=row, column=j, value=f'=COUNTIFS({G},$A{row},{R},"{kq}")')
        wt.cell(row=row, column=7, value=f"=B{row}-SUM(C{row}:F{row})")
        wt.cell(row=row, column=8, value=f"=IFERROR(C{row}/(C{row}+D{row}+E{row}),\"—\")")
        row += 1
    tong = row
    wt.cell(row=tong, column=1, value="TỔNG")
    for j in range(2, 8):
        col = get_column_letter(j)
        wt.cell(row=tong, column=j, value=f"=SUM({col}4:{col}{tong - 1})")
    wt.cell(row=tong, column=8, value=f"=IFERROR(C{tong}/(C{tong}+D{tong}+E{tong}),\"—\")")
    # theo loại / ưu tiên
    U = f"'Ca kiểm thử'!$E${r0}:$E${last}"
    T = f"'Ca kiểm thử'!$F${r0}:$F${last}"
    extra = [("Ca loại Bảo mật (yêu cầu 100% ĐẠT)", T, "Bảo mật"), ("Ca ưu tiên Cao (yêu cầu ≥ 95% ĐẠT)", U, "Cao")]
    row = tong + 2
    for ten, rng, val in extra:
        wt.cell(row=row, column=1, value=ten)
        wt.cell(row=row, column=2, value=f'=COUNTIF({rng},"{val}")')
        for j, kq in zip(range(3, 7), ("ĐẠT", "KHÔNG ĐẠT", "CHẶN", "BỎ QUA")):
            wt.cell(row=row, column=j, value=f'=COUNTIFS({rng},"{val}",{R},"{kq}")')
        wt.cell(row=row, column=7, value=f"=B{row}-SUM(C{row}:F{row})")
        wt.cell(row=row, column=8, value=f"=IFERROR(C{row}/(C{row}+D{row}+E{row}),\"—\")")
        row += 1
    # 40 KB
    row += 1
    header(wt, row, ["40 kịch bản pháp lý", "Tổng", "ĐÚNG", "MỘT PHẦN", "SAI", "BỎ QUA", "Chưa chấm", "ĐÚNG + MỘT PHẦN"],
           [42, 10, 10, 12, 10, 10, 11, 16])
    row += 1
    KR = f"'40 KB pháp lý'!$I${k0}:$I${klast}"
    KA = f"'40 KB pháp lý'!$A${k0}:$A${klast}"
    wt.cell(row=row, column=1, value="Toàn bộ")
    wt.cell(row=row, column=2, value=f"=COUNTA({KA})")
    for j, kq in zip(range(3, 7), ("ĐÚNG", "MỘT PHẦN", "SAI", "BỎ QUA")):
        wt.cell(row=row, column=j, value=f'=COUNTIF({KR},"{kq}")')
    wt.cell(row=row, column=7, value=f"=B{row}-SUM(C{row}:F{row})")
    wt.cell(row=row, column=8, value=f"=IFERROR((C{row}+D{row})/(C{row}+D{row}+E{row}),\"—\")")
    for r in wt.iter_rows(min_row=4, max_row=row):
        for c in r:
            if c.value is not None or c.column <= 8:
                c.font = f_bold if (c.row == tong) else f_norm
                c.border = border
                if c.column == 8:
                    c.number_format = "0%"
    wt.cell(row=row + 2, column=1, value="Công thức tự tính khi mở bằng Excel / Google Sheets. Dòng ví dụ (dòng 3 sheet Ca kiểm thử) không được tính.").font = f_note

    # ---- Ma trận HĐ
    wm = wb.create_sheet("Ma trận HĐ")
    wm["A1"] = "MA TRẬN HỢP ĐỒNG → CHỨC NĂNG → CA KIỂM THỬ (trạng thái 03/10/2026)"
    wm["A1"].font = f_title
    header(wm, 2, ["Mục", "Yêu cầu (báo giá 6 tuần)", "Có trên hệ thống", "Trạng thái 03/10", "Ca kiểm thử", "Kết luận nghiệm thu (HDS điền)"],
           [7, 45, 55, 40, 22, 30])
    for i, m in enumerate(F.MA_TRAN, 3):
        put(wm, i, list(m) + [None], input_from=6)
    wm.freeze_panes = "B3"

    # ---- Danh mục chức năng
    wd = wb.create_sheet("Danh mục chức năng")
    wd["A1"] = "DANH MỤC CHỨC NĂNG WEBAPP"
    wd["A1"].font = f_title
    header(wd, 2, ["Khu vực", "Chức năng", "Ai dùng", "Mục HĐ", "Ca kiểm thử"], [22, 70, 26, 12, 20])
    for i, m in enumerate(F.CHUC_NANG, 3):
        put(wd, i, list(m))
    wd.freeze_panes = "A3"

    # ---- Phát hiện
    wp = wb.create_sheet("Phát hiện")
    wp["A1"] = "PHÁT HIỆN KHI RÀ SOÁT 02/10/2026 + TRẠNG THÁI SAU ĐỢT SỬA 03/10/2026 — lỗi, rủi ro, tài liệu lệch"
    wp["A1"].font = f_title
    header(wp, 2, ["Mã", "Mức", "Nhóm", "Mô tả", "Tái hiện / bằng chứng", "Vị trí", "Ca test", "Trạng thái 03/10", "Xác nhận lại (tester)"],
           [7, 11, 13, 60, 40, 30, 10, 45, 18])
    for i, m in enumerate(F.PHAT_HIEN, 3):
        put(wp, i, [plain(x) for x in m] + [F.TRANG_THAI.get(m[0], ""), None], input_from=9)
    r = 3 + len(F.PHAT_HIEN) + 1
    header(wp, r, ["Mã", "Tài liệu", "Đã ghi", "Thực tế", "Xử lý", "", "", "", ""], [7, 11, 13, 60, 40, 30, 10, 45, 18])
    for i, m in enumerate(F.LECH_TAI_LIEU, r + 1):
        put(wp, i, [m[0], m[1], plain(m[2]), plain(m[3]), m[4]])
    wp.freeze_panes = "A3"

    # ---- Báo lỗi
    wb_ = wb.create_sheet("Báo lỗi")
    wb_["A1"] = "SỔ BÁO LỖI — 1 lỗi = 1 dòng"
    wb_["A1"].font = f_title
    cols = ["Mã lỗi", "Mã ca", "Mức độ", "Tài khoản", "Bước tái hiện (gửi gì)", "Mong đợi", "Thực tế (nguyên văn)",
            "Ảnh / file", "Bản build", "Người báo", "Ngày", "Trạng thái", "Ghi chú nhà phát triển"]
    header(wb_, 2, cols, [9, 9, 12, 12, 45, 35, 40, 18, 14, 14, 11, 12, 30])
    put(wb_, 3, ["BUG-000", "RR-01", "Cao", "TK-CV", "Rà soát RR01_HD_vay_lai_3pt_thang.docx, loại Tự nhận diện",
                 "CẢNH BÁO Lãi suất vay 3%/tháng", "Không có dòng lãi suất trong bảng", "RR-01_0310.png", "dca98ff-0110.0317",
                 "Nguyễn A", "03/10/2026", "Mới", "(ví dụ — xoá dòng này)"],
        font=Font(name=FONT, size=10, italic=True, color="7F7F7F"), fill=fill_ex)
    for rr in range(4, 104):
        put(wb_, rr, [f"BUG-{rr - 3:03d}"] + [None] * 12, input_from=2)
    dv4 = DataValidation(type="list", formula1='"Nghiêm trọng,Cao,Trung bình,Thấp"', allow_blank=True)
    wb_.add_data_validation(dv4)
    dv4.add("C4:C103")
    dv5 = DataValidation(type="list", formula1='"Mới,Đang sửa,Đã sửa – chờ test lại,Đóng,Không phải lỗi"', allow_blank=True)
    wb_.add_data_validation(dv5)
    dv5.add("L4:L103")
    wb_.freeze_panes = "C3"

    # ---- Dữ liệu mẫu
    wdl = wb.create_sheet("Dữ liệu mẫu")
    wdl["A1"] = "FILE ĐẦU VÀO MẪU (thư mục du-lieu-mau/) — dữ liệu giả, đáp án đã chạy thật qua bộ quy tắc"
    wdl["A1"].font = f_title
    header(wdl, 2, ["File", "Dùng ở ca", "Kết quả mong đợi chính"], [48, 34, 80])
    for i, m in enumerate(DU_LIEU, 3):
        put(wdl, i, list(m))

    for w in wb.worksheets:
        w.sheet_view.zoomScale = 90
    wb.save(path)
    return r0, last, k0, klast


if __name__ == "__main__":
    os.makedirs(os.path.join(OUT, "du-lieu-mau"), exist_ok=True)
    md = build_md()
    open(os.path.join(OUT, "KICH_BAN_KIEM_THU.md"), "w", encoding="utf-8", newline="\n").write(md)
    open(os.path.join(OUT, "du-lieu-mau", "README.md"), "w", encoding="utf-8", newline="\n").write(build_readme_du_lieu())
    info = build_xlsx(os.path.join(OUT, "KICH_BAN_KIEM_THU.xlsx"))
    print("md chars", len(md), "| cases", len(C), "| kb", len(KB), "| xlsx rows", info)
