# -*- coding: utf-8 -*-
"""Sinh KET_QUA_CAU_HOI_MAU_<ngày>.md + .xlsx từ lượt chạy THẬT bộ câu hỏi mẫu
(tu_dong/kt_cau_hoi_mau.py → ket_qua_cau_hoi_mau_0510.json) + bảng chấm
cham_cau_hoi_mau_0510.json ({mã: {"kq": Đúng|Một phần|Sai|Bịa|Không chạy, "ghi": …},
"_tong_ket": [...], "_luot": "…"}).

    python deploy/kiem-thu/nguon/build_ket_qua_cau_hoi_mau.py deploy/kiem-thu
"""
import json
import os
import re
import statistics
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import cau_hoi_mau as Q  # noqa: E402
import kb_cases  # noqa: E402
from build_docs import bang, plain  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(HERE)
KQ = json.load(open(os.path.join(HERE, "ket_qua_cau_hoi_mau_0510.json"), encoding="utf-8"))
CH = json.load(open(os.path.join(HERE, "cham_cau_hoi_mau_0510.json"), encoding="utf-8"))
NGAY_FILE = "06-10-2026"   # ngày chạy (bộ câu lập 05/10)
MUC = ["Đúng", "Một phần", "Sai", "Bịa", "Không chạy"]
TEN_VN = {m: t for m, t, _b in Q.NHOM_VN}
TEN_EN = {m: t for m, t, _g in Q.NHOM_EN}
KB = {"KB" + k["id"]: k for k in kb_cases.load() if k["id"] not in Q.KB_DA_DUNG}


def dong_cau():
    """[(bộ, mã, nhóm, câu, đạt khi, căn cứ/ghi chú)] theo đúng thứ tự bộ câu."""
    ra = []
    for r in Q.VN:
        ra.append(("VN", r[0], TEN_VN[r[1]], r[4], r[5], r[6]))
    for ma, k in KB.items():
        ra.append(("KB", ma, f"{k['linh_vuc']} · {k['tieu_de']}", k["cau_hoi"], k["tieu_chi"], k["dieu"]))
    for r in Q.EN:
        ra.append(("EN", r[0], TEN_EN[r[1]], r[3], r[4], r[5]))
    return ra


DONG = dong_cau()


def ket(ma):
    return (CH.get(ma) or {}).get("kq") or "Chưa chấm"


def ghi(ma):
    return (CH.get(ma) or {}).get("ghi") or ""


def tra_loi(ma):
    return ((KQ["cau"].get(ma) or {}).get("answer") or "").strip()


def giay(ma):
    return (KQ["cau"].get(ma) or {}).get("giay")


def nguon(ma, n=6):
    ds = (KQ["cau"].get(ma) or {}).get("nguon") or []
    return "; ".join(f"[{s.get('n')}] {(s.get('title') or s.get('attachment_name') or '')[:60]}"
                     + (f" — {s['section_title'][:40]}" if s.get("section_title") else "") for s in ds[:n])


def dem(bo=None):
    c = {m: 0 for m in MUC}
    for b, ma, *_ in DONG:
        if bo and b != bo:
            continue
        k = ket(ma)
        c[k] = c.get(k, 0) + 1
    return c


def thoi_gian():
    ds = [g for g in (giay(ma) for _b, ma, *_ in DONG) if isinstance(g, (int, float)) and g > 0]
    return (statistics.median(ds), max(ds)) if ds else (0, 0)


def build_md():
    tv, tmax = thoi_gian()
    L = [f"# KẾT QUẢ CHẠY THẬT BỘ CÂU HỎI MẪU 05/10 — chạy {NGAY_FILE.replace('-', '/')}", "",
         f"*{CH.get('_luot', '')}*", "",
         "Mỗi câu được gửi qua **đúng API của web** (như nhân viên gõ), máy chủ thật, model qwen3:14b, sau đợt sửa "
         "P1–P8 ngày 05/10. Câu trả lời do Claude chấm theo cột *Đạt khi* của bộ câu (đã đối chiếu kho). "
         "**Luật sư vẫn nên chấm lại** các câu Một phần / Sai trong phiếu xlsx (cột nền vàng).", "",
         "## 1. Tổng hợp", ""]
    rows = []
    for bo, ten in (("VN", "Câu hỏi mẫu tiếng Việt"), ("KB", "24 tình huống dài (KB)"), ("EN", "Tài liệu tiếng Anh (SEC)")):
        c = dem(bo)
        tong = sum(c.values())
        rows.append((ten, tong, *[c.get(m, 0) for m in MUC]))
    c = dem()
    rows.append(("**Tổng**", sum(c.values()), *[c.get(m, 0) for m in MUC]))
    L += [bang(["Bộ", "Số câu", *MUC], rows), "",
          f"Thời gian trả lời: trung vị {tv:.0f} giây, lâu nhất {tmax:.0f} giây.", ""]
    if CH.get("_tong_ket"):
        L += ["## 2. Điều rút ra", ""] + [f"- {x}" for x in CH["_tong_ket"]] + [""]
    L += ["## 3. Từng câu", ""]
    for bo, ten in (("VN", "Câu hỏi mẫu tiếng Việt"), ("KB", "Tình huống dài"), ("EN", "Tài liệu tiếng Anh")):
        L += [f"### {ten}", "",
              bang(["Mã", "Nhóm", "Kết quả", "Nhận xét", "Giây"],
                   [(ma, nhom[:50], ket(ma), ghi(ma), giay(ma) or "—") for b, ma, nhom, *_ in DONG if b == bo]), ""]
    L += ["## 4. Toàn văn câu trả lời", "", "Bản đầy đủ (kèm danh sách nguồn) ở sheet từng bộ trong "
          f"`KET_QUA_CAU_HOI_MAU_{NGAY_FILE}.xlsx`.", ""]
    for b, ma, nhom, cau, dat, _cc in DONG:
        L += [f"#### {ma} — {ket(ma)}", "", f"> {plain(cau)}", "", f"**Đạt khi:** {plain(dat)}", ""]
        if ghi(ma):
            L += [f"**Nhận xét:** {ghi(ma)}", ""]
        L += ["```text", tra_loi(ma) or "(không có câu trả lời)", "```", ""]
        if nguon(ma):
            L += [f"*Nguồn:* {nguon(ma)}", ""]
    return "\n".join(L)


def build_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    fb, fn, ft = Font(bold=True), Font(size=10), Font(bold=True, size=13)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="BBBBBB")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="DDE7F3")
    vang = PatternFill("solid", fgColor="FFF4C2")
    mau = {"Đúng": "D9F2D9", "Một phần": "FFF0C8", "Sai": "F8D0D0", "Bịa": "E8B0B0", "Không chạy": "E6E6E6"}

    ws = wb.active
    ws.title = "Tổng hợp"
    ws["A1"] = f"KẾT QUẢ CHẠY THẬT BỘ CÂU HỎI MẪU 05/10 — chạy {NGAY_FILE.replace('-', '/')}"
    ws["A1"].font = ft
    ws["A2"] = CH.get("_luot", "")
    ws["A2"].font = Font(italic=True, size=10)
    hdr = ["Bộ", "Số câu", *MUC]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font, c.fill, c.border = fb, head, bd
    r = 5
    for bo, ten in (("VN", "Câu hỏi mẫu tiếng Việt"), ("KB", "24 tình huống dài"), ("EN", "Tài liệu tiếng Anh (SEC)"), (None, "Tổng")):
        cc = dem(bo)
        for j, v in enumerate((ten, sum(cc.values()), *[cc.get(m, 0) for m in MUC]), 1):
            c = ws.cell(row=r, column=j, value=v)
            c.border, c.font = bd, (fb if bo is None else fn)
        r += 1
    r += 1
    for x in CH.get("_tong_ket") or []:
        c = ws.cell(row=r, column=1, value="• " + x)
        c.alignment, c.font = wrap, fn
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 46
        r += 1
    for col, w in zip("ABCDEFG", (34, 10, 10, 10, 10, 10, 12)):
        ws.column_dimensions[col].width = w

    cols = ["Mã", "Nhóm", "Câu gửi bot", "Đạt khi", "Kết quả (Claude chấm)", "Nhận xét", "Luật sư chấm lại",
            "Ghi chú luật sư", "Câu trả lời của bot (toàn văn)", "Nguồn bot dùng", "Giây"]
    rong = [8, 22, 44, 50, 13, 50, 13, 30, 90, 50, 7]
    dv_ket = '"Đúng,Một phần,Sai,Bịa"'
    for bo, ten in (("VN", "Câu hỏi mẫu (VN)"), ("KB", "Tình huống dài (KB)"), ("EN", "Tài liệu tiếng Anh")):
        w = wb.create_sheet(ten)
        for j, (h, wd) in enumerate(zip(cols, rong), 1):
            c = w.cell(row=1, column=j, value=h)
            c.font, c.fill, c.border, c.alignment = fb, head, bd, wrap
            w.column_dimensions[c.column_letter].width = wd
        dv = DataValidation(type="list", formula1=dv_ket, allow_blank=True)
        w.add_data_validation(dv)
        i = 2
        for b, ma, nhom, cau, dat, _cc in DONG:
            if b != bo:
                continue
            vals = [ma, nhom, cau, plain(dat), ket(ma), ghi(ma), "", "", tra_loi(ma)[:32000],
                    nguon(ma, 12), giay(ma)]
            for j, v in enumerate(vals, 1):
                c = w.cell(row=i, column=j, value=v)
                c.font, c.alignment, c.border = fn, wrap, bd
            w.cell(row=i, column=5).fill = PatternFill("solid", fgColor=mau.get(ket(ma), "FFFFFF"))
            for j in (7, 8):
                w.cell(row=i, column=j).fill = vang
            dv.add(f"G{i}")
            w.row_dimensions[i].height = 180
            i += 1
        w.freeze_panes = "C2"
        w.auto_filter.ref = f"A1:K{i - 1}"
    for w in wb.worksheets:
        w.sheet_view.zoomScale = 90
    wb.save(path)


if __name__ == "__main__":
    md = os.path.join(OUT, f"KET_QUA_CAU_HOI_MAU_{NGAY_FILE}.md")
    with open(md, "w", encoding="utf-8", newline="\n") as f:
        f.write(build_md() + "\n")
    build_xlsx(os.path.join(OUT, f"KET_QUA_CAU_HOI_MAU_{NGAY_FILE}.xlsx"))
    print("đã sinh", md)
