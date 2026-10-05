# -*- coding: utf-8 -*-
"""Sinh KIEM_THU_CHO_NHAN_VIEN.md + PHIEU_KIEM_THU_NHAN_VIEN.xlsx từ nhan_vien.py.
Chạy: python deploy/kiem-thu/nguon/build_nhan_vien.py deploy/kiem-thu"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import nhan_vien as N  # noqa: E402
import kb_cases  # noqa: E402
from build_docs import bang, plain, NGAY  # noqa: E402

OUT = sys.argv[1]
TEN_NHOM = dict(N.NHOM)
_GOI_Y = []
for _p in N.PHAN_CONG:
    _GOI_Y += [x.strip() for x in _p[2].split(",") if x.strip() not in ("—", "")]
KB_B4 = [k for k in kb_cases.load() if k["id"] in _GOI_Y]


def build_md():
    L = ["# DÙNG THỬ & KIỂM TRA TRỢ LÝ AI — DÀNH CHO NHÂN VIÊN HDS", "",
         f"*Lập ngày {NGAY} · {len(N.BAI)} bài · mỗi người 60–90 phút*", "",
         "Tài liệu này để **luật sư, chuyên viên, trợ lý** tự kiểm tra trợ lý AI có làm đúng việc của công ty luật không, "
         "ở ba việc dùng hằng ngày: **truy vấn** (hỏi luật, hỏi hồ sơ), **tạo tài liệu** (soạn đơn, hợp đồng, điền mẫu) và "
         "**kiểm tra pháp lý** (soi hợp đồng khách gửi). Mỗi bài ghi rõ **gõ gì / đính kèm gì** và **thế nào là đạt**. "
         "Bộ kiểm thử kỹ thuật đầy đủ cho tester nằm ở [KICH_BAN_KIEM_THU.md](KICH_BAN_KIEM_THU.md).", "",
         "> **Nguyên tắc:** AI là công cụ tra cứu và soạn thảo — luật sư ký tên là người chịu trách nhiệm. "
         "Mục đích buổi thử là tìm chỗ bot **sai** và **bịa** trước khi dùng cho khách thật, không phải để bot \"qua bài\".", "",
         "## 1. Chuẩn bị", "",
         "- Tài khoản công ty của bạn (đổi mật khẩu tạm ở lần đăng nhập đầu — bài A1).",
         "- Thư mục **du-lieu-mau** (IT gửi kèm): hợp đồng mẫu cài sẵn lỗi `RR01…RR05`, CCCD/CV giả `TD01`, `TD02`, tài liệu thử `AT01`, `CHAT01`, `KTMT01`. "
         "Toàn bộ là dữ liệu giả — không phải khách thật.",
         "- Phiếu **PHIEU_KIEM_THU_NHAN_VIEN.xlsx**: điền họ tên, phòng, và kết quả từng bài (ô nền vàng).",
         "- Bài **E** dùng việc thật của bạn: **xoá tên khách, MST, số CCCD, số tài khoản** trước khi đưa vào.", "",
         "## 2. Cách chấm", "",
         bang(["Kết quả", "Khi nào"], N.CHAM), "",
         "Với mỗi câu trả lời, nhìn theo thứ tự: **huy hiệu kiểm chứng** (🟢 / 🟡 / 🔴) → **bấm một chip [Nguồn n] → Xem trước** "
         "xem đúng văn bản, đúng Điều chưa → rồi mới đọc nội dung. Ghi thêm **thời gian** (bấm đồng hồ cạnh câu trả lời) nếu chậm hơn 2 phút.", "",
         "## 3. Phân công theo phòng", "",
         "Ai cũng làm nhóm A; còn lại làm theo phòng (thời gian còn thì làm thêm bài khác).", "",
         bang(["Phòng / vai", "Bài làm", "Kịch bản B4 gợi ý (Phụ lục)"], N.PHAN_CONG), ""]
    for ma, ten in N.NHOM:
        ds = [b for b in N.BAI if b[1] == ma]
        L += [f"## 4.{'ABCDE'.index(ma) + 1} {ten}", ""]
        for b in ds:
            L += [f"### {b[0]} — {b[2]}", "",
                  f"- **Ai làm:** {b[3]}",
                  f"- **Làm / gõ gì:** {b[4]}",
                  f"- **Đạt khi:** {b[5]}"]
            if b[6]:
                L.append(f"- **Lưu ý:** {b[6]}")
            L += ["- [ ] Đúng  [ ] Một phần  [ ] Sai  [ ] Bịa — ghi chú: ……", ""]
    L += ["## 5. Báo lỗi thế nào", "",
          "1. Ngay dưới câu trả lời sai: bấm **👎**, ghi **sai ở đâu + căn cứ đúng**, ví dụ "
          "*\"Dẫn Điều 47 LDN là đúng nhưng thiếu Điều 35 về thủ tục chuyển quyền sở hữu tài sản góp vốn\"*.",
          "2. Ghi mã bài + kết quả vào phiếu; chụp màn hình bài **Sai / Bịa**, đặt tên `<mã bài>_<tên bạn>.png`.",
          "3. Báo **ngay** cho admin (không đợi hết buổi) khi: thấy tài liệu của khách khác lẫn vào câu trả lời; "
          "bot làm theo câu lệnh nằm trong file đính kèm; bot trả lời được câu hỏi về hồ sơ mà tài khoản của bạn không được xem.",
          "4. Bot chậm > 2 phút, báo *network error*, trang trắng → báo IT kèm mã bản ở chân trang (bài A3).", "",
          "## 6. Những điều không làm khi thử", "",
          "- Không dán hồ sơ khách thật chưa che tên vào khung chat khi chưa cần; không dùng kênh chat **trên website** để thử hồ sơ nội bộ.",
          "- Không ép bot trả lời câu nó đã báo 🔴 *Không tìm thấy căn cứ* bằng cách hỏi vòng — làm vậy chỉ tăng khả năng nhận câu bịa.",
          "- Không gửi ra ngoài bất kỳ file nào bot tạo trong buổi thử; file bot tạo tự xoá sau 7 ngày.", "",
          "## Phụ lục — Câu hỏi cho bài B4 (trích *Kịch bản test AI* của HDS)", "",
          "Dán **nguyên văn** cột *Câu gửi bot* vào một cuộc trò chuyện mới. Cột *Phải dẫn* là căn cứ HDS đặt làm tiêu chí — "
          "kho đang dùng bản luật mới nhất nên số điều có thể lệch bản 2020; chấm theo bản luật công ty đã chốt.", "",
          bang(["Mã", "Lĩnh vực · Chủ đề", "Câu gửi bot", "Phải dẫn"],
               [(k["id"], f"{k['linh_vuc']} · {k['tieu_de']}", k["cau_hoi"], k["dieu"]) for k in KB_B4]), ""]
    return "\n".join(L)


def build_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule

    F = "Arial"
    navy = "1F3A5F"
    fn, fb = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    fh, ft = Font(name=F, size=10, bold=True, color="FFFFFF"), Font(name=F, size=14, bold=True, color=navy)
    fnote = Font(name=F, size=9, italic=True, color="595959")
    head = PatternFill("solid", fgColor=navy)
    vang = PatternFill("solid", fgColor="FFF2CC")
    xam = PatternFill("solid", fgColor="EDEDED")
    s = Side(style="thin", color="BFBFBF")
    bd = Border(left=s, right=s, top=s, bottom=s)
    wrap = Alignment(wrap_text=True, vertical="top")
    ctr = Alignment(horizontal="center", vertical="top", wrap_text=True)

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    def header(ws, row, cols, widths):
        for j, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(row=row, column=j, value=h)
            c.font, c.fill, c.alignment, c.border = fh, head, ctr, bd
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[row].height = 30

    # ---- Phiếu
    ws = wb.active
    ws.title = "Phiếu"
    ws["A1"] = "PHIẾU DÙNG THỬ TRỢ LÝ AI — NHÂN VIÊN HDS"
    ws["A1"].font = ft
    ws["A2"] = "Điền các ô nền vàng. Hướng dẫn từng bài: KIEM_THU_CHO_NHAN_VIEN.md. File đầu vào: thư mục du-lieu-mau."
    ws["A2"].font = fnote
    info = [("Họ tên", "C3"), ("Phòng", "C4"), ("Vai", "C5"), ("Ngày thử", "C6"), ("Mã bản (chân trang)", "C7")]
    for i, (k, _) in enumerate(info, 3):
        ws.cell(row=i, column=2, value=k).font = fb
        c = ws.cell(row=i, column=3)
        c.fill, c.border = vang, bd
    cols = ["Mã", "Nhóm", "Bài", "Ai làm", "Làm / gõ gì", "Đạt khi", "Kết quả", "Thời gian (giây)", "Ghi chú (sai ở đâu, căn cứ đúng)"]
    widths = [6, 18, 26, 18, 60, 60, 12, 10, 40]
    header(ws, 9, cols, widths)
    r0 = 10
    for i, b in enumerate(N.BAI):
        vals = [b[0], TEN_NHOM[b[1]], b[2], b[3], plain(b[4]), plain(b[5] + ((" — Lưu ý: " + b[6]) if b[6] else "")), None, None, None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r0 + i, column=j, value=v)
            c.font, c.alignment, c.border = fn, wrap, bd
            if j >= 7:
                c.fill = vang
    last = r0 + len(N.BAI) - 1
    dv = DataValidation(type="list", formula1='"Đúng,Một phần,Sai,Bịa,Bỏ qua"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G{r0}:G{last}")
    for val, color in (("Đúng", "C6EFCE"), ("Một phần", "FFEB9C"), ("Sai", "FFC7CE"), ("Bịa", "FF7C80")):
        ws.conditional_formatting.add(f"G{r0}:G{last}", CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                                     fill=PatternFill("solid", fgColor=color)))
    ws.freeze_panes = f"D{r0}"
    ws.auto_filter.ref = f"A9:I{last}"
    # tổng hợp ngay trên phiếu
    t = last + 2
    ws.cell(row=t, column=3, value="TỔNG HỢP").font = fb
    for k, kq in enumerate(("Đúng", "Một phần", "Sai", "Bịa", "Bỏ qua"), 1):
        ws.cell(row=t + k, column=3, value=kq).font = fn
        ws.cell(row=t + k, column=4, value=f'=COUNTIF($G${r0}:$G${last},"{kq}")').font = fn
    ws.cell(row=t + 6, column=3, value="Chưa làm").font = fn
    ws.cell(row=t + 6, column=4, value=f"={len(N.BAI)}-SUM(D{t + 1}:D{t + 5})").font = fn
    ws.cell(row=t + 7, column=3, value="Tỉ lệ Đúng + Một phần (trên bài đã chấm)").font = fb
    c = ws.cell(row=t + 7, column=4, value=f'=IFERROR((D{t + 1}+D{t + 2})/(D{t + 1}+D{t + 2}+D{t + 3}+D{t + 4}),"—")')
    c.font, c.number_format = fb, "0%"

    # ---- Việc thật (E1/E2 chi tiết)
    we = wb.create_sheet("Việc thật")
    we["A1"] = "THỬ BẰNG VIỆC THẬT (bài E1, E2, D9) — đã che tên khách"
    we["A1"].font = ft
    cols = ["STT", "Loại việc", "Mô tả việc (đã che tên)", "Luật sư đã kết luận / làm gì", "Bot trả lời / làm gì",
            "Kết quả", "Điểm bot nêu thêm mà luật sư bỏ sót", "Ước % hoàn thiện (soạn thảo)", "Tiết kiệm (phút)"]
    header(we, 3, cols, [5, 18, 40, 40, 40, 12, 30, 12, 10])
    loai = DataValidation(type="list", formula1='"Câu hỏi khách,Soạn văn bản,Rà soát hợp đồng"', allow_blank=True)
    kq = DataValidation(type="list", formula1='"Đúng,Một phần,Sai,Bịa"', allow_blank=True)
    we.add_data_validation(loai)
    we.add_data_validation(kq)
    ex = [None, "Câu hỏi khách", "VÍ DỤ — Khách hỏi thời hạn góp vốn khi góp bằng quyền sử dụng đất (xoá dòng này)",
          "90 ngày, Điều 47 LDN; thủ tục chuyển quyền Điều 35", "Nêu 90 ngày Điều 47, thiếu Điều 35", "Một phần",
          "Nhắc hậu quả điều chỉnh vốn điều lệ", None, 15]
    for j, v in enumerate(ex, 1):
        c = we.cell(row=4, column=j, value=v)
        c.font, c.fill, c.alignment, c.border = Font(name=F, size=10, italic=True, color="7F7F7F"), xam, wrap, bd
    for r in range(5, 15):
        we.cell(row=r, column=1, value=r - 4).font = fn
        for j in range(1, 10):
            c = we.cell(row=r, column=j)
            c.border, c.alignment = bd, wrap
            if j >= 2:
                c.fill = vang
    loai.add("B5:B14")
    kq.add("F5:F14")

    # ---- Cách chấm + phân công
    wc = wb.create_sheet("Cách chấm & phân công")
    wc["A1"] = "CÁCH CHẤM"
    wc["A1"].font = ft
    header(wc, 2, ["Kết quả", "Khi nào"], [16, 100])
    for i, (a, b) in enumerate(N.CHAM, 3):
        for j, v in enumerate((a, b), 1):
            c = wc.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = fn, wrap, bd
    r = 3 + len(N.CHAM) + 1
    wc.cell(row=r, column=1, value="PHÂN CÔNG THEO PHÒNG").font = ft
    header(wc, r + 1, ["Phòng / vai", "Bài làm", "Kịch bản B4 gợi ý"], [16, 100, 30])
    for i, row in enumerate(N.PHAN_CONG, r + 2):
        for j, v in enumerate(row, 1):
            c = wc.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = fn, wrap, bd
    wq = wb.create_sheet("Câu hỏi B4", 1)
    wq["A1"] = "CÂU HỎI CHO BÀI B4 — dán nguyên văn cột C vào một cuộc trò chuyện mới"
    wq["A1"].font = ft
    header(wq, 2, ["Mã", "Lĩnh vực · Chủ đề", "Câu gửi bot", "Phải dẫn (tiêu chí HDS)"], [7, 30, 90, 40])
    for i, k in enumerate(KB_B4, 3):
        for j, v in enumerate((k["id"], f"{k['linh_vuc']} · {k['tieu_de']}", k["cau_hoi"], k["dieu"]), 1):
            c = wq.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = fn, wrap, bd
    wq.freeze_panes = "C3"
    for w in wb.worksheets:
        w.sheet_view.zoomScale = 90
    wb.save(path)


if __name__ == "__main__":
    md = build_md()
    open(os.path.join(OUT, "KIEM_THU_CHO_NHAN_VIEN.md"), "w", encoding="utf-8", newline="\n").write(md)
    build_xlsx(os.path.join(OUT, "PHIEU_KIEM_THU_NHAN_VIEN.xlsx"))
    print("bài", len(N.BAI), "| md", len(md))
